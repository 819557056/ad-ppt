"""Build an editable presentation directly from page content, without OCR.

This is intentionally a *new* export mode. It does not pretend that a generated
full-slide bitmap can be converted back into editable PowerPoint objects.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(17, 29, 58)
INK = RGBColor(34, 48, 75)
MUTED = RGBColor(94, 107, 130)
PAPER = RGBColor(247, 249, 253)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(58, 91, 225)
TEAL = RGBColor(21, 174, 164)
AMBER = RGBColor(236, 168, 69)
FONT = "Microsoft YaHei"


@dataclass(frozen=True)
class SlidePlan:
    title: str
    points: tuple[str, ...]
    section: str
    source_page_id: str


def _clean(value: object) -> str:
    text = re.sub(r"!\[[^]]*\]\([^)]*\)", "", str(value or ""))
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _point_text(value: object) -> str:
    if isinstance(value, dict):
        for key in ("text", "content", "point", "title"):
            if value.get(key):
                return _clean(value[key])
        return ""
    return _clean(value)


def plans_from_pages(pages: list) -> list[SlidePlan]:
    """Use persisted outline/description data as the structured source of truth."""
    plans = []
    for index, page in enumerate(pages, 1):
        outline = page.get_outline_content() or {}
        description = page.get_description_content() or {}
        if not isinstance(outline, dict):
            outline = {}
        if not isinstance(description, dict):
            description = {}
        title = _clean(outline.get("title")) or f"第 {index} 页"
        raw_points = outline.get("points") or []
        if isinstance(raw_points, str):
            raw_points = raw_points.splitlines()
        if not isinstance(raw_points, list):
            raw_points = []
        points = tuple(filter(None, (_point_text(item) for item in raw_points)))
        if not points:
            desc = description.get("text") or ""
            # Descriptions may contain image-generation instructions. Only use
            # them when the page has no outline points, and keep them concise.
            points = tuple(filter(None, (_clean(s) for s in re.split(r"[。；;\n]", desc))))
        if len(points) > 8:
            raise ValueError(f"Page {index} has {len(points)} points; structured export supports at most 8 per slide")
        plans.append(SlidePlan(title, points, _clean(page.part), page.id))
    return plans


def _rect(slide, x, y, w, h, color, *, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _text(slide, value, x, y, w, h, *, size, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def create_structured_pptx(project, pages: list, output_path: str | Path) -> list[SlidePlan]:
    """Create native PowerPoint text/shapes; never call MinerU, Baidu or an LLM."""
    plans = plans_from_pages(pages)
    if not plans:
        raise ValueError("No pages selected for structured export")

    prs = Presentation()
    ratio = (project.image_aspect_ratio or "16:9").strip()
    width = 10.0 if ratio == "4:3" else 13.333
    height = 7.5
    prs.slide_width = Inches(width)
    prs.slide_height = Inches(height)
    prs.core_properties.author = "banana-slides"
    prs.core_properties.title = _clean(project.project_title) or "Editable presentation"
    project_title = _clean(project.project_title) or "BANANA SLIDES"

    for index, plan in enumerate(plans, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _rect(slide, 0, 0, width, height, PAPER)
        _rect(slide, 0, 0, width, 0.11, BLUE)
        _rect(slide, 0.62, 0.53, 0.09, 0.42, TEAL)
        kicker = plan.section or project_title
        _text(slide, kicker.upper(), 0.86, 0.43, width - 1.5, 0.32, size=11, color=MUTED, bold=True)
        title_size = 30 if len(plan.title) < 35 else 25
        _text(slide, plan.title, 0.68, 1.03, width - 1.36, 0.88, size=title_size, color=NAVY, bold=True)
        _rect(slide, 0.68, 2.02, width - 1.36, 0.025, RGBColor(219, 226, 240))

        points = plan.points or ("此页尚无正文，请在 PowerPoint 中编辑。",)
        count = len(points)
        # Two-column cards keep short outlines readable; denser pages use
        # full-width rows and auto-fit to avoid clipped text.
        columns = 2 if count <= 6 and width > 11 else 1
        rows = (count + columns - 1) // columns
        content_top, content_bottom = 2.27, 6.82
        gap_x, gap_y = 0.18, 0.15
        card_w = (width - 1.36 - gap_x * (columns - 1)) / columns
        card_h = (content_bottom - content_top - gap_y * (rows - 1)) / rows
        for point_index, point in enumerate(points):
            column, row = point_index % columns, point_index // columns
            x = 0.68 + column * (card_w + gap_x)
            y = content_top + row * (card_h + gap_y)
            _rect(slide, x, y, card_w, card_h, WHITE, rounded=True)
            _rect(slide, x + 0.17, y + 0.2, 0.055, max(0.17, card_h - 0.4),
                  (BLUE, TEAL, AMBER)[point_index % 3])
            font_size = 17 if len(point) < 95 else 14 if len(point) < 220 else 11
            _text(slide, point, x + 0.37, y + 0.16, card_w - 0.58, card_h - 0.32,
                  size=font_size, color=INK)

        _text(slide, project_title, 0.68, 7.08, width - 2.1, 0.2, size=9, color=MUTED)
        _text(slide, f"{index:02d} / {len(plans):02d}", width - 1.6, 7.08, 0.9, 0.2,
              size=9, color=MUTED, align=PP_ALIGN.RIGHT)

    prs.save(str(output_path))
    return plans


def inspect_pptx(path: str | Path, *, expected_slides: int | None = None,
                 require_editable: bool = False) -> dict:
    """Read back the artifact and verify OOXML container and native objects."""
    path = Path(path)
    if not zipfile.is_zipfile(path):
        raise ValueError("Export is not a valid PPTX ZIP package")
    with zipfile.ZipFile(path) as archive:
        broken_member = archive.testzip()
        if broken_member:
            raise ValueError(f"Corrupt PPTX member: {broken_member}")
    prs = Presentation(str(path))
    if expected_slides is not None and len(prs.slides) != expected_slides:
        raise ValueError(f"Expected {expected_slides} slides, got {len(prs.slides)}")
    text_counts = []
    shape_counts = []
    image_counts = []
    warnings = []
    for index, slide in enumerate(prs.slides, 1):
        text_count = sum(bool(shape.has_text_frame and shape.text.strip()) for shape in slide.shapes)
        image_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        text_counts.append(text_count)
        shape_counts.append(len(slide.shapes))
        image_counts.append(image_count)
        if require_editable and not text_count:
            raise ValueError(f"Slide {index} has no editable text")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 10000 or shape.top + shape.height > prs.slide_height + 10000:
                warnings.append(f"Slide {index} contains an out-of-bounds shape")
                break
    return {
        "slide_count": len(prs.slides),
        "editable_text_shapes": sum(text_counts),
        "native_shape_count": sum(shape_counts) - sum(image_counts),
        "image_shapes": sum(image_counts),
        "warnings": warnings,
    }
