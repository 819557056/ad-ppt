"""Render an immutable Scene snapshot to native-text PPTX."""
from io import BytesIO
from services.scene.telemetry import timed_stage

import posixpath
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt, Centipoints

from services.scene.store import checked_bytes
from services.scene.layout_contract import soft_breaks as layout_soft_breaks, validate_layout
from services.exports.report import page_inventory


def _color(value):
    return RGBColor.from_string(value.lstrip('#'))


def _picture_bytes(asset, element):
    image = Image.open(BytesIO(checked_bytes(asset))).convert('RGBA')
    crop = element.get('crop', {'x': 0, 'y': 0, 'w': 1, 'h': 1})
    left = round(crop['x'] * image.width)
    top = round(crop['y'] * image.height)
    right = round((crop['x'] + crop['w']) * image.width)
    bottom = round((crop['y'] + crop['h']) * image.height)
    image = image.crop((left, top, max(left + 1, right), max(top + 1, bottom)))
    if element.get('opacity', 1) != 1:
        image.putalpha(image.getchannel('A').point(lambda v: round(v * element['opacity'])))
    stream = BytesIO()
    image.save(stream, format='PNG')
    stream.seek(0)
    return stream, image.size


def _add_picture(slide, element, asset):
    frame = element['frame']
    stream, (iw, ih) = _picture_bytes(asset, element)
    x, y, w, h = (frame[k] for k in ('x', 'y', 'w', 'h'))
    if element.get('fit') == 'contain':
        ratio = min(w / iw, h / ih)
        pw, ph = iw * ratio, ih * ratio
        x += (w - pw) / 2
        y += (h - ph) / 2
        shape = slide.shapes.add_picture(stream, Pt(x), Pt(y), Pt(pw), Pt(ph))
    else:
        shape = slide.shapes.add_picture(stream, Pt(x), Pt(y), Pt(w), Pt(h))
        source_ratio, box_ratio = iw / ih, w / h
        if source_ratio > box_ratio:
            cut = (1 - box_ratio / source_ratio) / 2
            shape.crop_left = shape.crop_right = cut
        else:
            cut = (1 - source_ratio / box_ratio) / 2
            shape.crop_top = shape.crop_bottom = cut
    shape.name = f"scene:{element['id']}"
    shape.rotation = frame.get('rotation_deg', 0)


def _add_text(slide, element, measured=None):
    soft_breaks = layout_soft_breaks(measured) if measured else ()
    frame = element['frame']
    shape = slide.shapes.add_textbox(Pt(frame['x']), Pt(frame['y']), Pt(frame['w']), Pt(frame['h']))
    shape.name = f"scene:{element['id']}"
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = measured is None
    tf.auto_size = MSO_AUTO_SIZE.NONE
    style = element['style']
    padding = measured['padding_pt'] if measured else style['padding_pt']
    tf.margin_left = tf.margin_right = Pt(padding)
    tf.margin_top = tf.margin_bottom = Pt(padding)
    tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE,
                          'bottom': MSO_ANCHOR.BOTTOM}[style['vertical_align']]
    # An explicit <a:br> is an editable soft break. A control character in a
    # run would instead be escaped as literal `_x000B_` by python-pptx.
    source_offset = 0
    for index, line in enumerate(element['text'].split('\n')):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                               'right': PP_ALIGN.RIGHT}[style['align']]
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        # OOXML spcPts has 1/100 pt precision, unlike Chromium CSS subpixels.
        paragraph.line_spacing = Centipoints(round(measured['line_height_pt'] * 100)) if measured else style['line_height']
        # Empty paragraphs still need the same font metrics, not the theme size.
        paragraph.font.name = 'Noto Sans CJK SC'
        paragraph.font.size = Pt(style['font_size_pt'])
        paragraph.font.bold = style['font_weight'] == 700
        line_breaks = [offset - source_offset for offset in soft_breaks
                       if source_offset < offset < source_offset + len(line)]
        starts = [0, *line_breaks]
        ends = [*line_breaks, len(line)]
        for part_index, (start, end) in enumerate(zip(starts, ends)):
            if part_index:
                paragraph.add_line_break()
            run = paragraph.add_run()
            run.text = line[start:end]
            # Match the bundled font's internal family, not its historical file name.
            # Office resolves this name on the recipient's machine.
            run.font.name = 'Noto Sans CJK SC'
            # python-pptx writes only a:latin. Office otherwise resolves CJK
            # characters through the theme's East Asian font, even when the
            # Latin typeface is installed on the recipient's computer.
            east_asian = OxmlElement('a:ea')
            east_asian.set('typeface', 'Noto Sans CJK SC')
            run._r.get_or_add_rPr().append(east_asian)
            run.font.size = Pt(style['font_size_pt'])
            run.font.bold = style['font_weight'] == 700
            run.font.color.rgb = _color(style['color'])
        source_offset += len(line) + 1


@timed_stage('pptx_compile')
def render_pptx(snapshot, revisions, assets, text_layout=None):
    if text_layout is not None:
        text_layout = validate_layout(text_layout, snapshot, revisions)
    manifest = snapshot.manifest_json
    presentation = Presentation()
    presentation.slide_width = Pt(manifest['canvas']['width_pt'])
    presentation.slide_height = Pt(manifest['canvas']['height_pt'])
    layout = presentation.slide_layouts[6]
    for page_index, item in enumerate(manifest['pages']):
        scene = revisions[item['revision_id']].scene_json
        slide = presentation.slides.add_slide(layout)
        bg = scene['background']
        if bg['kind'] == 'solid':
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _color(bg['color'])
        else:
            _add_picture(slide, {'id': f"background:{item['page_id']}",
                                 'frame': {'x': 0, 'y': 0, 'w': scene['canvas']['width_pt'],
                                           'h': scene['canvas']['height_pt'], 'rotation_deg': 0},
                                 'crop': bg['crop'], 'fit': 'cover', 'opacity': 1}, assets[bg['asset_id']])
        for element in scene['elements']:
            if element['kind'] == 'text':
                breaks = (text_layout['pages'][page_index][element['id']]
                          if text_layout is not None else None)
                _add_text(slide, element, breaks)
            else:
                _add_picture(slide, element, assets[element['asset_id']])
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _verify_native_text(shape, element, measured=None):
    frame, style = element['frame'], element['style']
    if shape.rotation != 0:
        raise ValueError('PPTX text rotation mismatch')
    if any(abs(actual - Pt(frame[key])) > 2 for actual, key in (
            (shape.left, 'x'), (shape.top, 'y'), (shape.width, 'w'), (shape.height, 'h'))):
        raise ValueError('PPTX text frame geometry mismatch')
    text_frame = shape.text_frame
    if (text_frame.auto_size != MSO_AUTO_SIZE.NONE or text_frame.word_wrap is not (measured is None) or
            text_frame.vertical_anchor != {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE,
                                           'bottom': MSO_ANCHOR.BOTTOM}[style['vertical_align']] or
            any(abs(margin - Pt(style['padding_pt'])) > 2 for margin in (
                text_frame.margin_left, text_frame.margin_right,
                text_frame.margin_top, text_frame.margin_bottom))):
        raise ValueError('PPTX text frame layout mismatch')
    alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                 'right': PP_ALIGN.RIGHT}[style['align']]
    for paragraph in text_frame.paragraphs:
        if (paragraph.alignment != alignment or
                paragraph.line_spacing != (Centipoints(round(measured['line_height_pt'] * 100)) if measured else style['line_height']) or
                paragraph.space_before != Pt(0) or paragraph.space_after != Pt(0)):
            raise ValueError('PPTX paragraph style mismatch')
        for run in paragraph.runs:
            properties = run._r.rPr
            east_asian = properties.find(qn('a:ea')) if properties is not None else None
            try:
                font_rgb = run.font.color.rgb
            except (AttributeError, TypeError, ValueError):
                font_rgb = None
            if (run.font.name != 'Noto Sans CJK SC' or east_asian is None or
                    east_asian.get('typeface') != 'Noto Sans CJK SC' or
                    run.font.size != Pt(style['font_size_pt']) or
                    run.font.bold != (style['font_weight'] == 700) or
                    font_rgb != _color(style['color'])):
                raise ValueError('PPTX native font mismatch')


def _verify_opc_package(payload):
    relationship_tag = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    with ZipFile(BytesIO(payload)) as package:
        names = package.namelist()
        parts = set(names)
        if (len(names) != len(parts) or '[Content_Types].xml' not in parts or
                '_rels/.rels' not in parts or package.testzip() is not None):
            raise ValueError('PPTX ZIP/OPC package is invalid')
        for name in names:
            if name.startswith('/') or '\\' in name or '..' in name.split('/'):
                raise ValueError('PPTX package path is invalid')
            if not name.endswith('.rels'):
                continue
            if name == '_rels/.rels':
                source_directory = ''
            else:
                directory, leaf = posixpath.split(name)
                if posixpath.basename(directory) != '_rels' or not leaf.endswith('.rels'):
                    raise ValueError('PPTX relationship path is invalid')
                source_directory = posixpath.dirname(directory)
                source = posixpath.join(source_directory, leaf[:-5])
                if source not in parts:
                    raise ValueError('PPTX relationship source is missing')
            root = fromstring(package.read(name))
            if root.tag != relationship_tag + 'Relationships':
                raise ValueError('PPTX relationship XML is invalid')
            seen_ids = set()
            for relation in root:
                if relation.tag != relationship_tag + 'Relationship':
                    raise ValueError('PPTX relationship XML is invalid')
                relation_id, target = relation.get('Id'), relation.get('Target')
                if not relation_id or relation_id in seen_ids or not target:
                    raise ValueError('PPTX relationship entry is invalid')
                seen_ids.add(relation_id)
                if relation.get('TargetMode', 'Internal') != 'Internal':
                    raise ValueError('PPTX external relationship is forbidden')
                parsed = urlsplit(target)
                if parsed.scheme or parsed.netloc or parsed.query or not parsed.path:
                    raise ValueError('PPTX relationship target is invalid')
                path = unquote(parsed.path)
                resolved = posixpath.normpath(path.lstrip('/') if path.startswith('/') else
                    posixpath.join(source_directory, path))
                if resolved.startswith('../') or resolved not in parts:
                    raise ValueError('PPTX relationship target is missing')


def _verify_picture(shape, element, asset=None):
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        raise ValueError('PPTX image is not a picture object')
    frame = element['frame']
    if abs(shape.rotation - frame.get('rotation_deg', 0)) > .01:
        raise ValueError('PPTX image rotation mismatch')
    x, y, width, height = (frame[key] for key in ('x', 'y', 'w', 'h'))
    crops = (0, 0, 0, 0)
    if asset is not None:
        source, (image_width, image_height) = _picture_bytes(asset, element)
        if shape.image.blob != source.getvalue():
            raise ValueError('PPTX image asset bytes mismatch')
        if element.get('fit') == 'contain':
            ratio = min(width / image_width, height / image_height)
            placed_width, placed_height = image_width * ratio, image_height * ratio
            x += (width - placed_width) / 2
            y += (height - placed_height) / 2
            width, height = placed_width, placed_height
        else:
            source_ratio, box_ratio = image_width / image_height, width / height
            if source_ratio > box_ratio:
                cut = (1 - box_ratio / source_ratio) / 2
                crops = (cut, 0, cut, 0)
            else:
                cut = (1 - source_ratio / box_ratio) / 2
                crops = (0, cut, 0, cut)
    elif element.get('fit') == 'contain':
        # Exact contain geometry depends on source dimensions; it is checked
        # when the worker supplies the frozen asset bytes.
        x, y = shape.left / 12700, shape.top / 12700
        width, height = shape.width / 12700, shape.height / 12700
        if (x < frame['x'] - .01 or y < frame['y'] - .01 or
                x + width > frame['x'] + frame['w'] + .01 or
                y + height > frame['y'] + frame['h'] + .01):
            raise ValueError('PPTX image frame mismatch')
    if any(abs(actual - Pt(expected)) > 2 for actual, expected in (
            (shape.left, x), (shape.top, y), (shape.width, width), (shape.height, height))):
        raise ValueError('PPTX image frame mismatch')
    if asset is not None and any(abs(actual - expected) > .00002 for actual, expected in zip(
            (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom), crops)):
        raise ValueError('PPTX image crop mismatch')


@timed_stage('pptx_verify')
def verify_pptx(payload, snapshot, revisions, text_layout=None, assets=None):
    if text_layout is not None:
        text_layout = validate_layout(text_layout, snapshot, revisions)
    _verify_opc_package(payload)
    presentation = Presentation(BytesIO(payload))
    pages = snapshot.manifest_json['pages']
    canvas = snapshot.manifest_json['canvas']
    if (abs(presentation.slide_width - Pt(canvas['width_pt'])) > 2 or
            abs(presentation.slide_height - Pt(canvas['height_pt'])) > 2):
        raise ValueError('PPTX canvas size mismatch')
    if len(presentation.slides) != len(pages):
        raise ValueError('PPTX page count mismatch')
    page_reports = []
    for page_index, (slide, item) in enumerate(zip(presentation.slides, pages)):
        scene = revisions[item['revision_id']].scene_json
        background = scene['background']
        expected_order = ([f"scene:background:{item['page_id']}"]
                          if background['kind'] == 'image' else [])
        expected_order += [f"scene:{element['id']}" for element in scene['elements']]
        if [shape.name for shape in slide.shapes] != expected_order:
            raise ValueError('PPTX scene shape order or mapping mismatch')
        by_name = {shape.name: shape for shape in slide.shapes}
        if background['kind'] == 'solid':
            if (slide.background.fill.type != MSO_FILL.SOLID or
                    slide.background.fill.fore_color.rgb != _color(background['color'])):
                raise ValueError('PPTX solid background mismatch')
        if scene['background']['kind'] == 'image':
            if slide.background.fill.type != MSO_FILL.BACKGROUND:
                raise ValueError('PPTX image background underlay mismatch')
            background_element = {'frame': {'x': 0, 'y': 0,
                'w': scene['canvas']['width_pt'], 'h': scene['canvas']['height_pt'],
                'rotation_deg': 0}, 'crop': background['crop'], 'fit': 'cover', 'opacity': 1}
            _verify_picture(by_name[f"scene:background:{item['page_id']}"], background_element,
                assets[background['asset_id']] if assets is not None else None)
        for element in scene['elements']:
            shape = by_name.get(f"scene:{element['id']}")
            if shape is None:
                raise ValueError('PPTX missing scene element')
            if element['kind'] == 'text':
                expected = element['text']
                if text_layout is not None:
                    for offset in reversed(layout_soft_breaks(text_layout['pages'][page_index][element['id']])):
                        expected = expected[:offset] + '\v' + expected[offset:]
                if not shape.has_text_frame or (shape.text if text_layout is not None else
                        shape.text.replace('\v', '')) != expected:
                    raise ValueError('PPTX native text mismatch')
                _verify_native_text(shape, element, text_layout['pages'][page_index][element['id']] if text_layout else None)
            if element['kind'] == 'image':
                _verify_picture(shape, element,
                    assets[element['asset_id']] if assets is not None else None)
        page_reports.append(page_inventory(item, scene))
    return {'page_count': len(pages), 'native_text_count': sum(p['text_count'] for p in page_reports),
            'opc_relationships_checked': True, 'native_text_style_checked': True,
            'image_asset_bytes_checked': assets is not None, 'pages': page_reports}
