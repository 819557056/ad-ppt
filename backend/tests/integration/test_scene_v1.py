"""Offline Scene API regression tests; no paid Provider calls."""
from io import BytesIO
import base64
import hashlib
import hmac
import os
import re
import tarfile
from datetime import timedelta
from types import SimpleNamespace
from uuid import uuid4
import time
import threading

from models import db, Page, Project
from ops.scene_asset_audit import audit_scene
from ops.scene_volume_archive import archive_tree, restore_tree
from models.scene_v1 import (ApiIdempotencyRecord, Asset, DeckSnapshot, GenerationPlan, ModelCredential, PageSceneRevision, Principal, RevisionAsset, SceneExport, TemplateDocument,
                             SnapshotPage,
                             SceneCandidate, SceneTaskAttempt, SceneTaskItem, SceneWorkerHeartbeat, now)
from services.exports.scene_pptx import render_pptx, verify_pptx
from services.exports.quality_gate import validate_text_fit
from services.scene.validation import digest
from services.scene.validation import SlideScene, blank_scene, normalize_scene
from services.scene.commands import CommandError, assert_locked_preserved
from services.scene.generation import _edit_scene
from services.scene.openapi import route_contracts
from services.scene.render_jobs import _read_render_output, run_render_job
import json
from pathlib import Path
from services.scene.versioning import LOCAL_OWNER_ID, SceneError, accept_candidates
from PIL import Image
import fitz
import pytest
from services.exports.scene_pdf import measure_text_layout, verify_pdf
from services.templates.style_profile import normalize_style_profile
from workers.scene_runner import WORKER_ID, claim, renew_lease, run_once


def test_volume_archive_roundtrip_and_path_traversal_rejection(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    (source / 'nested').mkdir()
    (source / 'nested' / '数据.txt').write_text('Scene asset 123', encoding='utf-8')
    archive = tmp_path / 'volume.tar'
    archive_tree(source, archive)
    target = tmp_path / 'target'
    target.mkdir()
    restore_tree(archive, target)
    assert (target / 'nested' / '数据.txt').read_text(encoding='utf-8') == 'Scene asset 123'
    bad = tmp_path / 'bad.tar'
    with tarfile.open(bad, 'w') as output:
        entry = tarfile.TarInfo('../escape.txt')
        entry.size = 1
        output.addfile(entry, BytesIO(b'x'))
    empty = tmp_path / 'empty'
    empty.mkdir()
    with pytest.raises(ValueError, match='Unsafe archive member'):
        restore_tree(bad, empty)
    assert not (tmp_path / 'escape.txt').exists()
    with pytest.raises(ValueError, match='already exists'):
        archive_tree(source, archive)
    with pytest.raises(ValueError, match='empty'):
        restore_tree(archive, target)
    try:
        (source / 'linked').symlink_to(source / 'nested', target_is_directory=True)
    except (OSError, NotImplementedError):
        pass  # Windows test accounts may not have symlink privilege.
    else:
        with pytest.raises(ValueError, match='Symlinks'):
            archive_tree(source, tmp_path / 'with-link.tar')


def test_scene_openapi_covers_live_routes_and_access_errors(client, app, monkeypatch):
    expected = route_contracts()
    actual = {(method, rule.rule): rule.endpoint.rsplit('.', 1)[-1]
              for rule in app.url_map.iter_rules() if rule.endpoint.startswith('scene_v2.')
              for method in rule.methods if method in ('GET', 'POST', 'PATCH', 'DELETE')}
    assert set(actual) == set(expected)
    assert all(actual[key] == meta['endpoint'] for key, meta in expected.items())
    monkeypatch.setenv('ACCESS_CODE', 'scene-contract-test')
    denied = client.get('/api/v2/openapi.json')
    assert denied.status_code == 401
    assert denied.json['error']['code'] == 'AUTH_REQUIRED'
    assert denied.json['error']['stage'] == 'access'
    monkeypatch.setitem(app.config, 'SCENE_EDITOR_ENABLED', True)
    headers = {'X-Access-Code': 'scene-contract-test'}
    response = client.get('/api/v2/openapi.json', headers=headers)
    assert response.status_code == 200
    spec = response.json
    assert spec['openapi'] == '3.1.0'
    assert sum(len(methods) for methods in spec['paths'].values()) == len(expected)
    for (method, route), meta in expected.items():
        path = route.replace('<', '{').replace('>', '}')
        operation = spec['paths'][path][method.lower()]
        assert operation['operationId'] == meta['endpoint']
        assert str(meta['status']) in operation['responses']
        if meta['endpoint'] in ('asset_content', 'scene_font'):
            continue
        assert 'request_id' in operation['responses'][str(meta['status'])]['content']['application/json']['schema']['required'] or meta['endpoint'] == 'scene_openapi'
        concrete = re.sub(r'<[^>]+>', str(uuid4()), route)
        blocked = client.open(concrete, method=method)
        assert blocked.status_code == 401, (method, route, blocked.json)
        assert set(blocked.json) == {'error', 'request_id'}
    assert spec['paths']['/api/v2/readiness']['get']['responses']['503']['content']['application/json']
    assert spec['paths']['/api/v2/projects/{project_id}/pages/{page_id}/scene']['patch']['requestBody']
    page_body = spec['paths']['/api/v2/projects/{project_id}/pages']['post']['requestBody'][
        'content']['application/json']['schema']
    assert {'insert_at', 'copy_from_page_id', 'copy_from_page_version',
            'copy_from_revision_id'} <= set(page_body['properties'])
    for (method, route), meta in expected.items():
        if method in ('POST', 'PATCH', 'DELETE'):
            path = route.replace('<', '{').replace('>', '}')
            operation = spec['paths'][path][method.lower()]
            assert any(parameter['name'] == 'Idempotency-Key' and parameter['required']
                       for parameter in operation['parameters']), meta['endpoint']

    def check_references(node):
        if isinstance(node, dict):
            for name, value in node.items():
                if name == '$ref':
                    assert value.startswith('#/')
                    cursor = spec
                    for part in value[2:].split('/'):
                        cursor = cursor[part.replace('~1', '/').replace('~0', '~')]
                else:
                    check_references(value)
        elif isinstance(node, list):
            for value in node:
                check_references(value)
    check_references(spec)


def test_scene_font_manifest_matches_download_and_pptx_family(client, app, monkeypatch):
    from fontTools.ttLib import TTFont
    from pptx import Presentation

    monkeypatch.setenv('ACCESS_CODE', 'scene-font-test')
    monkeypatch.setitem(app.config, 'SCENE_EDITOR_ENABLED', True)
    denied = client.get('/api/v2/font-manifest')
    assert denied.status_code == 401
    response = client.get('/api/v2/font-manifest', headers={'X-Access-Code': 'scene-font-test'})
    assert response.status_code == 200, response.json
    info = response.json['data']
    assert info['font_manifest_id'] == 'fonts-v1'
    assert info['family'] == 'Noto Sans CJK SC'
    downloaded = client.get(info['download_url'])
    assert downloaded.status_code == 200
    assert 'attachment' in downloaded.headers['Content-Disposition']
    assert downloaded.headers['Content-Type'].startswith('font/otf')
    assert hashlib.sha256(downloaded.data).hexdigest() == info['sha256']
    font = TTFont(BytesIO(downloaded.data))
    try:
        families = {item.toUnicode() for item in font['name'].names if item.nameID == 1}
        assert info['family'] in families
    finally:
        font.close()
    element = text_element('中文 English 123')
    snapshot = SimpleNamespace(manifest_json={'canvas': {'width_pt': 960, 'height_pt': 540},
        'pages': [{'page_id': 'page', 'revision_id': 'revision'}]})
    revisions = {'revision': SimpleNamespace(scene_json={'background': {'kind': 'solid',
        'color': '#FFFFFF'}, 'elements': [element]})}
    pptx = Presentation(BytesIO(render_pptx(snapshot, revisions, {})))
    run = pptx.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == info['family']
    assert f'<a:ea typeface="{info["family"]}"' in run._r.xml


def test_pptx_quality_gate_rejects_native_text_style_or_geometry_tampering():
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    element = text_element('中文 English 123')
    snapshot = SimpleNamespace(manifest_json={'canvas': {'width_pt': 960, 'height_pt': 540},
        'pages': [{'page_id': 'page', 'revision_id': 'revision'}]})
    revisions = {'revision': SimpleNamespace(scene_json={'background': {'kind': 'solid',
        'color': '#FFFFFF'}, 'elements': [element]})}
    original = render_pptx(snapshot, revisions, {})
    assert verify_pptx(original, snapshot, revisions)['native_text_count'] == 1
    for changed_field, expected_error in (
        ('east_asian_font', 'native font'), ('size', 'native font'),
        ('geometry', 'geometry'), ('auto_shrink', 'layout'), ('line_spacing', 'paragraph style')):
        deck = Presentation(BytesIO(original))
        shape = deck.slides[0].shapes[0]
        run = shape.text_frame.paragraphs[0].runs[0]
        if changed_field == 'east_asian_font':
            run._r.rPr.find(qn('a:ea')).set('typeface', 'Arial')
        elif changed_field == 'size':
            run.font.size = Pt(12)
        elif changed_field == 'geometry':
            shape.left += Pt(10)
        elif changed_field == 'auto_shrink':
            shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        else:
            shape.text_frame.paragraphs[0].line_spacing = 1.5
        changed = BytesIO()
        deck.save(changed)
        with pytest.raises(ValueError, match=expected_error):
            verify_pptx(changed.getvalue(), snapshot, revisions)


def test_pptx_quality_gate_rejects_external_or_broken_opc_relationships():
    from xml.etree import ElementTree as ET
    from zipfile import ZIP_DEFLATED, ZipFile

    element = text_element('关系完整性')
    snapshot = SimpleNamespace(manifest_json={'canvas': {'width_pt': 960, 'height_pt': 540},
        'pages': [{'page_id': 'page', 'revision_id': 'revision'}]})
    revisions = {'revision': SimpleNamespace(scene_json={'background': {'kind': 'solid',
        'color': '#FFFFFF'}, 'elements': [element]})}
    original = render_pptx(snapshot, revisions, {})
    with ZipFile(BytesIO(original)) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    relationship_tag = '{http://schemas.openxmlformats.org/package/2006/relationships}'

    def repack(rels):
        changed = BytesIO()
        with ZipFile(changed, 'w', ZIP_DEFLATED) as archive:
            for name, content in members.items():
                archive.writestr(name, ET.tostring(rels) if name == '_rels/.rels' else content)
        return changed.getvalue()

    external = ET.fromstring(members['_rels/.rels'])
    ET.SubElement(external, relationship_tag + 'Relationship', {
        'Id': 'rIdExternal', 'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        'Target': 'https://example.invalid/outside', 'TargetMode': 'External'})
    with pytest.raises(ValueError, match='external relationship'):
        verify_pptx(repack(external), snapshot, revisions)
    missing = ET.fromstring(members['_rels/.rels'])
    missing[0].set('Target', 'missing-part.xml')
    with pytest.raises(ValueError, match='target is missing'):
        verify_pptx(repack(missing), snapshot, revisions)


def test_credential_creation_replays_without_storing_plain_key_fingerprint(client, app):
    create(client, app)
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'k' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    payload = {'label': 'My gateway', 'base_url': 'https://model.example/v1',
               'api_key': 'short-sandbox-secret'}
    headers = key()
    first = client.post('/api/v2/model-credentials', json=payload, headers=headers)
    replayed = client.post('/api/v2/model-credentials', json=payload, headers=headers)
    assert first.status_code == replayed.status_code == 201
    assert first.json['data'] == replayed.json['data']
    with app.app_context():
        assert ModelCredential.query.filter_by(owner_id=LOCAL_OWNER_ID).count() == 1
        record = ApiIdempotencyRecord.query.filter_by(owner_id=LOCAL_OWNER_ID,
            operation='add_model_credential').one()
        assert record.request_hash != digest(payload)
        assert payload['api_key'] not in json.dumps(record.response_json)
    assert client.post('/api/v2/model-credentials', json={**payload, 'api_key': 'another-secret'},
                       headers=headers).status_code == 409
    revoke_url = f'/api/v2/model-credentials/{first.json["data"]["credential_id"]}'
    revoke_key = key()
    revoked = client.delete(revoke_url, headers=revoke_key)
    replay_revoke = client.delete(revoke_url, headers=revoke_key)
    assert revoked.status_code == replay_revoke.status_code == 200
    assert revoked.json['data'] == replay_revoke.json['data']
    assert revoked.json['data']['status'] == 'revoked'


@pytest.mark.parametrize(('field', 'invalid'), [
    ('label', {'not': 'text'}), ('api_key', {'not': 'text'}),
    ('base_url', {'not': 'text'}), ('base_url', 'https://model.example/v1\n'),
    ('api_key', 'bad\nkey'),
])
def test_credential_rejects_non_text_or_control_input(client, app, field, invalid):
    create(client, app)
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'k' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    payload = {'label': 'Mine', 'base_url': 'https://model.example/v1',
               'api_key': 'sandbox-secret'}
    response = client.post('/api/v2/model-credentials',
        json={**payload, field: invalid}, headers=key())
    assert response.status_code == 422
    assert response.json['error']['code'] in {'CREDENTIAL_INVALID', 'GATEWAY_INVALID'}
    with app.app_context():
        assert ModelCredential.query.count() == 0


def test_generation_plan_idempotency_replays_lost_response(client, app):
    project, _ = create(client, app)
    url = f'/api/v2/projects/{project["project_id"]}/generation-plans'
    request_headers = key()
    payload = {'base_project_version': project['project_version']}
    first = client.post(url, json=payload, headers=request_headers)
    changed = client.patch(f'/api/v2/projects/{project["project_id"]}/model-config',
        json={'base_project_version': project['project_version'],
              'model_config': {'text_model': 'later-model'}}, headers=key())
    assert changed.status_code == 200
    again = client.post(url, json=payload, headers=request_headers)
    assert first.status_code == again.status_code == 201
    assert first.json['data'] == again.json['data']
    assert client.post(url, json={'base_project_version': 999}, headers=request_headers).status_code == 409
    assert client.post(url, json=payload).status_code == 422


def test_generation_plan_rejects_changed_page_even_when_project_version_is_unchanged(client, app):
    project, page = create(client, app)
    project_id, page_id = project['project_id'], page['page_id']
    plan_url = f'/api/v2/projects/{project_id}/generation-plans'
    expected = [{'page_id': page_id, 'page_version': page['page_version'],
                 'revision_id': page['revision_id']}]
    saved = client.patch(f'/api/v2/projects/{project_id}/pages/{page_id}/scene',
        json={'base_revision_id': page['revision_id'], 'base_page_version': page['page_version'],
              'commands': [{'op': 'add_element', 'element': text_element()}]}, headers=key())
    assert saved.status_code == 200, saved.json
    stale = client.post(plan_url, json={'base_project_version': project['project_version'],
        'expected_pages': expected}, headers=key())
    assert stale.status_code == 409
    assert stale.json['error']['code'] == 'PAGE_VERSION_CONFLICT'
    current = saved.json['data']
    expected[0].update(page_version=current['page_version'], revision_id=current['revision_id'])
    fresh = client.post(plan_url, json={'base_project_version': project['project_version'],
        'expected_pages': expected}, headers=key())
    assert fresh.status_code == 201, fresh.json


def test_ten_page_plan_enqueues_only_selected_page_and_rejects_malformed_targets(client, app):
    app.config['SCENE_EDITOR_ENABLED'] = True
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'g' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    created = client.post('/api/v2/projects', json={'title': '十页回归',
        'pages': [{'title': f'第 {index + 1} 页', 'points': []} for index in range(10)]}, headers=key())
    assert created.status_code == 201, created.json
    project = created.json['data']
    project_id = project['project_id']
    assert len(project['pages']) == 10
    credential = client.post('/api/v2/model-credentials', json={'label': 'Ten pages',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key())
    assert credential.status_code == 201, credential.json
    configured = client.patch(f'/api/v2/projects/{project_id}/model-config', json={
        'base_project_version': project['project_version'],
        'model_config': {'text_model': 'fixture-text'}}, headers=key())
    assert configured.status_code == 200, configured.json
    plan = client.post(f'/api/v2/projects/{project_id}/generation-plans', json={
        'base_project_version': configured.json['data']['project_version']}, headers=key())
    assert plan.status_code == 201, plan.json
    page = project['pages'][3]
    target = {'page_id': page['page_id'], 'base_revision_id': page['revision_id'],
              'base_page_version': page['page_version']}
    url = f'/api/v2/projects/{project_id}/generation-tasks'
    payload = {'plan_id': plan.json['data']['plan_id'],
               'credential_id': credential.json['data']['credential_id'], 'targets': [target]}
    malformed = client.post(url, json={**payload, 'targets': ['not an object']}, headers=key())
    assert malformed.status_code == 422 and malformed.json['error']['code'] == 'TARGETS_INVALID'
    malformed = client.post(url, json={**payload, 'targets': [{**target, 'page_id': []}]}, headers=key())
    assert malformed.status_code == 422 and malformed.json['error']['code'] == 'TARGETS_INVALID'
    submitted = client.post(url, json=payload, headers=key())
    assert submitted.status_code == 202, submitted.json
    assert [item['page_id'] for item in submitted.json['data']['tasks']] == [page['page_id']]
    with app.app_context():
        queued = SceneTaskItem.query.filter_by(project_id=project_id, operation='generate_scene').all()
        assert len(queued) == 1 and queued[0].page_id == page['page_id']
    changed = patch(client, project, page, [{'op': 'add_element', 'element': text_element('人工先改字')}])
    assert changed.status_code == 200, changed.json
    stale = client.post(url, json=payload, headers=key())
    assert stale.status_code == 409 and stale.json['error']['code'] == 'PAGE_VERSION_CONFLICT'
    with app.app_context():
        assert SceneTaskItem.query.filter_by(project_id=project_id, operation='generate_scene').count() == 1


def test_ai_edit_rejects_unhashable_selection_before_paid_task(client, app):
    project, page = create(client, app)
    response = client.post(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/ai-edits',
        json={'base_revision_id': page['revision_id'], 'base_page_version': page['page_version'],
              'instruction': '修改当前页', 'element_ids': [{'invalid': True}]}, headers=key())
    assert response.status_code == 422 and response.json['error']['code'] == 'SCOPE_INVALID'
    with app.app_context():
        assert SceneTaskItem.query.filter_by(project_id=project['project_id'], operation='ai_edit').count() == 0


def test_outline_and_candidate_inputs_reject_unhashable_ids(client, app):
    project, _ = create(client, app)
    project_id = project['project_id']
    invalid_outline = client.patch(f'/api/v2/projects/{project_id}/outline', json={
        'base_project_version': project['project_version'],
        'pages': [{'page_id': [], 'outline': {'title': 'Invalid', 'points': []}}]}, headers=key())
    assert invalid_outline.status_code == 422
    assert invalid_outline.json['error']['code'] == 'OUTLINE_INVALID'
    invalid_candidates = client.post(f'/api/v2/projects/{project_id}/candidates/accept',
        json={'candidate_ids': [{}]}, headers=key())
    assert invalid_candidates.status_code == 422
    assert invalid_candidates.json['error']['code'] == 'CANDIDATES_INVALID'


def test_candidate_accept_rechecks_archived_project_under_lock(client, app):
    project, page = create(client, app)
    with app.app_context():
        proposal_id, candidate_id = str(uuid4()), str(uuid4())
        revision = db.session.get(PageSceneRevision, page['revision_id'])
        proposal = dict(revision.scene_json)
        proposal['background'] = {'kind': 'solid', 'color': '#ABCDEF'}
        db.session.add(PageSceneRevision(id=proposal_id, owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], page_id=page['page_id'], seq=2,
            parent_revision_id=revision.id, scene_json=proposal, scene_hash=digest(proposal),
            origin='ai_generate', created_by=LOCAL_OWNER_ID))
        db.session.add(SceneCandidate(id=candidate_id, owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], page_id=page['page_id'],
            proposed_revision_id=proposal_id, base_revision_id=revision.id,
            base_page_version=page['page_version'], scope_json={'kind': 'page'}, state='pending'))
        db.session.commit()
        stale_project = db.session.get(Project, project['project_id'])
        stale_project.deleted_at = now()
        db.session.commit()
        with pytest.raises(SceneError) as raised:
            accept_candidates(stale_project, LOCAL_OWNER_ID, [candidate_id])
        assert raised.value.code == 'NOT_FOUND'
        assert db.session.get(SceneCandidate, candidate_id).state == 'pending'


def test_snapshot_then_export_response_loss_reuses_both_resources(client, app):
    project, page = create(client, app)
    project_id = project['project_id']
    snapshot_url = f'/api/v2/projects/{project_id}/snapshots'
    snapshot_body = {'project_version': project['project_version'],
        'pages': [{'page_id': page['page_id'], 'page_version': page['page_version'],
                   'revision_id': page['revision_id']}]}
    snapshot_headers = key()
    first_snapshot = client.post(snapshot_url, json=snapshot_body, headers=snapshot_headers)
    assert first_snapshot.status_code == 201, first_snapshot.json
    # The browser received the snapshot ID, then lost the create-export response.
    export_url = f'/api/v2/projects/{project_id}/exports'
    export_body = {'snapshot_id': first_snapshot.json['data']['snapshot_id'], 'format': 'pdf',
                   'options': {'quality_profile': 'standard'}}
    export_headers = key()
    first_export = client.post(export_url, json=export_body, headers=export_headers)
    assert first_export.status_code == 202, first_export.json
    repeated_snapshot = client.post(snapshot_url, json=snapshot_body, headers=snapshot_headers)
    repeated_export = client.post(export_url, json=export_body, headers=export_headers)
    assert repeated_snapshot.json['data'] == first_snapshot.json['data']
    assert repeated_export.json['data'] == first_export.json['data']
    with app.app_context():
        assert DeckSnapshot.query.filter_by(project_id=project_id).count() == 1
        assert SceneExport.query.filter_by(project_id=project_id).count() == 1
        assert SceneTaskItem.query.filter_by(project_id=project_id, operation='export').count() == 1
    assert client.post(export_url, json={**export_body, 'format': 'pptx'},
                       headers=export_headers).status_code == 409


def test_snapshot_and_export_report_missing_font_or_corrupt_asset(client, app, tmp_path, monkeypatch):
    project, page = create(client, app)
    project_id = project['project_id']
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'assets')
    image = BytesIO(); Image.new('RGB', (20, 20), '#224466').save(image, format='PNG'); image.seek(0)
    uploaded = client.post(f'/api/v2/projects/{project_id}/assets',
        data={'file': (image, 'figure.png')}, content_type='multipart/form-data', headers=key())
    assert uploaded.status_code == 201
    asset_id = uploaded.json['data']['asset_id']
    element = {'id': str(uuid4()), 'kind': 'image', 'role': 'illustration', 'asset_id': asset_id,
        'frame': {'x': 100, 'y': 120, 'w': 300, 'h': 200, 'rotation_deg': 0},
        'crop': {'x': 0, 'y': 0, 'w': 1, 'h': 1}, 'fit': 'contain', 'opacity': 1,
        'locked': False, 'alt_text': ''}
    saved = patch(client, project, page, [{'op': 'add_element', 'element': element}])
    assert saved.status_code == 200
    head = saved.json['data']
    snapshot_url = f'/api/v2/projects/{project_id}/snapshots'
    snapshot_body = {'project_version': project['project_version'], 'pages': [{
        'page_id': page['page_id'], 'page_version': head['page_version'],
        'revision_id': head['revision_id']}]}
    import services.scene.versioning as versioning
    font_dir = versioning.FONT_DIR
    monkeypatch.setattr(versioning, 'FONT_DIR', tmp_path / 'missing-font')
    missing_font = client.post(snapshot_url, json=snapshot_body, headers=key())
    assert missing_font.status_code == 503
    assert missing_font.json['error']['code'] == 'FONT_UNAVAILABLE'
    monkeypatch.setattr(versioning, 'FONT_DIR', font_dir)
    from services.scene.store import path_for
    with app.app_context():
        asset_path = path_for(db.session.get(Asset, asset_id))
        original = asset_path.read_bytes()
        asset_path.write_bytes(b'corrupt image')
    corrupt = client.post(snapshot_url, json=snapshot_body, headers=key())
    assert corrupt.status_code == 503
    assert corrupt.json['error']['code'] == 'ASSET_UNAVAILABLE'
    with app.app_context():
        assert DeckSnapshot.query.filter_by(project_id=project_id).count() == 0
        asset_path.write_bytes(original)
    frozen = client.post(snapshot_url, json=snapshot_body, headers=key())
    assert frozen.status_code == 201
    queued = client.post(f'/api/v2/projects/{project_id}/exports', json={
        'snapshot_id': frozen.json['data']['snapshot_id'], 'format': 'pptx', 'options': {}}, headers=key())
    assert queued.status_code == 202
    asset_path.write_bytes(b'corrupt after snapshot')
    assert run_once()
    task = client.get(f'/api/v2/tasks/{queued.json["data"]["task_id"]}').json['data']
    assert task['state'] == 'failed'
    assert task['error_code'] == 'ASSET_UNAVAILABLE'


def test_snapshot_refuses_corrupt_persisted_scene_revision(client, app):
    project, page = create(client, app)
    with app.app_context():
        revision = db.session.get(PageSceneRevision, page['revision_id'])
        scene = dict(revision.scene_json)
        scene['background'] = {'kind': 'solid', 'color': '#123456'}
        revision.scene_json = scene
        db.session.commit()
    frozen = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': project['project_version'], 'pages': [{
            'page_id': page['page_id'], 'page_version': page['page_version'],
            'revision_id': page['revision_id']}]}, headers=key())
    assert frozen.status_code == 503
    assert frozen.json['error']['code'] == 'SCENE_HASH_MISMATCH'


def test_add_page_and_restore_replay_after_response_loss(client, app):
    project, first_page = create(client, app)
    project_id = project['project_id']
    add_url = f'/api/v2/projects/{project_id}/pages'
    add_body = {'base_project_version': project['project_version'],
                'outline': {'title': '第二页', 'points': []}}
    add_headers = key()
    added = client.post(add_url, json=add_body, headers=add_headers)
    assert added.status_code == 201, added.json
    repeated = client.post(add_url, json=add_body, headers=add_headers)
    assert repeated.status_code == 201 and repeated.json['data'] == added.json['data']
    assert len(client.get(f'/api/v2/projects/{project_id}').json['data']['pages']) == 2
    saved = patch(client, project, first_page,
                  [{'op': 'add_element', 'element': text_element('需要撤销的文字')}])
    assert saved.status_code == 200
    restore_url = f'/api/v2/projects/{project_id}/pages/{first_page["page_id"]}/restore'
    restore_body = {'base_revision_id': saved.json['data']['revision_id'],
                    'base_page_version': saved.json['data']['page_version'],
                    'target_revision_id': first_page['revision_id']}
    restore_headers = key()
    restored = client.post(restore_url, json=restore_body, headers=restore_headers)
    assert restored.status_code == 200, restored.json
    replayed = client.post(restore_url, json=restore_body, headers=restore_headers)
    assert replayed.status_code == 200 and replayed.json['data'] == restored.json['data']
    with app.app_context():
        assert PageSceneRevision.query.filter_by(page_id=first_page['page_id']).count() == 3


def test_restore_cas_rejects_interleaving_page_update(client, app, monkeypatch):
    project, page = create(client, app)
    saved = patch(client, project, page, [{'op': 'add_element',
        'element': text_element('第二版')}]).json['data']
    from sqlalchemy import update
    from services.scene import versioning
    original_check = versioning.assert_locked_preserved
    def interleave(old, new):
        original_check(old, new)
        # Simulate another writer changing the row after the initial base
        # check, without synchronizing the already loaded Page identity.
        db.session.execute(update(Page).where(Page.id == page['page_id'])
            .values(row_version=Page.row_version + 1)
            .execution_options(synchronize_session=False))
    monkeypatch.setattr(versioning, 'assert_locked_preserved', interleave)
    restored = client.post(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/restore',
        json={'base_revision_id': saved['revision_id'],
              'base_page_version': saved['page_version'],
              'target_revision_id': page['revision_id']}, headers=key())
    assert restored.status_code == 409
    assert restored.json['error']['code'] == 'SCENE_VERSION_CONFLICT'
    with app.app_context():
        assert PageSceneRevision.query.filter_by(page_id=page['page_id']).count() == 2
        assert db.session.get(Page, page['page_id']).head_revision_id == saved['revision_id']


def test_insert_and_copy_page_preserve_scene_assets_but_allocate_new_element_ids(client, app, tmp_path):
    project, source = create(client, app)
    project_id = project['project_id']
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'copy-assets')
    image = BytesIO(); Image.new('RGB', (20, 20), '#446688').save(image, format='PNG'); image.seek(0)
    uploaded = client.post(f'/api/v2/projects/{project_id}/assets',
        data={'file': (image, 'figure.png')}, content_type='multipart/form-data', headers=key())
    assert uploaded.status_code == 201
    asset_id = uploaded.json['data']['asset_id']
    image_element = {'id': str(uuid4()), 'kind': 'image', 'role': 'illustration', 'asset_id': asset_id,
        'frame': {'x': 400, 'y': 160, 'w': 200, 'h': 150, 'rotation_deg': 0},
        'crop': {'x': 0, 'y': 0, 'w': 1, 'h': 1}, 'fit': 'contain', 'opacity': 1,
        'locked': False, 'alt_text': '参考配图'}
    saved = patch(client, project, source, [
        {'op': 'add_element', 'element': text_element('可复制的文字')},
        {'op': 'add_element', 'element': image_element}])
    assert saved.status_code == 200
    source_head = saved.json['data']
    second = client.post(f'/api/v2/projects/{project_id}/pages',
        json={'base_project_version': 0, 'outline': {'title': '后页', 'points': []}}, headers=key())
    assert second.status_code == 201
    copy_url = f'/api/v2/projects/{project_id}/pages'
    payload = {'base_project_version': 1, 'insert_at': 1,
        'copy_from_page_id': source['page_id'], 'copy_from_page_version': source_head['page_version'],
        'copy_from_revision_id': source_head['revision_id']}
    assert client.post(copy_url, json={**payload, 'insert_at': 4}, headers=key()).status_code == 422
    assert client.post(copy_url, json={**payload, 'copy_from_page_version': 1}, headers=key()).status_code == 409
    copy_key = key()
    copied = client.post(copy_url, json=payload, headers=copy_key)
    replay = client.post(copy_url, json=payload, headers=copy_key)
    assert copied.status_code == replay.status_code == 201
    assert copied.json['data'] == replay.json['data']
    assert client.post(copy_url, json=payload, headers=key()).status_code == 409
    copied_page = copied.json['data']
    deck = client.get(f'/api/v2/projects/{project_id}').json['data']
    assert [p['page_id'] for p in deck['pages']] == [source['page_id'], copied_page['page_id'], second.json['data']['page_id']]
    assert deck['pages'][2]['page_version'] == 2
    assert copied_page['outline_content']['title'] == source['outline_content']['title']
    cloned_scene = client.get(f'/api/v2/projects/{project_id}/pages/{copied_page["page_id"]}/scene').json['data']
    assert [e['kind'] for e in cloned_scene['scene']['elements']] == ['text', 'image']
    assert {e['id'] for e in cloned_scene['scene']['elements']}.isdisjoint(
        {e['id'] for e in source_head['scene']['elements']})
    assert cloned_scene['scene']['elements'][1]['asset_id'] == asset_id
    assert client.get(cloned_scene['asset_urls'][asset_id]).status_code == 200
    edited_source = patch(client, project, source_head, [{'op': 'set_text',
        'element_id': source_head['scene']['elements'][0]['id'], 'text': '后来修改的源页'}])
    assert edited_source.status_code == 200
    assert client.get(f'/api/v2/projects/{project_id}/pages/{copied_page["page_id"]}/scene').json[
        'data']['scene']['elements'][0]['text'] == '可复制的文字'
    with app.app_context():
        revision = db.session.get(PageSceneRevision, copied_page['revision_id'])
        assert revision.origin == 'copy'
        assert revision.validation_json['copied_from_revision_id'] == source_head['revision_id']


def test_page_deletion_and_project_archive_replay_after_response_loss(client, app):
    project, _ = create(client, app)
    project_id = project['project_id']
    added = client.post(f'/api/v2/projects/{project_id}/pages',
        json={'base_project_version': 0, 'outline': {'title': '第二页', 'points': []}},
        headers=key())
    assert added.status_code == 201
    page = added.json['data']
    delete_url = f'/api/v2/projects/{project_id}/pages/{page["page_id"]}'
    payload = {'base_project_version': 1, 'base_page_version': page['page_version']}
    delete_key = key()
    deleted = client.delete(delete_url, json=payload, headers=delete_key)
    replay = client.delete(delete_url, json=payload, headers=delete_key)
    assert deleted.status_code == replay.status_code == 200
    assert deleted.json['data'] == replay.json['data']
    assert len(replay.json['data']['pages']) == 1
    assert client.delete(delete_url, json={**payload, 'base_project_version': 2}, headers=delete_key).status_code == 409
    archive_url = f'/api/v2/projects/{project_id}'
    archive_key = key()
    archived = client.delete(archive_url, headers=archive_key)
    replay_archive = client.delete(archive_url, headers=archive_key)
    assert archived.status_code == replay_archive.status_code == 200
    assert archived.json['data'] == replay_archive.json['data']
    assert client.get(archive_url).status_code == 404


def test_outline_and_model_config_replay_without_second_version_increment(client, app):
    project, page = create(client, app)
    project_id = project['project_id']
    config_url = f'/api/v2/projects/{project_id}/model-config'
    config_body = {'base_project_version': 0, 'model_config': {'text_model': 'model-a'}}
    config_headers = key()
    configured = client.patch(config_url, json=config_body, headers=config_headers)
    repeated_config = client.patch(config_url, json=config_body, headers=config_headers)
    assert configured.status_code == repeated_config.status_code == 200
    assert configured.json['data'] == repeated_config.json['data']
    outline_url = f'/api/v2/projects/{project_id}/outline'
    outline_body = {'base_project_version': 1, 'pages': [
        {'page_id': page['page_id'], 'outline': {'title': '更新标题', 'points': ['要点']}},
        {'outline': {'title': '新增页面', 'points': []}}]}
    outline_headers = key()
    saved = client.patch(outline_url, json=outline_body, headers=outline_headers)
    repeated_outline = client.patch(outline_url, json=outline_body, headers=outline_headers)
    assert saved.status_code == repeated_outline.status_code == 200
    assert saved.json['data'] == repeated_outline.json['data']
    current = client.get(f'/api/v2/projects/{project_id}').json['data']
    assert current['project_version'] == 2
    assert len(current['pages']) == 2


def test_image_upload_over_limit_returns_structured_413(client, app, monkeypatch):
    project, _ = create(client, app)
    monkeypatch.setitem(app.config, 'MAX_UPLOAD_BYTES', 8)
    uploaded = client.post(f'/api/v2/projects/{project["project_id"]}/assets',
        data={'file': (BytesIO(b'longer-than-eight-bytes'), 'figure.png')},
        content_type='multipart/form-data', headers=key())
    assert uploaded.status_code == 413
    assert uploaded.json['error']['code'] == 'UPLOAD_TOO_LARGE'


def test_scene_replay_renews_short_lived_asset_urls(client, app, tmp_path, monkeypatch):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'replay-assets')
    image = BytesIO(); Image.new('RGB', (30, 30), '#334455').save(image, format='PNG'); image.seek(0)
    image_bytes = image.getvalue()
    upload_key = key()
    uploaded = client.post(f'/api/v2/projects/{project["project_id"]}/assets',
        data={'file': (image, 'figure.png')}, content_type='multipart/form-data', headers=upload_key)
    assert uploaded.status_code == 201
    asset_id = uploaded.json['data']['asset_id']
    replay_upload = client.post(f'/api/v2/projects/{project["project_id"]}/assets',
        data={'file': (BytesIO(image_bytes), 'figure.png')}, content_type='multipart/form-data', headers=upload_key)
    assert replay_upload.status_code == 201
    assert replay_upload.json['data']['asset_id'] == asset_id
    changed_upload = client.post(f'/api/v2/projects/{project["project_id"]}/assets',
        data={'file': (BytesIO(image_bytes), 'other.png')}, content_type='multipart/form-data', headers=upload_key)
    assert changed_upload.status_code == 409
    with app.app_context():
        assert Asset.query.filter_by(project_id=project['project_id'], kind='upload').count() == 1
    element = {'id': str(uuid4()), 'kind': 'image', 'role': 'illustration', 'asset_id': asset_id,
        'frame': {'x': 60, 'y': 80, 'w': 300, 'h': 200, 'rotation_deg': 0},
        'crop': {'x': 0, 'y': 0, 'w': 1, 'h': 1}, 'fit': 'contain', 'opacity': 1,
        'locked': False, 'alt_text': ''}
    url = f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/scene'
    payload = {'base_revision_id': page['revision_id'], 'base_page_version': page['page_version'],
               'commands': [{'op': 'add_element', 'element': element}]}
    headers = key()
    first = client.patch(url, json=payload, headers=headers)
    assert first.status_code == 200, first.json
    first_url = first.json['data']['asset_urls'][asset_id]
    import services.scene.store as asset_store
    real_time = asset_store.time.time()
    monkeypatch.setattr(asset_store.time, 'time', lambda: real_time + 400)
    replayed = client.patch(url, json=payload, headers=headers)
    assert replayed.status_code == 200
    assert replayed.json['data']['revision_id'] == first.json['data']['revision_id']
    renewed_url = replayed.json['data']['asset_urls'][asset_id]
    assert renewed_url != first_url
    assert client.get(first_url).status_code == 404
    assert client.get(renewed_url).status_code == 200


def test_asset_audit_verifies_migrated_database_and_bytes(tmp_path):
    from alembic import command
    from alembic.config import Config as AlembicConfig
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    root = Path(__file__).resolve().parents[3]
    database = tmp_path / 'backup-audit.db'
    url = f'sqlite:///{database.as_posix()}'
    config = AlembicConfig(str(root / 'backend' / 'alembic.ini'))
    config.set_main_option('sqlalchemy.url', url)
    command.upgrade(config, 'head')
    assets_root = tmp_path / 'assets'
    assets_root.mkdir()
    owner_id, project_id, asset_id = str(uuid4()), str(uuid4()), str(uuid4())
    page_id, revision_id, snapshot_id, credential_id = (str(uuid4()) for _ in range(4))
    key = f'owners/{owner_id}/projects/{project_id}/assets/{asset_id}/content.bin'
    path = assets_root / key
    path.parent.mkdir(parents=True)
    path.write_bytes(b'immutable scene asset')
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            session.add(Principal(id=owner_id, kind='local_owner', display_name='Owner'))
            session.add(Project(id=project_id, owner_id=owner_id, editor_mode='scene_v1',
                project_title='Backup test', canvas_width_pt=960, canvas_height_pt=540,
                font_manifest_id='fonts-v1'))
            session.flush()
            session.add(Asset(id=asset_id, owner_id=owner_id, project_id=project_id,
                kind='upload', storage_key=key, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                byte_size=path.stat().st_size, mime_type='application/octet-stream', state='ready'))
            scene = blank_scene()
            session.add(Page(id=page_id, project_id=project_id, order_index=0,
                head_revision_id=revision_id))
            session.add(PageSceneRevision(id=revision_id, owner_id=owner_id,
                project_id=project_id, page_id=page_id, seq=1, scene_json=scene,
                scene_hash=digest(scene), origin='manual', created_by=owner_id))
            manifest = {'pages': [{'page_id': page_id, 'revision_id': revision_id,
                                   'scene_hash': digest(scene)}], 'assets': {}}
            session.add(DeckSnapshot(id=snapshot_id, owner_id=owner_id,
                project_id=project_id, confirmed_by=owner_id,
                manifest_json=manifest, manifest_hash=digest(manifest)))
            session.add(SnapshotPage(snapshot_id=snapshot_id, ordinal=1,
                page_id=page_id, revision_id=revision_id))
            from cryptography.hazmat.primitives.ciphers.aead import AESGCM
            secret, nonce, plaintext = os.urandom(32), os.urandom(12), b'sandbox-key'
            associated = f'{owner_id}:{credential_id}'.encode()
            session.add(ModelCredential(id=credential_id, owner_id=owner_id,
                provider_kind='sub2api_openai_compatible', base_url='https://fixture.invalid/v1',
                label='Audit test', key_suffix='-key', key_version='v1', nonce=nonce,
                encrypted_secret=AESGCM(secret).encrypt(nonce, plaintext, associated),
                secret_fingerprint=hmac.new(secret, associated + plaintext, hashlib.sha256).hexdigest(),
                capabilities_json={}))
            session.commit()
    finally:
        engine.dispose()
    encoded_key = base64.urlsafe_b64encode(secret).decode()
    report = audit_scene(url, assets_root, check_credentials=True, encryption_key=encoded_key)
    assert report['schema_revision'] == '13e79b401ca8'
    assert report['asset_count'] == 1 and report['assets'][0]['asset_id'] == asset_id
    assert report['snapshot_count'] == 1 and report['credential_count'] == 1
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            session.add(RevisionAsset(revision_id=revision_id, asset_id=asset_id,
                                      role='illustration', project_id=project_id))
            session.commit()
    finally:
        engine.dispose()
    with pytest.raises(ValueError, match='Scene asset references mismatch'):
        audit_scene(url, assets_root)
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            session.delete(session.get(RevisionAsset, (revision_id, asset_id, 'illustration')))
            session.commit()
    finally:
        engine.dispose()

    export_id, file_id, report_id, proof_id = (str(uuid4()) for _ in range(4))
    manifest_hash = digest(manifest)
    report_bytes = json.dumps({'snapshot_id': snapshot_id, 'snapshot_hash': manifest_hash,
        'format': 'pptx', 'structural_and_text': 'passed',
        'visual_comparison': 'needs_review',
        'waivable_warning_codes': ['VISUAL_DIFF_UNCALIBRATED'],
        'visual': {'pages': [{'page_id': page_id, 'contact_sheet_asset_id': proof_id}]}}).encode()
    files = ((file_id, b'export bytes', 'export', 'pptx'),
             (report_id, report_bytes, 'export_report', 'json'),
             (proof_id, b'proof bytes', 'visual_comparison', 'png'))
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            for identifier, content, kind, extension in files:
                storage_key = f'owners/{owner_id}/projects/{project_id}/assets/{identifier}/content.{extension}'
                asset_path = assets_root / storage_key
                asset_path.parent.mkdir(parents=True)
                asset_path.write_bytes(content)
                session.add(Asset(id=identifier, owner_id=owner_id, project_id=project_id,
                    kind=kind, storage_key=storage_key,
                    sha256=hashlib.sha256(content).hexdigest(), byte_size=len(content),
                    mime_type='application/octet-stream', state='ready',
                    provenance_json={'snapshot_id': snapshot_id, 'snapshot_hash': manifest_hash,
                                     'report_asset_id': report_id}
                        if kind == 'export' else {}))
            session.add(SceneExport(id=export_id, owner_id=owner_id, project_id=project_id,
                snapshot_id=snapshot_id, format='pptx', options_json={}, options_hash=digest({}),
                status='succeeded', file_asset_id=file_id, report_asset_id=report_id,
                review_report_sha256=hashlib.sha256(report_bytes).hexdigest(),
                reviewed_by=owner_id, reviewed_at=now()))
            session.commit()
    finally:
        engine.dispose()
    assert audit_scene(url, assets_root)['export_count'] == 1
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            session.get(SceneExport, export_id).review_report_sha256 = '0' * 64
            session.commit()
    finally:
        engine.dispose()
    with pytest.raises(ValueError, match='Export review hash mismatch'):
        audit_scene(url, assets_root)
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            session.get(SceneExport, export_id).review_report_sha256 = hashlib.sha256(report_bytes).hexdigest()
            session.commit()
    finally:
        engine.dispose()
    report_path = assets_root / f'owners/{owner_id}/projects/{project_id}/assets/{report_id}/content.json'
    tampered_report = json.loads(report_bytes)
    tampered_report['visual']['pages'][0]['contact_sheet_asset_id'] = str(uuid4())
    tampered_bytes = json.dumps(tampered_report).encode()
    report_path.write_bytes(tampered_bytes)
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            report_asset = session.get(Asset, report_id)
            report_asset.sha256 = hashlib.sha256(tampered_bytes).hexdigest()
            report_asset.byte_size = len(tampered_bytes)
            session.get(SceneExport, export_id).review_report_sha256 = report_asset.sha256
            session.commit()
    finally:
        engine.dispose()
    with pytest.raises(ValueError, match='Export visual proof mismatch'):
        audit_scene(url, assets_root)
    report_path.write_bytes(report_bytes)
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            report_asset = session.get(Asset, report_id)
            report_asset.sha256 = hashlib.sha256(report_bytes).hexdigest()
            report_asset.byte_size = len(report_bytes)
            session.get(SceneExport, export_id).review_report_sha256 = report_asset.sha256
            session.commit()
    finally:
        engine.dispose()
    with pytest.raises(ValueError, match='Credential cannot be decrypted'):
        audit_scene(url, assets_root, check_credentials=True,
                    encryption_key=base64.urlsafe_b64encode(os.urandom(32)).decode())
    path.write_bytes(b'tampered')
    with pytest.raises(ValueError, match='Asset size/hash mismatch'):
        audit_scene(url, assets_root)
    path.write_bytes(b'immutable scene asset')
    engine = create_engine(url)
    try:
        with Session(engine) as session:
            snapshot_page = session.get(SnapshotPage, (snapshot_id, 1))
            snapshot_page.revision_id = str(uuid4())
            session.commit()
    finally:
        engine.dispose()
    with pytest.raises(ValueError, match='Snapshot revision mismatch'):
        audit_scene(url, assets_root)


def test_renderer_identity_must_be_complete_before_job_creation(app, tmp_path, monkeypatch):
    with app.app_context():
        monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', str(tmp_path))
        monkeypatch.setitem(app.config, 'SCENE_RENDER_UID', 10001)
        monkeypatch.setitem(app.config, 'SCENE_RENDER_GID', -1)
        with pytest.raises(SceneError) as raised:
            run_render_job('pdf', b'<html></html>')
        assert raised.value.code == 'RENDERER_CONFIG_INVALID'
        assert not [path for path in tmp_path.iterdir() if path.is_dir()]


def test_unique_renderer_identity_fails_closed_without_linux_root(app, tmp_path, monkeypatch):
    if os.name == 'posix' and os.geteuid() == 0:
        pytest.skip('This negative-path test needs a non-root host')
    with app.app_context():
        monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', str(tmp_path))
        monkeypatch.setitem(app.config, 'SCENE_RENDER_UID', -1)
        monkeypatch.setitem(app.config, 'SCENE_RENDER_GID', -1)
        monkeypatch.setitem(app.config, 'SCENE_RENDER_UNIQUE_UID', True)
        with pytest.raises(SceneError) as raised:
            run_render_job('pdf', b'<html></html>')
        assert raised.value.code == 'RENDERER_CONFIG_INVALID'
        assert not [path for path in tmp_path.iterdir() if path.is_dir()]


def test_renderer_output_must_be_bounded_regular_file(tmp_path):
    job = tmp_path / 'job'
    job.mkdir()
    (job / 'output.pdf').write_bytes(b'%PDF-safe')
    assert _read_render_output(job, 'output.pdf', 100) == b'%PDF-safe'
    with pytest.raises(SceneError) as raised:
        _read_render_output(job, 'output.pdf', 2)
    assert raised.value.code == 'RENDER_OUTPUT_INVALID'
    (job / 'directory.pdf').mkdir()
    with pytest.raises(SceneError) as raised:
        _read_render_output(job, 'directory.pdf', 100)
    assert raised.value.code == 'RENDER_OUTPUT_INVALID'
    if hasattr(os, 'O_NOFOLLOW'):
        (job / 'linked.pdf').symlink_to(job / 'output.pdf')
        with pytest.raises(SceneError) as raised:
            _read_render_output(job, 'linked.pdf', 100)
        assert raised.value.code == 'RENDER_OUTPUT_INVALID'


def test_renderer_does_not_expose_a_new_job_beside_an_unfinished_orphan(app, tmp_path, monkeypatch):
    monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', str(tmp_path))
    (tmp_path / 'orphan').mkdir()
    with app.app_context(), pytest.raises(SceneError) as raised:
        run_render_job('pdf', b'<html></html>')
    assert raised.value.code == 'RENDERER_ORPHANED_JOB'
    assert [path.name for path in tmp_path.iterdir() if path.is_dir()] == ['orphan']


def test_scene_readiness_checks_worker_renderer_storage_and_font(client, app, tmp_path, monkeypatch):
    from backend.tests.scene_layout_fixtures import identity
    monkeypatch.setitem(app.config, 'SCENE_EDITOR_ENABLED', True)
    assets = tmp_path / 'assets'
    jobs = tmp_path / 'jobs'
    assets.mkdir(); jobs.mkdir()
    monkeypatch.setitem(app.config, 'ASSET_STORE_ROOT', str(assets))
    monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', str(jobs))
    monkeypatch.setitem(app.config, 'SCENE_MIN_FREE_BYTES', 1)
    assert client.get('/api/v2/readiness').status_code == 503
    with app.app_context():
        db.session.add(SceneWorkerHeartbeat(worker_id='test-worker', heartbeat_at=now()))
        db.session.commit()
    (jobs / '.renderer_status.json').write_text(json.dumps({
        'updated_at': time.time(), 'chromium_ok': True, 'libreoffice_ok': True}))
    assert client.get('/api/v2/readiness').status_code == 503  # Old TextLayout protocol is not ready.
    (jobs / '.renderer_status.json').write_text(json.dumps({
        'updated_at': time.time(), 'chromium_ok': True, 'libreoffice_ok': True, 'text_layout_identity': identity()}))
    assert client.get('/api/v2/readiness').status_code == 503  # No actual PPTX-to-PDF capability proof.
    (jobs / '.renderer_status.json').write_text(json.dumps({
        'updated_at': time.time(), 'chromium_ok': True, 'libreoffice_ok': True,
        'pptx_pdf_ok': True, 'text_layout_identity': identity()}))
    ready = client.get('/api/v2/readiness')
    assert ready.status_code == 200 and ready.json['data']['ready'] is True
    assert all(ready.json['data']['checks'].values())
    monkeypatch.setitem(app.config, 'SCENE_RENDER_UNIQUE_UID', True)
    assert client.get('/api/v2/readiness').status_code == 503
    (jobs / '.renderer_status.json').write_text(json.dumps({
        'updated_at': time.time(), 'chromium_ok': True, 'libreoffice_ok': True,
        'isolation_mode': 'unique_uid', 'pptx_pdf_ok': True, 'text_layout_identity': identity()}))
    assert client.get('/api/v2/readiness').status_code == 200
    (jobs / '.renderer_status.json').write_text(json.dumps({
        'updated_at': time.time() - 120, 'chromium_ok': True, 'libreoffice_ok': True}))
    stale = client.get('/api/v2/readiness')
    assert stale.status_code == 503 and stale.json['data']['checks']['renderer'] is False


def test_worker_expiry_unknown_and_heartbeat_fence(client, app):
    project, page = create(client, app)
    with app.app_context():
        paid = SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], page_id=page['page_id'],
            operation='generate_scene', resource_id=str(uuid4()), input_json={},
            input_hash='0' * 64, state='running', attempt_count=1,
            dispatch_state='may_have_been_sent', lease_owner='crashed-worker',
            lease_expires_at=now() - timedelta(seconds=5), fence_token=7)
        db.session.add(paid)
        db.session.add(SceneTaskAttempt(task_item_id=paid.id, attempt_no=1,
            fence_token=7, worker_id='crashed-worker',
            dispatch_state='may_have_been_sent'))
        db.session.commit()
        assert claim() is None
        db.session.refresh(paid)
        assert paid.state == 'outcome_unknown'
        assert not renew_lease(paid.id, 7, 'crashed-worker', 60)
        assert SceneTaskAttempt.query.filter_by(task_item_id=paid.id).one().finished_at is not None

        local = SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], operation='export',
            resource_id=str(uuid4()), input_json={}, input_hash='1' * 64,
            max_attempts=2)
        db.session.add(local)
        db.session.commit()
        assert claim() == (local.id, 1)
        assert renew_lease(local.id, 0, WORKER_ID, 60) is False
        assert renew_lease(local.id, 1, WORKER_ID, 60) is True
        local.cancel_requested_at = now()
        db.session.commit()
        assert renew_lease(local.id, 1, WORKER_ID, 60) is False


def test_worker_claim_scans_past_busy_owner_queue(client, app, monkeypatch):
    project, _ = create(client, app)
    other_owner = str(uuid4())
    other_project = str(uuid4())
    monkeypatch.setitem(app.config, 'OWNER_MODEL_CONCURRENCY', 1)
    with app.app_context():
        db.session.add(Principal(id=other_owner, kind='sub2api_user', display_name='Other'))
        db.session.add(Project(id=other_project, owner_id=other_owner,
                               editor_mode='scene_v1', project_title='Other'))
        base_time = now() - timedelta(minutes=5)
        db.session.add(SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], operation='generate_outline',
            resource_id=str(uuid4()), input_json={}, input_hash='0' * 64,
            state='running', lease_owner='busy-worker',
            lease_expires_at=now() + timedelta(minutes=5), created_at=base_time))
        for index in range(51):
            db.session.add(SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
                project_id=project['project_id'], operation='generate_outline',
                resource_id=str(uuid4()), input_json={}, input_hash='1' * 64,
                created_at=base_time + timedelta(seconds=index + 1)))
        available = SceneTaskItem(id=str(uuid4()), owner_id=other_owner,
            project_id=other_project, operation='generate_outline',
            resource_id=str(uuid4()), input_json={}, input_hash='2' * 64,
            created_at=base_time + timedelta(seconds=52))
        db.session.add(available)
        db.session.commit()
        assert claim() == (available.id, 1)
        assert SceneTaskItem.query.filter_by(owner_id=LOCAL_OWNER_ID,
            state='queued').count() == 51


def test_local_worker_failure_retries_only_within_attempt_budget(client, app, monkeypatch):
    project, _ = create(client, app)
    with app.app_context():
        item = SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], operation='export',
            resource_id=str(uuid4()), input_json={}, input_hash='2' * 64,
            max_attempts=2)
        db.session.add(item)
        db.session.commit()
        monkeypatch.setattr('workers.scene_runner.perform',
                            lambda *_: (_ for _ in ()).throw(OSError('temporary IO')))
        assert run_once()
        db.session.refresh(item)
        assert item.state == 'queued'
        assert item.attempt_count == 1
        assert item.next_run_at is not None
        item.next_run_at = now() - timedelta(seconds=1)
        db.session.commit()
        assert run_once()
        db.session.refresh(item)
        assert item.state == 'failed'
        assert item.attempt_count == 2


def test_worker_does_not_claim_asset_child_after_parent_is_cancelled(client, app):
    project, page = create(client, app)
    with app.app_context():
        group_id = str(uuid4())
        parent = SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], page_id=page['page_id'], group_id=group_id,
            logical_key='generate:parent', operation='generate_scene',
            resource_id=page['page_id'], input_json={}, input_hash='1' * 64,
            state='cancelled')
        child = SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], page_id=page['page_id'], group_id=group_id,
            logical_key='asset:0', operation='generate_asset', resource_id=page['page_id'],
            input_json={'parent_task_id': parent.id}, input_hash='2' * 64)
        db.session.add_all([parent, child])
        db.session.commit()
        assert claim() is None
        db.session.refresh(child)
        assert child.state == 'cancelled'
        assert child.error_code == 'PARENT_UNAVAILABLE'
        assert child.attempt_count == 0


def test_paid_model_retry_requires_charge_ack(client, app):
    project, page = create(client, app)
    with app.app_context():
        item = SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], page_id=page['page_id'],
            operation='generate_scene', resource_id=page['page_id'],
            input_json={}, input_hash='3' * 64, state='failed',
            attempt_count=1, max_attempts=2, dispatch_state='may_have_been_sent')
        db.session.add(item)
        db.session.commit()
        task_id = item.id
    status = client.get(f'/api/v2/tasks/{task_id}')
    assert status.json['data']['possible_charge'] is True
    listed = client.get(f'/api/v2/projects/{project["project_id"]}/tasks')
    assert listed.status_code == 200
    assert listed.json['data']['items'][0]['task_id'] == task_id
    denied = client.post(f'/api/v2/tasks/{task_id}/retry', json={}, headers=key())
    assert denied.status_code == 422
    assert denied.json['error']['code'] == 'POSSIBLE_CHARGE_ACK_REQUIRED'
    retry_headers = key()
    accepted = client.post(f'/api/v2/tasks/{task_id}/retry',
                           json={'acknowledge_possible_charge': True}, headers=retry_headers)
    assert accepted.status_code == 200
    assert accepted.json['data']['state'] == 'queued'
    with app.app_context():
        item = db.session.get(SceneTaskItem, task_id)
        item.state = 'failed'  # A rapid second failure before the browser receives the first response.
        db.session.commit()
    replayed = client.post(f'/api/v2/tasks/{task_id}/retry',
                           json={'acknowledge_possible_charge': True}, headers=retry_headers)
    assert replayed.status_code == 200 and replayed.json['data'] == accepted.json['data']
    with app.app_context():
        assert db.session.get(SceneTaskItem, task_id).state == 'failed'
    assert client.post(f'/api/v2/tasks/{task_id}/retry', json={}, headers=retry_headers).status_code == 409


def key():
    return {'Idempotency-Key': str(uuid4())}


def create(client, app):
    app.config['SCENE_EDITOR_ENABLED'] = True
    response = client.post('/api/v2/projects', json={'title': '中文测试', 'pages': [{'title': '第一页', 'points': []}]}, headers=key())
    assert response.status_code == 201, response.json
    project = response.json['data']
    return project, project['pages'][0]


@pytest.mark.parametrize(('override', 'code'), [
    ({'canvas': {'width_pt': 'nan', 'height_pt': 540}}, 'CANVAS_INVALID'),
    ({'canvas': None}, 'CANVAS_INVALID'),
    ({'pages': [{'title': 'Bad', 'points': [123]}]}, 'OUTLINE_INVALID'),
    ({'model_config': {'text_model': {'nested': 'invalid'}}}, 'MODEL_CONFIG_INVALID'),
    ({'prompt': {'not': 'text'}}, 'PROJECT_INVALID'),
])
def test_project_creation_rejects_malformed_inputs(client, app, override, code):
    app.config['SCENE_EDITOR_ENABLED'] = True
    payload = {'title': 'Input test', 'pages': [{'title': '第一页', 'points': []}], **override}
    response = client.post('/api/v2/projects', json=payload, headers=key())
    assert response.status_code == 422
    assert response.json['error']['code'] == code
    assert response.json['error']['stage'] == 'create_project'


def test_standard_four_three_canvas_exports_at_selected_size(client, app):
    app.config['SCENE_EDITOR_ENABLED'] = True
    created = client.post('/api/v2/projects', json={'title': '4:3 测试',
        'canvas': {'width_pt': 720, 'height_pt': 540},
        'pages': [{'title': '单页', 'points': []}]}, headers=key())
    assert created.status_code == 201
    project = created.json['data']
    page = project['pages'][0]
    snap = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': 0, 'pages': [{'page_id': page['page_id'],
            'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    assert snap.status_code == 201
    with app.app_context():
        snapshot = db.session.get(DeckSnapshot, snap.json['data']['snapshot_id'])
        revision = db.session.get(PageSceneRevision, page['revision_id'])
        from pptx import Presentation
        exported = Presentation(BytesIO(render_pptx(snapshot, {revision.id: revision}, {})))
        assert exported.slide_width == 720 * 12700
        assert exported.slide_height == 540 * 12700


def text_element(text='季度汇报'):
    return {'id': str(uuid4()), 'kind': 'text', 'role': 'title',
            'frame': {'x': 40, 'y': 40, 'w': 700, 'h': 90, 'rotation_deg': 0},
            'text': text, 'style': {'font_family_id': 'noto-sans-sc',
            'font_size_pt': 32, 'font_weight': 700, 'color': '#17324D',
            'align': 'left', 'vertical_align': 'top', 'line_height': 1.2,
            'padding_pt': 0}, 'locked': False}


def test_ai_candidate_cannot_cover_locked_title_with_new_text():
    title = text_element('锁定标题')
    title['locked'] = True
    added = text_element('遮挡文字')
    added['frame'] = dict(title['frame'])
    old = {'elements': [title]}
    with pytest.raises(CommandError, match='obscures locked text'):
        assert_locked_preserved(old, {'elements': [title, added]})
    added['frame'] = {**added['frame'], 'y': 250}
    assert_locked_preserved(old, {'elements': [title, added]})


def test_ai_selected_scope_cannot_reorder_unselected_elements():
    selected = text_element('选中')
    first = text_element('未选 A')
    second = text_element('未选 B')
    for index, element in enumerate((selected, first, second)):
        element['frame']['y'] = 40 + index * 140
    scene = blank_scene()
    scene['elements'] = [selected, first, second]
    source = {'scope': {'kind': 'selected_elements', 'element_ids': [selected['id']]}}
    project = SimpleNamespace(canvas_width_pt=960, canvas_height_pt=540)
    with pytest.raises(SceneError, match='reordered unselected'):
        _edit_scene(scene, [{'op': 'reorder_elements', 'element_ids': [
            selected['id'], second['id'], first['id']]}], source, project)
    changed = _edit_scene(scene, [
        {'op': 'set_text', 'element_id': selected['id'], 'text': '修改选中对象'},
        {'op': 'reorder_elements', 'element_ids': [first['id'], selected['id'], second['id']]},
    ], source, project)
    assert [item['id'] for item in changed['elements']] == [first['id'], selected['id'], second['id']]


@pytest.mark.parametrize('command', [
    {'op': 'add_element', 'element': {'id': 7}},
    {'op': 'reorder_elements', 'element_ids': [{'bad': 'id'}]},
])
def test_malformed_editor_commands_return_422(client, app, command):
    project, page = create(client, app)
    response = patch(client, project, page, [command])
    assert response.status_code == 422
    assert response.json['error']['code'] == 'INPUT_INVALID'
    current = client.get(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/scene')
    assert current.json['data']['revision_id'] == page['revision_id']


def test_pdf_quality_gate_checks_text_region_not_only_full_page():
    element = text_element('ABC 123')
    element['frame'] = {'x': 80, 'y': 150, 'w': 500, 'h': 130, 'rotation_deg': 0}
    snapshot = SimpleNamespace(manifest_json={
        'canvas': {'width_pt': 960, 'height_pt': 540},
        'pages': [{'page_id': str(uuid4()), 'revision_id': 'revision'}]})
    revisions = {'revision': SimpleNamespace(scene_json={'elements': [element]})}
    document = fitz.open()
    page = document.new_page(width=960, height=540)
    page.insert_text((90, 180), 'ABC 123', fontsize=30)
    assert verify_pdf(document.tobytes(), snapshot, revisions)['searchable_text']
    misplaced = fitz.open()
    page = misplaced.new_page(width=960, height=540)
    page.insert_text((600, 180), 'ABC 123', fontsize=30)
    with pytest.raises(ValueError, match='outside frame'):
        verify_pdf(misplaced.tobytes(), snapshot, revisions)


def test_text_preflight_blocks_horizontal_glyph_clipping_and_empty_inner_box():
    element = text_element('W')
    element['frame']['w'] = 5
    element['frame']['h'] = 120
    with pytest.raises(ValueError, match='glyph wider than frame'):
        validate_text_fit({'elements': [element]})
    element['frame']['w'] = 100
    element['style']['padding_pt'] = 51
    with pytest.raises(ValueError, match='no inner width'):
        validate_text_fit({'elements': [element]})
    element['frame']['w'] = 500
    element['frame']['h'] = 30
    element['style']['padding_pt'] = 16
    with pytest.raises(ValueError, match='no inner height'):
        validate_text_fit({'elements': [element]})
    element['style']['padding_pt'] = 0
    element['text'] = '长段落' * 1000
    with pytest.raises(ValueError, match='Text overflow'):
        validate_text_fit({'elements': [element]})
    element['text'] = ''
    validate_text_fit({'elements': [element]})


def test_chromium_type3_cjk_pdf_has_copyable_unicode():
    element = text_element('中文 English 123')
    snapshot = SimpleNamespace(manifest_json={
        'canvas': {'width_pt': 960, 'height_pt': 540},
        'pages': [{'page_id': str(uuid4()), 'revision_id': 'revision'}]})
    revisions = {'revision': SimpleNamespace(scene_json={'elements': [element]})}
    fixture = Path(__file__).resolve().parents[3] / 'docs' / 'fixtures' / 'scene_pdf_chromium_cjk.pdf'
    report = verify_pdf(fixture.read_bytes(), snapshot, revisions)
    assert report['searchable_text'] and report['pages'][0]['text_count'] == 1


def test_fixed_font_blocks_unsupported_glyph_before_snapshot(client, app):
    supported = text_element('中文 English 123 ★ ™')
    validate_text_fit({'elements': [supported]})
    unsupported = text_element('中文 😀')
    with pytest.raises(SceneError) as raised:
        validate_text_fit({'elements': [unsupported]})
    assert raised.value.code == 'FONT_GLYPH_UNAVAILABLE'
    assert 'U+1F600' in raised.value.message
    project, page = create(client, app)
    saved = patch(client, project, page, [{'op': 'add_element', 'element': unsupported}])
    assert saved.status_code == 200, saved.json
    head = saved.json['data']
    frozen = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': project['project_version'], 'pages': [{
            'page_id': page['page_id'], 'page_version': head['page_version'],
            'revision_id': head['revision_id']}]}, headers=key())
    assert frozen.status_code == 422
    assert frozen.json['error']['code'] == 'FONT_GLYPH_UNAVAILABLE'
    with app.app_context():
        assert DeckSnapshot.query.filter_by(project_id=project['project_id']).count() == 0


def test_snapshot_preflight_reports_blocker_without_creating_snapshot(client, app):
    project, first = create(client, app)
    project_id = project['project_id']
    url = f'/api/v2/projects/{project_id}/snapshots/preflight'
    initial = client.get(url, query_string={'page_ids': first['page_id']})
    assert initial.status_code == 200 and initial.json['data']['ready']
    assert initial.headers['Cache-Control'] == 'no-store'
    overflow = text_element('中文段落' * 200)
    overflow['frame']['h'] = 45
    changed = patch(client, project, first, [{'op': 'add_element', 'element': overflow}])
    assert changed.status_code == 200, changed.json
    head = changed.json['data']
    blocked = client.get(url, query_string={'page_ids': first['page_id']})
    assert blocked.status_code == 200 and not blocked.json['data']['ready']
    assert blocked.json['data']['blockers'][0]['code'] == 'TEXT_OVERFLOW'
    assert blocked.json['data']['blockers'][0]['page_id'] == first['page_id']
    frozen = client.post(f'/api/v2/projects/{project_id}/snapshots', json={
        'project_version': project['project_version'], 'pages': [{
            'page_id': first['page_id'], 'page_version': head['page_version'],
            'revision_id': head['revision_id']}]}, headers=key())
    assert frozen.status_code == 422 and frozen.json['error']['code'] == 'TEXT_OVERFLOW'
    second = client.post(f'/api/v2/projects/{project_id}/pages', json={
        'base_project_version': project['project_version'],
        'outline': {'title': '可导出页', 'points': []}}, headers=key())
    assert second.status_code == 201
    selected = client.get(url, query_string={'page_ids': second.json['data']['page_id']})
    assert selected.status_code == 200 and selected.json['data']['ready']
    second_page = second.json['data']
    missing_glyph = patch(client, project, second_page,
        [{'op': 'add_element', 'element': text_element('😀')}])
    assert missing_glyph.status_code == 200
    combined = client.get(url, query_string={'page_ids': f"{first['page_id']},{second_page['page_id']}"})
    assert combined.status_code == 200 and not combined.json['data']['ready']
    assert [(item['page_id'], item['code']) for item in combined.json['data']['blockers']] == [
        (first['page_id'], 'TEXT_OVERFLOW'), (second_page['page_id'], 'FONT_GLYPH_UNAVAILABLE')]
    with app.app_context():
        assert DeckSnapshot.query.filter_by(project_id=project_id).count() == 0


@pytest.mark.parametrize('invalid', ['A\tB', 'A\ud800B'])
def test_scene_rejects_invalid_unicode_or_control_character(invalid):
    scene = blank_scene()
    scene['elements'] = [text_element(invalid)]
    with pytest.raises(ValueError, match='control character'):
        normalize_scene(scene, 960, 540)


def test_chromium_text_layout_adds_editable_pptx_soft_breaks(app):
    if not os.environ.get('SCENE_CHROMIUM_PATH'):
        pytest.skip('Set SCENE_CHROMIUM_PATH to measure real-browser text wrapping')
    element = text_element('中文 English 123\n第二行继续文本')
    element['frame'] = {**element['frame'], 'w': 130, 'h': 200}
    element['style'] = {**element['style'], 'font_size_pt': 20, 'font_weight': 400}
    scene = {'canvas': {'width_pt': 960, 'height_pt': 540},
             'background': {'kind': 'solid', 'color': '#FFFFFF'}, 'elements': [element]}
    snapshot = SimpleNamespace(manifest_json={'canvas': scene['canvas'],
        'pages': [{'page_id': str(uuid4()), 'revision_id': 'revision'}]})
    revisions = {'revision': SimpleNamespace(scene_json=scene)}
    with app.app_context():
        layout = measure_text_layout(snapshot, revisions, {})
        breaks = [line['end'] for line in layout['pages'][0][element['id']]['lines'] if line['break_kind'] == 'soft']
        assert breaks and all(element['text'][offset - 1] != '\n' for offset in breaks)
        payload = render_pptx(snapshot, revisions, {}, layout)
        assert verify_pptx(payload, snapshot, revisions, layout)['native_text_count'] == 1
    from pptx import Presentation
    exported = Presentation(BytesIO(payload)).slides[0].shapes[0].text
    assert '\v' in exported
    assert exported.replace('\v', '') == element['text']


def test_text_layout_is_transported_through_renderer_job(app, tmp_path, monkeypatch):
    if not os.environ.get('SCENE_CHROMIUM_PATH'):
        pytest.skip('Set SCENE_CHROMIUM_PATH to exercise renderer job transport')
    from renderer.daemon import run_job
    from services.exports.scene_pdf import FONT
    root = tmp_path / 'renderer-layout-jobs'
    root.mkdir()
    monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', str(root))
    monkeypatch.setitem(app.config, 'SCENE_RENDER_UNIQUE_UID', False)
    monkeypatch.setitem(app.config, 'SCENE_RENDER_UID', -1)
    monkeypatch.setitem(app.config, 'SCENE_RENDER_GID', -1)
    failures = []

    def serve_once():
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline:
            for directory in root.iterdir():
                if directory.is_dir() and (directory / 'ready').is_file():
                    try:
                        result = run_job(directory)
                    except Exception as exc:
                        failures.append(exc)
                        result = {'status': 'failed'}
                    (directory / 'result.json').write_text(json.dumps(result), encoding='utf-8')
                    return
            time.sleep(.05)
        failures.append(TimeoutError('renderer job did not appear'))

    html = ('<!doctype html><meta charset="utf-8"><style>'
            f'@font-face{{font-family:NotoScene;src:url("{FONT.as_uri()}")}}'
            'body{font-family:NotoScene;line-height:1.2}[data-scene-id]{height:60pt}'
            '[data-scene-paragraph]{min-height:1lh}</style><section class="slide">'
            '<div data-scene-id="smoke"><div data-scene-content><div data-scene-paragraph>中文 Scene 123</div></div></div></section>')
    server = threading.Thread(target=serve_once, daemon=True)
    server.start()
    with app.app_context():
        payload, _ = run_render_job('text_layout', html.encode('utf-8'))
    server.join(timeout=2)
    assert not failures
    assert json.loads(payload)['pages'][0]['smoke']['lines'][0]['end'] == len('中文 Scene 123')


def patch(client, project, page, commands):
    return client.patch(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/scene',
        json={'base_revision_id': page['revision_id'], 'base_page_version': page['page_version'],
              'commands': commands}, headers=key())


def test_save_conflict_lock_restore_and_native_pptx(client, app):
    project, page = create(client, app)
    element = text_element('中文 English 123')
    saved = patch(client, project, page, [{'op': 'add_element', 'element': element}])
    assert saved.status_code == 200, saved.json
    stale = patch(client, project, page, [{'op': 'add_element', 'element': text_element()}])
    assert stale.status_code == 409
    current = saved.json['data']
    locked = patch(client, project, current, [{'op': 'set_locked', 'element_id': element['id'], 'locked': True}])
    assert locked.status_code == 200
    blocked = patch(client, project, locked.json['data'], [{'op': 'set_text', 'element_id': element['id'], 'text': '越权'}])
    assert blocked.status_code == 422
    restored = client.post(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/restore',
        json={'base_revision_id': locked.json['data']['revision_id'],
              'base_page_version': locked.json['data']['page_version'],
              'target_revision_id': page['revision_id']}, headers=key())
    assert restored.status_code == 422  # restore cannot silently unlock
    snapshot = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots',
        json={'project_version': 0, 'pages': [{'page_id': page['page_id'],
              'page_version': locked.json['data']['page_version'],
              'revision_id': locked.json['data']['revision_id']}]}, headers=key())
    assert snapshot.status_code == 201, snapshot.json
    with app.app_context():
        snap = db.session.get(DeckSnapshot, snapshot.json['data']['snapshot_id'])
        rev = db.session.get(PageSceneRevision, locked.json['data']['revision_id'])
        payload = render_pptx(snap, {rev.id: rev}, {})
        report = verify_pptx(payload, snap, {rev.id: rev})
        assert report['native_text_count'] == 1


def test_lock_gesture_undo_redo_requires_explicit_commands(client, app):
    project, page = create(client, app)
    element = text_element('锁定操作的撤销不丢文字')
    added = patch(client, project, page, [{'op': 'add_element', 'element': element}]).json['data']
    locked = patch(client, project, added, [
        {'op': 'set_locked', 'element_id': element['id'], 'locked': True}]).json['data']
    restore = client.post(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/restore',
        json={'base_revision_id': locked['revision_id'], 'base_page_version': locked['page_version'],
              'target_revision_id': added['revision_id']}, headers=key())
    assert restore.status_code == 422  # Historical snapshots must not implicitly unlock.
    undone = patch(client, project, locked, [
        {'op': 'set_locked', 'element_id': element['id'], 'locked': False}])
    assert undone.status_code == 200
    assert undone.json['data']['scene'] == added['scene']
    redone = patch(client, project, undone.json['data'], [
        {'op': 'set_locked', 'element_id': element['id'], 'locked': True}])
    assert redone.status_code == 200
    assert redone.json['data']['scene'] == locked['scene']
    assert redone.json['data']['page_version'] == locked['page_version'] + 2
    with app.app_context():
        assert db.session.get(PageSceneRevision, added['revision_id']).scene_json == added['scene']
        assert db.session.get(PageSceneRevision, locked['revision_id']).scene_json == locked['scene']


def test_snapshot_export_keeps_old_text_after_new_edit(client, app):
    project, page = create(client, app)
    element = text_element('快照中的旧文字')
    saved = patch(client, project, page, [{'op': 'add_element', 'element': element}]).json['data']
    snapshot = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': 0, 'pages': [{'page_id': page['page_id'],
                                        'page_version': saved['page_version'],
                                        'revision_id': saved['revision_id']}]}, headers=key())
    assert snapshot.status_code == 201
    changed = patch(client, project, saved, [{'op': 'set_text',
        'element_id': element['id'], 'text': '编辑后的新文字'}])
    assert changed.status_code == 200
    with app.app_context():
        snap = db.session.get(DeckSnapshot, snapshot.json['data']['snapshot_id'])
        revision = db.session.get(PageSceneRevision, saved['revision_id'])
        output = render_pptx(snap, {revision.id: revision}, {})
        verify_pptx(output, snap, {revision.id: revision})
        from pptx import Presentation
        presentation = Presentation(BytesIO(output))
        assert presentation.slides[0].shapes[0].text == '快照中的旧文字'


@pytest.mark.parametrize('fmt', ['pptx', 'pdf'])
def test_queued_export_keeps_snapshot_order_and_text_after_reorder(client, app, tmp_path, fmt):
    if not os.environ.get('SCENE_CHROMIUM_PATH'):
        pytest.skip('Set SCENE_CHROMIUM_PATH to run the real-browser text layout export regression')
    project, first = create(client, app)
    project_id = project['project_id']
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'snapshot-order-assets')
    second_response = client.post(f'/api/v2/projects/{project_id}/pages', json={
        'base_project_version': 0, 'outline': {'title': '第二页', 'points': []}}, headers=key())
    assert second_response.status_code == 201
    second = second_response.json['data']
    first_text = text_element('原始第一页')
    second_text = text_element('原始第二页')
    first_saved = patch(client, project, first, [{'op': 'add_element', 'element': first_text}]).json['data']
    second_saved = patch(client, project, second, [{'op': 'add_element', 'element': second_text}]).json['data']
    frozen = client.post(f'/api/v2/projects/{project_id}/snapshots', json={
        'project_version': 1, 'pages': [{
            'page_id': first['page_id'], 'page_version': first_saved['page_version'],
            'revision_id': first_saved['revision_id']}, {
            'page_id': second['page_id'], 'page_version': second_saved['page_version'],
            'revision_id': second_saved['revision_id']}]}, headers=key())
    assert frozen.status_code == 201
    queued = client.post(f'/api/v2/projects/{project_id}/exports', json={
        'snapshot_id': frozen.json['data']['snapshot_id'], 'format': fmt, 'options': {}}, headers=key())
    assert queued.status_code == 202
    reordered = client.patch(f'/api/v2/projects/{project_id}/outline', json={
        'base_project_version': 1, 'pages': [
            {'page_id': second['page_id'], 'outline': {'title': '现在第一页', 'points': []}},
            {'page_id': first['page_id'], 'outline': {'title': '现在第二页', 'points': []}}]}, headers=key())
    assert reordered.status_code == 200
    assert [page['page_id'] for page in reordered.json['data']['pages']] == [second['page_id'], first['page_id']]
    current_first = reordered.json['data']['pages'][1]
    changed = patch(client, project, current_first, [{'op': 'set_text',
        'element_id': first_text['id'], 'text': '快照之后的修改'}])
    assert changed.status_code == 200
    assert run_once()
    exported = client.get(f'/api/v2/projects/{project_id}/exports/{queued.json["data"]["export_id"]}').json['data']
    assert exported['status'] == 'needs_review'
    report = client.get(exported['report_url']).json
    inventory = report['checks']['pages'][0]
    assert inventory['scene_hash'] == first_saved['scene_hash']
    assert inventory['elements'][0]['element_id'] == first_text['id']
    assert inventory['elements'][0]['role'] == 'title'
    assert inventory['role_counts'] == {'text': {'title': 1}, 'image': {}}
    assert any(w['code'] == 'PPTX_FONT_NOT_EMBEDDED' for w in report['font_warnings']) == (fmt == 'pptx')
    payload = client.get(exported['review_file_url']).data
    if fmt == 'pptx':
        from pptx import Presentation
        deck = Presentation(BytesIO(payload))
        assert [slide.shapes[0].text for slide in deck.slides] == ['原始第一页', '原始第二页']
    else:
        import pypdfium2 as pdfium
        deck = pdfium.PdfDocument(payload)
        assert [deck[index].get_textpage().get_text_range().replace('\r', '').replace('\n', '')
                for index in range(len(deck))] == ['原始第一页', '原始第二页']


@pytest.mark.parametrize('fmt', ['pptx', 'pdf'])
def test_explicit_subset_export_skips_unselected_invalid_page(client, app, tmp_path, fmt):
    if not os.environ.get('SCENE_CHROMIUM_PATH'):
        pytest.skip('Set SCENE_CHROMIUM_PATH to run the real-browser text layout export regression')
    project, first = create(client, app)
    project_id = project['project_id']
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'subset-assets')
    added = client.post(f'/api/v2/projects/{project_id}/pages', json={
        'base_project_version': project['project_version'],
        'outline': {'title': '仅导出此页', 'points': []}}, headers=key())
    assert added.status_code == 201, added.json
    second = added.json['data']
    excluded = patch(client, project, first, [{'op': 'add_element', 'element': text_element('😀')}])
    assert excluded.status_code == 200, excluded.json
    selected = patch(client, project, second, [{'op': 'add_element', 'element': text_element('仅此页文字')}])
    assert selected.status_code == 200, selected.json
    head = selected.json['data']
    frozen = client.post(f'/api/v2/projects/{project_id}/snapshots', json={
        'project_version': project['project_version'] + 1,
        'pages': [{'page_id': second['page_id'], 'page_version': head['page_version'],
                   'revision_id': head['revision_id']}]}, headers=key())
    assert frozen.status_code == 201, frozen.json
    with app.app_context():
        snapshot = db.session.get(DeckSnapshot, frozen.json['data']['snapshot_id'])
        assert [item['page_id'] for item in snapshot.manifest_json['pages']] == [second['page_id']]
    queued = client.post(f'/api/v2/projects/{project_id}/exports', json={
        'snapshot_id': frozen.json['data']['snapshot_id'], 'format': fmt, 'options': {}}, headers=key())
    assert queued.status_code == 202, queued.json
    assert run_once()
    result = client.get(f'/api/v2/projects/{project_id}/exports/{queued.json["data"]["export_id"]}').json['data']
    assert result['status'] == 'needs_review'
    payload = client.get(result['review_file_url']).data
    if fmt == 'pptx':
        from pptx import Presentation
        slides = Presentation(BytesIO(payload)).slides
        assert len(slides) == 1 and slides[0].shapes[0].text == '仅此页文字'
    else:
        import pypdfium2 as pdfium
        document = pdfium.PdfDocument(payload)
        assert len(document) == 1
        assert '仅此页文字' in document[0].get_textpage().get_text_range()


def test_owner_idempotency_and_stale_candidate(client, app):
    project, page = create(client, app)
    path = f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/scene'
    command = {'base_revision_id': page['revision_id'], 'base_page_version': page['page_version'],
               'commands': [{'op': 'add_element', 'element': text_element()}]}
    same_key = key()
    first = client.patch(path, json=command, headers=same_key)
    assert first.status_code == 200
    repeated = client.patch(path, json=command, headers=same_key)
    assert repeated.status_code == 200
    assert repeated.json['data']['revision_id'] == first.json['data']['revision_id']
    changed = client.patch(path, json={**command, 'commands': [{'op': 'add_element', 'element': text_element('不同')}]}, headers=same_key)
    assert changed.status_code == 409
    with app.app_context():
        foreign = Principal(id=str(uuid4()), kind='sub2api_user', display_name='Other')
        db.session.add(foreign)
        db.session.flush()
        candidate = SceneCandidate(id=str(uuid4()), owner_id=LOCAL_OWNER_ID, project_id=project['project_id'],
            page_id=page['page_id'], proposed_revision_id=page['revision_id'],
            base_revision_id=page['revision_id'], base_page_version=page['page_version'],
            scope_json={'kind': 'page'}, change_summary_json={})
        db.session.add(candidate)
        db.session.commit()
        candidate_id = candidate.id
        foreign_id = foreign.id
    rejected = client.post(f'/api/v2/projects/{project["project_id"]}/candidates/accept',
                           json={'candidate_ids': [candidate_id]}, headers=key())
    assert rejected.status_code == 409
    with app.app_context():
        another = Project(id=str(uuid4()), owner_id=foreign_id, editor_mode='scene_v1',
                          project_title='Other', canvas_width_pt=960, canvas_height_pt=540,
                          font_manifest_id='fonts-v1')
        db.session.add(another); db.session.commit()
        another_id = another.id
    assert client.get(f'/api/v2/projects/{another_id}').status_code == 404
    assert all(item['project_id'] != another_id for item in client.get('/api/v2/projects').json['data']['items'])


def test_reject_candidate_replays_original_response(client, app):
    project, page = create(client, app)
    with app.app_context():
        candidate = SceneCandidate(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], page_id=page['page_id'],
            proposed_revision_id=page['revision_id'], base_revision_id=page['revision_id'],
            base_page_version=page['page_version'], scope_json={'kind': 'page'},
            change_summary_json={})
        db.session.add(candidate)
        db.session.commit()
        candidate_id = candidate.id
    url = f'/api/v2/projects/{project["project_id"]}/candidates/{candidate_id}/reject'
    headers = key()
    first = client.post(url, json={}, headers=headers)
    replay = client.post(url, json={}, headers=headers)
    assert first.status_code == replay.status_code == 200
    assert first.json['data'] == replay.json['data'] == {
        'candidate_id': candidate_id, 'state': 'rejected'}


def test_snapshot_requires_pending_candidate_choice_and_blocks_first_generation(client, app):
    def add_pending(project, page):
        with app.app_context():
            db.session.add(SceneCandidate(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
                project_id=project['project_id'], page_id=page['page_id'],
                proposed_revision_id=page['revision_id'], base_revision_id=page['revision_id'],
                base_page_version=page['page_version'], scope_json={'kind': 'page'},
                change_summary_json={}))
            db.session.commit()

    def snapshot(project, page, policy=None):
        body = {'project_version': project['project_version'], 'pages': [{
            'page_id': page['page_id'], 'page_version': page['page_version'],
            'revision_id': page['revision_id']}],
            **({'pending_candidate_policy': policy} if policy else {})}
        return client.post(f'/api/v2/projects/{project["project_id"]}/snapshots',
                           json=body, headers=key())

    blank_project, blank_page = create(client, app)
    add_pending(blank_project, blank_page)
    first = snapshot(blank_project, blank_page, 'export_accepted')
    assert first.status_code == 422
    assert first.json['error']['code'] == 'FIRST_GENERATION_UNACCEPTED'

    project, page = create(client, app)
    saved = patch(client, project, page, [{'op': 'add_element',
        'element': text_element('已接受的文字')}]).json['data']
    add_pending(project, saved)
    missing = snapshot(project, saved)
    assert missing.status_code == 422
    assert missing.json['error']['code'] == 'PENDING_CANDIDATE_CHOICE_REQUIRED'
    accepted = snapshot(project, saved, 'export_accepted')
    assert accepted.status_code == 201


def test_project_cursor_pagination_and_scope(client, app):
    first, _ = create(client, app)
    second, _ = create(client, app)
    page = client.get('/api/v2/projects?limit=1')
    assert page.status_code == 200
    assert len(page.json['data']['items']) == 1
    cursor = page.json['data']['next_cursor']
    assert cursor == page.json['data']['items'][0]['project_id']
    following = client.get(f'/api/v2/projects?limit=1&cursor={cursor}')
    assert following.status_code == 200
    assert following.json['data']['items'][0]['project_id'] in {
        first['project_id'], second['project_id']}
    assert following.json['data']['items'][0]['project_id'] != cursor
    assert following.json['data']['next_cursor'] is None
    assert client.get('/api/v2/projects?limit=0').status_code == 422
    assert client.get('/api/v2/projects?cursor=not-a-project').status_code == 422


def test_revision_candidate_task_and_credential_cursors(client, app):
    project, page = create(client, app)
    first = patch(client, project, page, [{'op': 'add_element',
        'element': text_element('版本一')}]).json['data']
    patch(client, project, first, [{'op': 'set_text',
        'element_id': first['scene']['elements'][0]['id'], 'text': '版本二'}])
    revisions_path = f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/revisions'
    revisions = client.get(revisions_path + '?limit=1').json['data']
    assert len(revisions['items']) == 1 and revisions['next_cursor']
    older = client.get(revisions_path + f'?limit=1&cursor={revisions["next_cursor"]}').json['data']
    assert older['items'][0]['seq'] < revisions['items'][0]['seq']
    assert client.get(revisions_path + '?cursor=foreign').status_code == 422

    with app.app_context():
        for _ in range(2):
            db.session.add(SceneCandidate(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
                project_id=project['project_id'], page_id=page['page_id'],
                proposed_revision_id=first['revision_id'], base_revision_id=page['revision_id'],
                base_page_version=page['page_version'], scope_json={'kind': 'page'},
                change_summary_json={}))
            db.session.add(SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
                project_id=project['project_id'], operation='export', resource_id=str(uuid4()),
                input_json={}, input_hash='a' * 64))
        db.session.commit()
    for endpoint in ('candidates', 'tasks'):
        path = f'/api/v2/projects/{project["project_id"]}/{endpoint}'
        first_page = client.get(path + '?limit=1').json['data']
        second_page = client.get(path + f'?limit=1&cursor={first_page["next_cursor"]}').json['data']
        assert len(first_page['items']) == len(second_page['items']) == 1
        assert first_page['items'] != second_page['items']
        assert second_page['next_cursor'] is None
    candidates_path = f'/api/v2/projects/{project["project_id"]}/candidates'
    assert len(client.get(candidates_path + '?state=pending').json['data']['items']) == 2
    assert client.get(candidates_path + '?state=unexpected').status_code == 422

    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'e' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    for index in range(2):
        response = client.post('/api/v2/model-credentials', json={'label': f'Key {index}',
            'base_url': 'https://model.example/v1', 'api_key': f'sandbox-key-{index}'}, headers=key())
        assert response.status_code == 201
    first_page = client.get('/api/v2/model-credentials?limit=1').json['data']
    second_page = client.get('/api/v2/model-credentials?limit=1&cursor=' +
        first_page['next_cursor']).json['data']
    assert first_page['items'][0]['credential_id'] != second_page['items'][0]['credential_id']
    assert second_page['next_cursor'] is None


def test_image_template_upload_rejects_spoofed_format_and_records_actual_mime(client, app, tmp_path):
    project, _ = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'template-source-assets')
    project_id = project['project_id']
    url = f'/api/v2/projects/{project_id}/template-documents'
    gif = BytesIO()
    Image.new('RGB', (8, 8), 'blue').save(gif, format='GIF')
    rejected = client.post(url, data={'file': (BytesIO(gif.getvalue()), 'reference.png')},
                           content_type='multipart/form-data', headers=key())
    assert rejected.status_code == 422
    assert rejected.json['error']['code'] == 'TEMPLATE_INVALID'
    jpeg = BytesIO()
    Image.new('RGB', (8, 8), 'blue').save(jpeg, format='JPEG')
    accepted = client.post(url, data={'file': (BytesIO(jpeg.getvalue()), 'reference.png')},
                           content_type='multipart/form-data', headers=key())
    assert accepted.status_code == 202, accepted.json
    with app.app_context():
        document = db.session.get(TemplateDocument, accepted.json['data']['document_id'])
        source = db.session.get(Asset, document.source_asset_id)
        assert source.mime_type == 'image/jpeg'
        assert source.storage_key.endswith('.jpg')


def test_image_template_import_and_private_preview(client, app, tmp_path):
    project, _ = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'scene-assets')
    image = Image.new('RGB', (640, 360), '#24415c')
    output = BytesIO(); image.save(output, format='PNG'); output.seek(0)
    source_bytes = output.getvalue()
    upload_headers = key()
    response = client.post(f'/api/v2/projects/{project["project_id"]}/template-documents',
        data={'file': (output, 'reference.png')}, content_type='multipart/form-data', headers=upload_headers)
    assert response.status_code == 202, response.json
    repeated = client.post(f'/api/v2/projects/{project["project_id"]}/template-documents',
        data={'file': (BytesIO(source_bytes), 'reference.png')}, content_type='multipart/form-data',
        headers=upload_headers)
    assert repeated.status_code == 202 and repeated.json['data'] == response.json['data']
    different = client.post(f'/api/v2/projects/{project["project_id"]}/template-documents',
        data={'file': (BytesIO(b'different'), 'reference.png')}, content_type='multipart/form-data',
        headers=upload_headers)
    assert different.status_code == 409
    with app.app_context():
        assert TemplateDocument.query.filter_by(project_id=project['project_id']).count() == 1
        assert SceneTaskItem.query.filter_by(project_id=project['project_id'], operation='import_template').count() == 1
    from workers.scene_runner import run_once
    assert run_once()
    doc = client.get(f'/api/v2/projects/{project["project_id"]}/template-documents/{response.json["data"]["document_id"]}')
    assert doc.json['data']['status'] == 'preview_ready'
    invalid_selection = client.post(
        f'/api/v2/projects/{project["project_id"]}/template-documents/{response.json["data"]["document_id"]}/analyze',
        json={'selected_page_indexes': [[]]}, headers=key())
    assert invalid_selection.status_code == 422
    assert invalid_selection.json['error']['code'] == 'PAGE_SELECTION_INVALID'
    assets = client.get(f'/api/v2/projects/{project["project_id"]}/template-assets').json['data']['items']
    assert len(assets) == 1
    assert client.get(assets[0]['preview_url']).status_code == 200
    assert client.get(assets[0]['preview_url'].split('?')[0]).status_code == 404
    valid_style = {'schema_version': 1, 'role': 'content',
        'palette': {'background': '#ffffff', 'text': '#172a42', 'accent': '#f7cc35'},
        'font_suggestions': [], 'layout_hints': [], 'decorative_hints': [],
        'content_density': 'medium', 'warnings': []}
    style_url = f'/api/v2/projects/{project["project_id"]}/template-assets/{assets[0]["template_asset_id"]}'
    invalid = client.patch(style_url, json={'base_analysis_revision': 0,
        'analysis': {**valid_style, 'palette': {'background': 'javascript:alert(1)',
            'text': '#172A42', 'accent': '#F7CC35'}}}, headers=key())
    assert invalid.status_code == 422
    assert invalid.json['error']['code'] == 'ANALYSIS_INVALID'
    style_headers = key()
    saved = client.patch(style_url, json={'base_analysis_revision': 0, 'analysis': valid_style},
                         headers=style_headers)
    assert saved.status_code == 200
    assert saved.json['data']['analysis_revision'] == 1
    replayed_style = client.patch(style_url, json={'base_analysis_revision': 0, 'analysis': valid_style},
                                  headers=style_headers)
    assert replayed_style.status_code == 200 and replayed_style.json['data'] == saved.json['data']
    page = project['pages'][0]
    bind_url = f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/template'
    bind_body = {'base_page_version': page['page_version'],
                 'template_asset_id': assets[0]['template_asset_id'], 'style_text': '简洁商务'}
    invalid_style = client.patch(bind_url, json={**bind_body, 'style_text': {'unbounded': 'object'}},
                                 headers=key())
    assert invalid_style.status_code == 422
    assert invalid_style.json['error']['code'] == 'STYLE_TEXT_INVALID'
    bind_headers = key()
    bound = client.patch(bind_url, json=bind_body, headers=bind_headers)
    repeated_bind = client.patch(bind_url, json=bind_body, headers=bind_headers)
    assert bound.status_code == repeated_bind.status_code == 200
    assert bound.json['data'] == repeated_bind.json['data']
    assert client.get(f'/api/v2/projects/{project["project_id"]}').json['data']['project_version'] == 1
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'wrong-worker-volume')
    missing = client.get(assets[0]['preview_url'])
    assert missing.status_code == 503
    assert missing.json['error']['code'] == 'ASSET_UNAVAILABLE'
    unavailable_plan = client.post(f'/api/v2/projects/{project["project_id"]}/generation-plans',
        json={'base_project_version': 1}, headers=key())
    assert unavailable_plan.status_code == 503
    assert unavailable_plan.json['error']['code'] == 'ASSET_UNAVAILABLE'


def test_unanalysed_template_is_sent_as_bounded_visual_reference(client, app, tmp_path, monkeypatch):
    project, page = create(client, app)
    project_id = project['project_id']
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'reference-generation-assets')
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'r' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Mine',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    source = BytesIO(); Image.new('RGB', (640, 360), '#24415c').save(source, format='PNG'); source.seek(0)
    upload = client.post(f'/api/v2/projects/{project_id}/template-documents',
        data={'file': (source, 'visual.png')}, content_type='multipart/form-data', headers=key())
    assert upload.status_code == 202
    assert run_once()
    reference = client.get(f'/api/v2/projects/{project_id}/template-assets').json['data']['items'][0]
    assert reference['analysis_status'] != 'completed'
    bound = client.patch(f'/api/v2/projects/{project_id}/pages/{page["page_id"]}/template',
        json={'base_page_version': page['page_version'],
              'template_asset_id': reference['template_asset_id'], 'style_text': ''}, headers=key())
    assert bound.status_code == 200
    configured = client.patch(f'/api/v2/projects/{project_id}/model-config', json={
        'base_project_version': 1,
        'model_config': {'text_model': 'vision-text', 'image_model': 'image-model'}}, headers=key())
    assert configured.status_code == 200
    plan = client.post(f'/api/v2/projects/{project_id}/generation-plans',
        json={'base_project_version': 2}, headers=key())
    assert plan.status_code == 201
    current_page = configured.json['data']['pages'][0]
    task = client.post(f'/api/v2/projects/{project_id}/generation-tasks', json={
        'plan_id': plan.json['data']['plan_id'], 'credential_id': credential['credential_id'],
        'targets': [{'page_id': page['page_id'], 'base_revision_id': current_page['revision_id'],
                     'base_page_version': current_page['page_version']}]}, headers=key())
    assert task.status_code == 202
    observed = []
    def fake_model_call(_item, _attempt, kind, prompt, *_args, **_kwargs):
        assert kind == 'scene'
        observed.append(prompt)
        return {'background': {'kind': 'solid', 'color': '#FFFFFF'},
                'elements': [text_element('新正文')], 'image_request': None}
    monkeypatch.setattr('services.scene.generation._model_call', fake_model_call)
    assert run_once()
    assert len(observed) == 1 and [part['type'] for part in observed[0]] == ['text', 'image_url']
    url = observed[0][1]['image_url']['url']
    assert url.startswith('data:image/jpeg;base64,')
    with Image.open(BytesIO(base64.b64decode(url.split(',', 1)[1]))) as image:
        assert image.size == (640, 360)
        assert image.getpixel((320, 180))[2] > image.getpixel((320, 180))[0]
    assert run_once()  # Finish the image-free DraftScene locally.
    analysis = client.post(f'/api/v2/projects/{project_id}/template-documents/{upload.json["data"]["document_id"]}/analyze',
        json={'credential_id': credential['credential_id'], 'selected_page_indexes': [1]}, headers=key())
    assert analysis.status_code == 202
    duplicate = client.post(f'/api/v2/projects/{project_id}/template-documents/{upload.json["data"]["document_id"]}/analyze',
        json={'credential_id': credential['credential_id'], 'selected_page_indexes': [1]}, headers=key())
    assert duplicate.status_code == 409
    assert duplicate.json['error']['code'] == 'ANALYSIS_IN_PROGRESS'
    with app.app_context():
        assert SceneTaskItem.query.filter_by(project_id=project_id, operation='analyze_template').count() == 1
    style = {'schema_version': 1, 'role': 'content',
        'palette': {'background': '#24415C', 'text': '#FFFFFF', 'accent': '#F7CC35'},
        'font_suggestions': [], 'layout_hints': [], 'decorative_hints': [],
        'content_density': 'medium', 'warnings': []}
    monkeypatch.setattr('services.templates.style_analysis._model_call',
                        lambda *_args, **_kwargs: style)
    assert run_once()
    analyzed = client.get(f'/api/v2/projects/{project_id}/template-assets').json['data']['items'][0]
    assert analyzed['analysis_status'] == 'completed'
    analyzed_plan = client.post(f'/api/v2/projects/{project_id}/generation-plans',
        json={'base_project_version': 2}, headers=key())
    assert analyzed_plan.status_code == 201
    analyzed_task = client.post(f'/api/v2/projects/{project_id}/generation-tasks', json={
        'plan_id': analyzed_plan.json['data']['plan_id'], 'credential_id': credential['credential_id'],
        'targets': [{'page_id': page['page_id'], 'base_revision_id': current_page['revision_id'],
                     'base_page_version': current_page['page_version']}]}, headers=key())
    assert analyzed_task.status_code == 202
    assert run_once()
    assert isinstance(observed[-1], str) and '#24415C' in observed[-1]
    assert run_once()
    retry = client.post(f'/api/v2/projects/{project_id}/generation-tasks', json={
        'plan_id': analyzed_plan.json['data']['plan_id'], 'credential_id': credential['credential_id'],
        'targets': [{'page_id': page['page_id'], 'base_revision_id': current_page['revision_id'],
                     'base_page_version': current_page['page_version']}]}, headers=key())
    assert retry.status_code == 202
    with app.app_context():
        from services.scene.store import path_for
        saved_plan = db.session.get(GenerationPlan, analyzed_plan.json['data']['plan_id'])
        preview_id = saved_plan.manifest_json['pages'][0]['reference']['preview_asset_id']
        path_for(db.session.get(Asset, preview_id)).write_bytes(b'corrupt after plan confirmation')
    assert run_once()
    assert len(observed) == 2  # No third paid request was dispatched.
    failed = client.get(f'/api/v2/tasks/{retry.json["data"]["tasks"][0]["task_id"]}').json['data']
    assert failed['state'] == 'failed' and failed['error_code'] == 'ASSET_UNAVAILABLE'


def test_template_analysis_submission_replays_after_response_loss(client, app, tmp_path, monkeypatch):
    project, _ = create(client, app)
    project_id = project['project_id']
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'template-analysis-assets')
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'a' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Mine',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    configured = client.patch(f'/api/v2/projects/{project_id}/model-config', json={
        'base_project_version': project['project_version'],
        'model_config': {'text_model': 'fake-text'}}, headers=key())
    assert configured.status_code == 200
    image = BytesIO(); Image.new('RGB', (640, 360), '#24415c').save(image, format='PNG'); image.seek(0)
    uploaded = client.post(f'/api/v2/projects/{project_id}/template-documents',
        data={'file': (image, 'reference.png')}, content_type='multipart/form-data', headers=key())
    assert uploaded.status_code == 202
    assert run_once()
    document_id = uploaded.json['data']['document_id']
    url = f'/api/v2/projects/{project_id}/template-documents/{document_id}/analyze'
    payload = {'credential_id': credential['credential_id'], 'selected_page_indexes': [1]}
    headers = key()
    first = client.post(url, json=payload, headers=headers)
    assert first.status_code == 202, first.json
    repeated = client.post(url, json=payload, headers=headers)
    assert repeated.status_code == 202
    assert repeated.json['data'] == first.json['data']
    with app.app_context():
        assert SceneTaskItem.query.filter_by(project_id=project_id, operation='analyze_template').count() == 1
    assert client.post(url, json={**payload, 'selected_page_indexes': [2]}, headers=headers).status_code == 409
    with app.app_context():
        from services.scene.store import path_for
        queued = SceneTaskItem.query.filter_by(project_id=project_id, operation='analyze_template').one()
        preview_id = queued.input_json['preview_asset_id']
        assert queued.input_json['preview_sha256'] == db.session.get(Asset, preview_id).sha256
        preview_path = path_for(db.session.get(Asset, preview_id))
        original_preview = preview_path.read_bytes()
        preview_path.write_bytes(b'corrupt after task submission')
    monkeypatch.setattr('services.templates.style_analysis._model_call',
                        lambda *_args, **_kwargs: pytest.fail('paid analysis must not run with corrupt input'))
    assert run_once()
    failed = client.get(f'/api/v2/tasks/{first.json["data"]["tasks"][0]["task_id"]}').json['data']
    assert failed['state'] == 'failed' and failed['error_code'] == 'ASSET_UNAVAILABLE'
    preview_path.write_bytes(original_preview)
    reference = client.get(f'/api/v2/projects/{project_id}/template-assets').json['data']['items'][0]
    profile = {'schema_version': 1, 'role': 'content',
        'palette': {'background': '#FFFFFF', 'text': '#000000', 'accent': '#FFCC00'},
        'font_suggestions': [], 'layout_hints': [], 'decorative_hints': [],
        'content_density': 'medium', 'warnings': []}
    corrected = client.patch(f'/api/v2/projects/{project_id}/template-assets/{reference["template_asset_id"]}',
        json={'base_analysis_revision': reference['analysis_revision'], 'analysis': profile}, headers=key())
    assert corrected.status_code == 200, corrected.json
    document = client.get(f'/api/v2/projects/{project_id}/template-documents/{document_id}').json['data']
    assert document['status'] == 'ready'


def test_pptx_upload_refuses_unisolated_renderer(client, app):
    project, _ = create(client, app)
    response = client.post(f'/api/v2/projects/{project["project_id"]}/template-documents',
        data={'file': (BytesIO(b'not even a pptx'), 'reference.pptx')},
        content_type='multipart/form-data', headers=key())
    assert response.status_code == 503
    assert response.json['error']['code'] == 'TEMPLATE_RENDERER_UNAVAILABLE'


def test_style_profile_rejects_unbounded_or_outside_regions():
    base = {'schema_version': 1, 'role': 'content',
        'palette': {'background': '#FFFFFF', 'text': '#000000', 'accent': '#FFCC00'},
        'font_suggestions': [], 'layout_hints': [], 'decorative_hints': [],
        'content_density': 'medium', 'warnings': []}
    assert normalize_style_profile(base)['role'] == 'content'
    with pytest.raises(SceneError) as raised:
        normalize_style_profile({**base, 'layout_hints': [{'region': 'title', 'x': .9,
            'y': .1, 'w': .3, 'h': .2}]})
    assert raised.value.code == 'ANALYSIS_INVALID'


@pytest.mark.parametrize('edit_before_publish', [False, True])
def test_generation_candidate_does_not_overwrite_manual_edit(client, app, monkeypatch, edit_before_publish):
    project, page = create(client, app)
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'x' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Mine',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key())
    assert credential.status_code == 201, credential.json
    model = client.patch(f'/api/v2/projects/{project["project_id"]}/model-config',
        json={'base_project_version': 0, 'model_config': {'text_model': 'fake-text', 'image_model': 'fake-image'}}, headers=key())
    assert model.status_code == 200
    plan = client.post(f'/api/v2/projects/{project["project_id"]}/generation-plans',
        json={'base_project_version': 1}, headers=key())
    assert plan.status_code == 201
    task = client.post(f'/api/v2/projects/{project["project_id"]}/generation-tasks', json={
        'plan_id': plan.json['data']['plan_id'], 'credential_id': credential.json['data']['credential_id'],
        'targets': [{'page_id': page['page_id'], 'base_revision_id': page['revision_id'],
                     'base_page_version': page['page_version']}]}, headers=key())
    assert task.status_code == 202, task.json
    text = text_element('AI 候选')
    monkeypatch.setattr('services.scene.generation._model_call', lambda *args, **kwargs: {
        'background': {'kind': 'solid', 'color': '#FFFFFF'},
        'elements': [text], 'image_request': None})
    from workers.scene_runner import run_once
    assert run_once()
    if edit_before_publish:
        assert patch(client, project, page, [{'op': 'add_element', 'element': text_element('人工修改')}]).status_code == 200
    assert run_once()  # Local finalize of the durable, image-free DraftScene.
    candidates = client.get(f'/api/v2/projects/{project["project_id"]}/candidates').json['data']['items']
    assert len(candidates) == 1 and candidates[0]['state'] == 'pending'
    summary = candidates[0]['change_summary']
    assert summary['base_revision_id'] == page['revision_id']
    assert summary['added'] == [{'element_id': text['id'], 'kind': 'text', 'role': text['role'], 'label': 'AI 候选'}]
    assert summary['removed'] == [] and summary['modified'] == []
    if not edit_before_publish:
        changed = patch(client, project, page, [{'op': 'add_element', 'element': text_element('人工修改')}])
        assert changed.status_code == 200
    accept = client.post(f'/api/v2/projects/{project["project_id"]}/candidates/accept',
                         json={'candidate_ids': [candidates[0]['candidate_id']]}, headers=key())
    assert accept.status_code == 409
    after = client.get(f'/api/v2/projects/{project["project_id"]}/candidates').json['data']['items']
    assert after[0]['state'] == 'stale'
    current = client.get(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/scene').json['data']
    assert current['scene']['elements'][0]['text'] == '人工修改'


def test_generated_background_and_illustration_are_private_assets(client, app, monkeypatch, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'generated-assets')
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'y' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Mine',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    model = client.patch(f'/api/v2/projects/{project["project_id"]}/model-config',
        json={'base_project_version': 0, 'model_config': {'text_model': 'fake-text', 'image_model': 'fake-image'}}, headers=key())
    assert model.status_code == 200
    plan = client.post(f'/api/v2/projects/{project["project_id"]}/generation-plans',
        json={'base_project_version': 1}, headers=key()).json['data']
    task = client.post(f'/api/v2/projects/{project["project_id"]}/generation-tasks', json={
        'plan_id': plan['plan_id'], 'credential_id': credential['credential_id'],
        'targets': [{'page_id': page['page_id'], 'base_revision_id': page['revision_id'],
                     'base_page_version': page['page_version']}]}, headers=key())
    assert task.status_code == 202
    buffer = BytesIO(); Image.new('RGB', (128, 128), '#224466').save(buffer, format='PNG')
    def fake_model_call(_item, _attempt, kind, *_args, **_kwargs):
        if kind == 'scene':
            return {'background': {'kind': 'solid', 'color': '#FFFFFF'},
                    'elements': [text_element('可编辑标题')],
                    'image_requests': [
                        {'role': 'background', 'prompt': '蓝色抽象背景'},
                        {'role': 'illustration', 'prompt': '简单插画',
                         'frame': {'x': 620, 'y': 150, 'w': 280, 'h': 250, 'rotation_deg': 0}}]}
        return buffer.getvalue()
    monkeypatch.setattr('services.scene.generation._model_call', fake_model_call)
    assert run_once()
    assert run_once()  # Background asset
    assert run_once()  # Illustration asset
    assert run_once()  # Resolve ready assets and publish the candidate.
    candidates = client.get(f'/api/v2/projects/{project["project_id"]}/candidates').json['data']['items']
    assert 'IMAGE_TEXT_UNVERIFIED' in candidates[0]['change_summary']['warnings']
    scene = client.get(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/scene',
                       query_string={'revision_id': candidates[0]['revision_id']}).json['data']
    assert scene['scene']['background']['kind'] == 'image'
    assert sum(e['kind'] == 'image' for e in scene['scene']['elements']) == 1
    assert len(scene['asset_urls']) == 2
    assert all(client.get(url).status_code == 200 for url in scene['asset_urls'].values())
    accept_key = key()
    accept_url = f'/api/v2/projects/{project["project_id"]}/candidates/accept'
    accept_body = {'candidate_ids': [candidates[0]['candidate_id']]}
    accepted = client.post(accept_url, json=accept_body, headers=accept_key)
    assert accepted.status_code == 200
    head = accepted.json['data'][0]
    assert len(head['asset_urls']) == 2
    restore_url = f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/restore'
    undone = client.post(restore_url, json={'base_revision_id': head['revision_id'],
        'base_page_version': head['page_version'], 'target_revision_id': page['revision_id']}, headers=key())
    assert undone.status_code == 200
    assert undone.json['data']['scene']['elements'] == []
    redone = client.post(restore_url, json={'base_revision_id': undone.json['data']['revision_id'],
        'base_page_version': undone.json['data']['page_version'], 'target_revision_id': head['revision_id']}, headers=key())
    assert redone.status_code == 200
    assert redone.json['data']['scene_hash'] == head['scene_hash']
    assert len(redone.json['data']['asset_urls']) == 2
    replay = client.post(accept_url, json=accept_body, headers=accept_key)
    assert replay.status_code == 200
    assert replay.json['data'][0]['revision_id'] == head['revision_id']
    assert len(replay.json['data'][0]['asset_urls']) == 2
    assert client.get(f'/api/v2/projects/{project["project_id"]}/candidates').json['data']['items'][0]['state'] == 'accepted'


def _queued_generation(client, app, project, page):
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'z' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Mine',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    assert client.patch(f'/api/v2/projects/{project["project_id"]}/model-config',
        json={'base_project_version': 0, 'model_config': {
            'text_model': 'fake-text', 'image_model': 'fake-image'}}, headers=key()).status_code == 200
    plan = client.post(f'/api/v2/projects/{project["project_id"]}/generation-plans',
        json={'base_project_version': 1}, headers=key()).json['data']
    response = client.post(f'/api/v2/projects/{project["project_id"]}/generation-tasks', json={
        'plan_id': plan['plan_id'], 'credential_id': credential['credential_id'],
        'targets': [{'page_id': page['page_id'], 'base_revision_id': page['revision_id'],
                     'base_page_version': page['page_version']}]}, headers=key())
    assert response.status_code == 202, response.json
    return response.json['data']['tasks'][0]['task_id']


def test_asset_dag_retry_only_failed_image_and_preserve_draft(client, app, monkeypatch, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'dag-assets')
    parent_id = _queued_generation(client, app, project, page)
    buffer = BytesIO(); Image.new('RGB', (64, 64), '#224466').save(buffer, format='PNG')
    calls = {'text': 0, 'image': []}

    def fake_chat(*_args):
        calls['text'] += 1
        return {'background': {'kind': 'solid', 'color': '#FFFFFF'},
            'elements': [text_element('持久草稿')], 'image_requests': [
                {'role': 'background', 'prompt': 'background'},
                {'role': 'illustration', 'prompt': 'illustration', 'frame': {
                    'x': 620, 'y': 150, 'w': 280, 'h': 250, 'rotation_deg': 0}}]}, 'text-request', {}

    def fake_image(_url, _key, _model, prompt):
        calls['image'].append(prompt)
        if 'illustration' in prompt and calls['image'].count(prompt) == 1:
            raise SceneError('MODEL_UNSUPPORTED', 'Rejected before billing', 422)
        return buffer.getvalue(), f'image-{len(calls["image"])}', {}

    monkeypatch.setattr('services.scene.generation.chat_json', fake_chat)
    monkeypatch.setattr('services.scene.generation.generate_image', fake_image)
    assert run_once()  # Draft and two child tasks are committed together.
    with app.app_context():
        parent = db.session.get(SceneTaskItem, parent_id)
        children = SceneTaskItem.query.filter_by(group_id=parent.group_id, operation='generate_asset').all()
        assert parent.state == 'waiting_assets' and len(children) == 2
        assert parent.result_json['draft_asset_id']
    assert run_once() and run_once()
    with app.app_context():
        children = SceneTaskItem.query.filter_by(group_id=parent.group_id, operation='generate_asset').order_by(
            SceneTaskItem.logical_key).all()
        assert [child.state for child in children] == ['succeeded', 'failed']
        first_asset_id = children[0].result_json['asset_id']
        failed_id = children[1].id
        assert db.session.get(SceneTaskItem, parent_id).state == 'waiting_assets'
        assert db.session.get(Asset, first_asset_id).state == 'ready'
    assert client.post(f'/api/v2/tasks/{failed_id}/retry', json={}, headers=key()).status_code == 200
    assert run_once() and run_once()  # Retry only image 2, then local finalize.
    with app.app_context():
        parent = db.session.get(SceneTaskItem, parent_id)
        assert parent.state == 'succeeded'
        assert parent.attempt_count == 2
        assert db.session.get(SceneTaskItem, children[0].id).attempt_count == 1
        assert db.session.get(SceneTaskItem, children[1].id).attempt_count == 2
        assert len(SceneCandidate.query.filter_by(page_id=page['page_id']).all()) == 1
    assert calls['text'] == 1 and len(calls['image']) == 3


def test_asset_unknown_requires_ack_and_cancel_prevents_publish(client, app, monkeypatch, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'unknown-assets')
    parent_id = _queued_generation(client, app, project, page)
    from services.scene.provider import ProviderOutcomeUnknown
    monkeypatch.setattr('services.scene.generation.chat_json', lambda *_: (
        {'background': {'kind': 'solid', 'color': '#FFFFFF'},
         'elements': [text_element('文字')], 'image_requests': [
             {'role': 'background', 'prompt': 'only background'}]}, 'text-request', {}))
    monkeypatch.setattr('services.scene.generation.generate_image',
                        lambda *_: (_ for _ in ()).throw(ProviderOutcomeUnknown()))
    assert run_once() and run_once()
    with app.app_context():
        parent = db.session.get(SceneTaskItem, parent_id)
        child = SceneTaskItem.query.filter_by(group_id=parent.group_id, operation='generate_asset').one()
        assert child.state == 'outcome_unknown'
        child_id, stale_fence, group_id = child.id, child.fence_token, parent.group_id
        assert parent.state == 'waiting_assets'
    assert client.post(f'/api/v2/tasks/{child_id}/retry', json={}, headers=key()).status_code == 422
    assert client.post(f'/api/v2/tasks/{child_id}/retry',
        json={'acknowledge_possible_charge': True}, headers=key()).status_code == 200
    with app.app_context():
        from services.scene.generation import generate_asset
        with pytest.raises(SceneError) as raised:
            generate_asset(child_id, stale_fence, WORKER_ID)
        assert raised.value.code == 'TASK_FENCE_LOST'
        db.session.rollback()
    cancel_key = key()
    first_cancel = client.post(f'/api/v2/tasks/{parent_id}/cancel', json={}, headers=cancel_key)
    replay_cancel = client.post(f'/api/v2/tasks/{parent_id}/cancel', json={}, headers=cancel_key)
    assert first_cancel.status_code == replay_cancel.status_code == 200
    assert first_cancel.json['data'] == replay_cancel.json['data']
    with app.app_context():
        assert db.session.get(SceneTaskItem, child_id).state == 'cancelled'
        assert db.session.get(SceneTaskItem, parent_id).state == 'cancelled'
        assert SceneCandidate.query.filter_by(page_id=page['page_id']).count() == 0
    assert client.get(f'/api/v2/task-groups/{group_id}').json['data']['state'] == 'cancelled'


def test_export_visual_review_binds_exact_report_hash(client, app, monkeypatch, tmp_path):
    from backend.tests.scene_layout_fixtures import single_line_layout
    monkeypatch.setattr('workers.scene_runner.measure_text_layout',
        lambda snapshot, revisions, assets: single_line_layout(snapshot, revisions))
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'review-assets')
    snapshot = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': 0, 'pages': [{'page_id': page['page_id'],
            'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    assert snapshot.status_code == 201, snapshot.json
    export = client.post(f'/api/v2/projects/{project["project_id"]}/exports', json={
        'snapshot_id': snapshot.json['data']['snapshot_id'], 'format': 'pptx', 'options': {}}, headers=key())
    assert export.status_code == 202, export.json
    export_id = export.json['data']['export_id']
    sheet = BytesIO(); Image.new('RGB', (16, 8), '#e0e0e0').save(sheet, format='PNG')
    monkeypatch.setattr('workers.scene_runner.compare_export', lambda *_: ({
        'status': 'needs_review', 'rendered_text_layout': {'version': 1, 'status': 'passed',
            'rendered_pdf_sha256': 'a' * 64, 'pages': [{'page_id': page['page_id']}]},
        'pages': [{'page_id': page['page_id'], 'mean_abs_rgb': 1.5}]}, [sheet.getvalue()]))
    assert run_once()
    result = client.get(f'/api/v2/projects/{project["project_id"]}/exports/{export_id}')
    assert result.status_code == 200
    state = result.json['data']
    assert state['status'] == 'needs_review' and state['download_url'] is None
    assert client.get(state['review_file_url']).status_code == 200
    assert client.get(state['report_url']).status_code == 200
    assert len(state['visual_evidence']) == 1
    assert client.get(state['visual_evidence'][0]['url']).status_code == 200
    rejected = client.post(f'/api/v2/projects/{project["project_id"]}/exports/{export_id}/review',
        json={'report_sha256': '0' * 64,
              'acknowledged_warning_codes': ['VISUAL_DIFF_UNCALIBRATED']}, headers=key())
    assert rejected.status_code == 409
    review_key = key()
    accepted = client.post(f'/api/v2/projects/{project["project_id"]}/exports/{export_id}/review',
        json={'report_sha256': state['report_sha256'],
              'acknowledged_warning_codes': ['VISUAL_DIFF_UNCALIBRATED']}, headers=review_key)
    assert accepted.status_code == 200, accepted.json
    repeated = client.post(f'/api/v2/projects/{project["project_id"]}/exports/{export_id}/review',
        json={'report_sha256': state['report_sha256'],
              'acknowledged_warning_codes': ['VISUAL_DIFF_UNCALIBRATED']}, headers=review_key)
    assert repeated.status_code == 200
    assert repeated.json['data'] == accepted.json['data']
    result = client.get(f'/api/v2/projects/{project["project_id"]}/exports/{export_id}').json['data']
    assert result['status'] == 'succeeded' and client.get(result['download_url']).status_code == 200


def test_export_history_cursor_recovers_queued_jobs(client, app):
    project, page = create(client, app)
    snapshot = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': 0, 'pages': [{'page_id': page['page_id'],
            'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    assert snapshot.status_code == 201
    for _ in range(3):
        response = client.post(f'/api/v2/projects/{project["project_id"]}/exports', json={
            'snapshot_id': snapshot.json['data']['snapshot_id'], 'format': 'pdf',
            'options': {}}, headers=key())
        assert response.status_code == 202
    path = f'/api/v2/projects/{project["project_id"]}/exports'
    first = client.get(path, query_string={'limit': 2}).json['data']
    assert len(first['items']) == 2 and first['next_cursor']
    second = client.get(path, query_string={'limit': 2,
        'cursor': first['next_cursor']}).json['data']
    assert len(second['items']) == 1 and second['next_cursor'] is None
    assert len({item['export_id'] for item in first['items'] + second['items']}) == 3


def test_visual_comparison_keeps_zero_diff_in_review(client, app, monkeypatch):
    from services.exports.visual_compare import compare_export
    document = fitz.open()
    document.new_page(width=100, height=100)
    payload = document.tobytes()
    image = document[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes('png')
    document.close()
    snapshot = SimpleNamespace(manifest_json={'pages': [{'page_id': str(uuid4())}]})
    monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', 'fake-renderer-job-root')
    monkeypatch.setattr('services.exports.visual_compare.scene_html', lambda *_: '<html></html>')
    monkeypatch.setattr('services.exports.visual_compare.run_render_job',
                        lambda *_args, **_kwargs: ([image], {'page_count': 1}))
    report, sheets = compare_export(payload, 'pdf', snapshot, {}, {})
    assert report['status'] == 'needs_review'
    assert report['pages'][0]['mean_abs_rgb'] == 0
    assert report['pages'][0]['changed_fraction_luma_gt_32'] == 0
    assert Image.open(BytesIO(sheets[0])).size[0] > 0


def test_credential_connection_check_does_not_claim_capabilities(client, app, monkeypatch):
    project, _ = create(client, app)
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'c' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Check',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    calls = []
    monkeypatch.setattr('workers.scene_runner.list_models',
                        lambda url, secret: calls.append((url, secret)) or ['text-model', 'image-model'])
    headers = key()
    first = client.post(f'/api/v2/model-credentials/{credential["credential_id"]}/check',
        json={'project_id': project['project_id']}, headers=headers)
    second = client.post(f'/api/v2/model-credentials/{credential["credential_id"]}/check',
        json={'project_id': project['project_id']}, headers=headers)
    assert first.status_code == second.status_code == 202
    task_id = first.json['data']['task_id']
    assert task_id == second.json['data']['task_id']
    assert run_once()
    assert calls == [('https://model.example/v1', 'sandbox-test-key')]
    task = client.get(f'/api/v2/tasks/{task_id}').json['data']
    assert task['state'] == 'succeeded' and task['possible_charge'] is False
    summary = client.get('/api/v2/model-credentials').json['data']['items'][0]
    assert summary['status'] == 'connected'
    assert summary['capabilities']['text_generation_verified'] is False
    assert summary['capabilities']['image_generation_verified'] is False


def test_paid_credential_capability_probes_require_ack_and_track_usage(client, app, monkeypatch):
    project, _ = create(client, app)
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'd' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Probe',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-probe-key'}, headers=key()).json['data']
    route = f'/api/v2/model-credentials/{credential["credential_id"]}/check'
    rejected = client.post(route, json={'project_id': project['project_id'],
        'kind': 'text_generation', 'model_id': 'text-model'}, headers=key())
    assert rejected.status_code == 422
    assert rejected.json['error']['code'] == 'POSSIBLE_CHARGE_ACK_REQUIRED'

    monkeypatch.setattr('services.scene.generation.chat_json', lambda *args: (
        {'ok': True}, 'text-request-id', {'total_tokens': 3}))
    text_task = client.post(route, json={'project_id': project['project_id'],
        'kind': 'text_generation', 'model_id': 'text-model',
        'acknowledge_possible_charge': True}, headers=key())
    assert text_task.status_code == 202 and text_task.json['data']['possible_charge'] is True
    assert run_once()
    text_id = text_task.json['data']['task_id']
    assert client.get(f'/api/v2/tasks/{text_id}').json['data']['state'] == 'succeeded'
    with app.app_context():
        attempt = SceneTaskAttempt.query.filter_by(task_item_id=text_id).one()
        assert attempt.usage_json['credential_text']['total_tokens'] == 3
        assert attempt.dispatch_state == 'resolved'

    picture = BytesIO()
    Image.new('RGB', (64, 64), 'yellow').save(picture, format='PNG')
    monkeypatch.setattr('services.scene.generation.generate_image', lambda *args: (
        picture.getvalue(), 'image-request-id', {'images': 1}))
    image_task = client.post(route, json={'project_id': project['project_id'],
        'kind': 'image_generation', 'model_id': 'image-model',
        'acknowledge_possible_charge': True}, headers=key())
    assert image_task.status_code == 202 and run_once()
    image_id = image_task.json['data']['task_id']
    assert client.get(f'/api/v2/tasks/{image_id}').json['data']['state'] == 'succeeded'
    summary = client.get('/api/v2/model-credentials').json['data']['items'][0]['capabilities']
    assert summary['text_generation_verified'] is True
    assert summary['text_generation_model_id'] == 'text-model'
    assert summary['image_generation_verified'] is True
    assert summary['image_generation_model_id'] == 'image-model'
    with app.app_context():
        expired = SceneTaskItem(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=project['project_id'], credential_id=credential['credential_id'],
            operation='check_credential', resource_id=credential['credential_id'],
            input_json={'kind': 'text_generation', 'model_id': 'text-model'},
            input_hash='f' * 64, state='running', attempt_count=1,
            dispatch_state='may_have_been_sent', lease_owner='crashed-probe',
            lease_expires_at=now() - timedelta(seconds=5), fence_token=3)
        db.session.add(expired)
        db.session.add(SceneTaskAttempt(task_item_id=expired.id, attempt_no=1,
            fence_token=3, worker_id='crashed-probe',
            dispatch_state='may_have_been_sent'))
        db.session.commit()
        assert claim() is None
        db.session.refresh(expired)
        assert expired.state == 'outcome_unknown'
        expired_id = expired.id
    assert client.get(f'/api/v2/tasks/{expired_id}').json['data']['possible_charge'] is True
    no_ack = client.post(f'/api/v2/tasks/{expired_id}/retry', json={}, headers=key())
    assert no_ack.status_code == 422


def test_generated_candidate_preserves_locked_title_without_model_copy(client, app, monkeypatch, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'locked-assets')
    title = text_element('人工锁定标题')
    saved = patch(client, project, page, [{'op': 'add_element', 'element': title}])
    assert saved.status_code == 200
    locked = patch(client, project, saved.json['data'], [
        {'op': 'set_locked', 'element_id': title['id'], 'locked': True}])
    assert locked.status_code == 200
    parent_id = _queued_generation(client, app, project, locked.json['data'])
    body = text_element('新增正文')
    body['frame'] = {**body['frame'], 'y': 170}
    monkeypatch.setattr('services.scene.generation.chat_json', lambda *_: (
        {'background': {'kind': 'solid', 'color': '#FFFFFF'},
         'elements': [body], 'image_requests': []}, 'text-request', {}))
    assert run_once() and run_once()
    with app.app_context():
        candidate_id = db.session.get(SceneTaskItem, parent_id).result_json['candidate_id']
        candidate = db.session.get(SceneCandidate, candidate_id)
        revision = db.session.get(PageSceneRevision, candidate.proposed_revision_id)
        assert revision.scene_json['elements'][0]['id'] == title['id']
        assert revision.scene_json['elements'][0]['text'] == '人工锁定标题'
        assert revision.scene_json['elements'][0]['locked'] is True
        assert revision.scene_json['elements'][1]['text'] == '新增正文'


def test_ai_image_edit_uses_asset_child_and_keeps_head_until_accept(client, app, monkeypatch, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'ai-edit-assets')
    old_image = BytesIO(); Image.new('RGB', (80, 80), '#884422').save(old_image, format='PNG')
    old_image.seek(0)
    upload = client.post(f'/api/v2/projects/{project["project_id"]}/assets',
        data={'file': (old_image, 'old.png')}, content_type='multipart/form-data', headers=key())
    assert upload.status_code == 201, upload.json
    old_asset_id = upload.json['data']['asset_id']
    image_id = str(uuid4())
    image_element = {'id': image_id, 'kind': 'image', 'role': 'illustration',
        'asset_id': old_asset_id, 'frame': {'x': 600, 'y': 160, 'w': 250, 'h': 230,
        'rotation_deg': 0}, 'crop': {'x': 0, 'y': 0, 'w': 1, 'h': 1},
        'fit': 'contain', 'opacity': 1, 'locked': False, 'alt_text': ''}
    saved = patch(client, project, page, [{'op': 'add_element', 'element': image_element}])
    assert saved.status_code == 200, saved.json
    base = saved.json['data']
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'e' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Edit',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    config = client.patch(f'/api/v2/projects/{project["project_id"]}/model-config',
        json={'base_project_version': 0, 'model_config': {
            'text_model': 'fake-text', 'image_model': 'fake-image'}}, headers=key())
    assert config.status_code == 200
    task = client.post(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}/ai-edits',
        json={'base_revision_id': base['revision_id'], 'base_page_version': base['page_version'],
              'credential_id': credential['credential_id'], 'instruction': '换成蓝色示意图',
              'element_ids': [image_id]}, headers=key())
    assert task.status_code == 202, task.json
    task_id = task.json['data']['task_id']
    new_image = BytesIO(); Image.new('RGB', (90, 90), '#224488').save(new_image, format='PNG')
    calls = {'text': 0, 'image': 0}
    def fake_chat(*_args):
        calls['text'] += 1
        return {'commands': [], 'image_requests': [{
            'element_id': image_id, 'prompt': 'blue illustration'}]}, 'edit-request', {}
    def fake_image(*_args):
        calls['image'] += 1
        return new_image.getvalue(), 'image-request', {}
    monkeypatch.setattr('services.scene.generation.chat_json', fake_chat)
    monkeypatch.setattr('services.scene.generation.generate_image', fake_image)
    assert run_once()
    with app.app_context():
        parent = db.session.get(SceneTaskItem, task_id)
        assert parent.state == 'waiting_assets'
        assert SceneTaskItem.query.filter_by(group_id=parent.group_id, operation='generate_asset').count() == 1
    assert run_once() and run_once()
    with app.app_context():
        parent = db.session.get(SceneTaskItem, task_id)
        candidate = db.session.get(SceneCandidate, parent.result_json['candidate_id'])
        proposed = db.session.get(PageSceneRevision, candidate.proposed_revision_id).scene_json
        assert proposed['elements'][0]['id'] == image_id
        assert proposed['elements'][0]['asset_id'] != old_asset_id
        assert db.session.get(Page, page['page_id']).head_revision_id == base['revision_id']
        assert calls == {'text': 1, 'image': 1}


def test_manual_image_crop_opacity_and_rotation_export(client, app, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'manual-image-assets')
    source = BytesIO()
    Image.new('RGB', (100, 80), '#884422').save(source, format='PNG')
    source.seek(0)
    uploaded = client.post(f'/api/v2/projects/{project["project_id"]}/assets',
        data={'file': (source, 'image.png')}, content_type='multipart/form-data', headers=key())
    assert uploaded.status_code == 201
    asset_id = uploaded.json['data']['asset_id']
    element_id = str(uuid4())
    element = {'id': element_id, 'kind': 'image', 'role': 'illustration',
        'asset_id': asset_id, 'frame': {'x': 50, 'y': 60, 'w': 300, 'h': 240,
        'rotation_deg': 0}, 'crop': {'x': 0, 'y': 0, 'w': 1, 'h': 1},
        'fit': 'contain', 'opacity': 1, 'locked': False, 'alt_text': '配图'}
    added = patch(client, project, page, [{'op': 'add_element', 'element': element}])
    assert added.status_code == 200
    base = added.json['data']
    assert asset_id in base['asset_urls']
    assert client.get(base['asset_urls'][asset_id]).status_code == 200
    missing = patch(client, project, base, [{'op': 'set_opacity', 'element_id': element_id}])
    assert missing.status_code == 422
    invalid = patch(client, project, base, [{'op': 'set_opacity',
        'element_id': element_id, 'opacity': 1.2}])
    assert invalid.status_code == 422
    changed = patch(client, project, base, [
        {'op': 'set_crop', 'element_id': element_id,
         'crop': {'x': .1, 'y': .1, 'w': .8, 'h': .8}, 'fit': 'cover'},
        {'op': 'set_opacity', 'element_id': element_id, 'opacity': .4},
        {'op': 'set_frame', 'element_id': element_id,
         'frame': {**element['frame'], 'rotation_deg': 25}},
    ])
    assert changed.status_code == 200, changed.json
    assert asset_id in changed.json['data']['asset_urls']
    edited = changed.json['data']['scene']['elements'][0]
    assert edited['crop'] == {'x': .1, 'y': .1, 'w': .8, 'h': .8}
    assert edited['opacity'] == .4 and edited['frame']['rotation_deg'] == 25
    snapshot = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': project['project_version'], 'pages': [{
            'page_id': page['page_id'], 'page_version': changed.json['data']['page_version'],
            'revision_id': changed.json['data']['revision_id']}]}, headers=key())
    assert snapshot.status_code == 201
    with app.app_context():
        snap = db.session.get(DeckSnapshot, snapshot.json['data']['snapshot_id'])
        revision = db.session.get(PageSceneRevision, changed.json['data']['revision_id'])
        asset = db.session.get(Asset, asset_id)
        from pptx import Presentation
        payload = render_pptx(snap, {revision.id: revision}, {asset_id: asset})
        verified = verify_pptx(payload, snap, {revision.id: revision}, assets={asset_id: asset})
        assert verified['pages'][0]['image_count'] == 1
        assert verified['opc_relationships_checked'] and verified['native_text_style_checked']
        assert verified['image_asset_bytes_checked']
        presentation = Presentation(BytesIO(payload))
        assert abs(presentation.slides[0].shapes[0].rotation - 25) < .01
        presentation.slides[0].shapes[0].crop_left = .1
        changed = BytesIO(); presentation.save(changed)
        with pytest.raises(ValueError, match='image crop mismatch'):
            verify_pptx(changed.getvalue(), snap, {revision.id: revision}, assets={asset_id: asset})
        presentation = Presentation(BytesIO(payload))
        presentation.slides[0].shapes[0].rotation = 45
        changed = BytesIO(); presentation.save(changed)
        with pytest.raises(ValueError, match='image rotation mismatch'):
            verify_pptx(changed.getvalue(), snap, {revision.id: revision}, assets={asset_id: asset})
        from zipfile import ZIP_DEFLATED, ZipFile
        replacement = BytesIO(); Image.new('RGBA', (80, 64), '#224488').save(replacement, format='PNG')
        changed = BytesIO()
        with ZipFile(BytesIO(payload)) as original, ZipFile(changed, 'w', ZIP_DEFLATED) as altered:
            for name in original.namelist():
                altered.writestr(name, replacement.getvalue() if name.startswith('ppt/media/') else original.read(name))
        with pytest.raises(ValueError, match='image asset bytes mismatch'):
            verify_pptx(changed.getvalue(), snap, {revision.id: revision}, assets={asset_id: asset})
        background_scene = {**revision.scene_json, 'background': {'kind': 'image',
            'asset_id': asset_id, 'crop': {'x': 0, 'y': 0, 'w': 1, 'h': 1}}}
        background_revisions = {revision.id: SimpleNamespace(scene_json=background_scene)}
        # This separate in-memory example has different content; do not label
        # it with the immutable database snapshot's original Scene hash.
        from copy import deepcopy
        background_snap = SimpleNamespace(manifest_json=deepcopy(snap.manifest_json))
        background_snap.manifest_json['pages'][0]['scene_hash'] = digest(background_scene)
        background_payload = render_pptx(background_snap, background_revisions, {asset_id: asset})
        assert verify_pptx(background_payload, background_snap, background_revisions,
            assets={asset_id: asset})['page_count'] == 1


def test_legacy_file_cookie_and_scene_write_isolation(client, app, monkeypatch):
    project, page = create(client, app)
    monkeypatch.setenv('ACCESS_CODE', 'sandbox-visit-code')
    asset_url = f'/files/{project["project_id"]}/template/missing.png'
    assert client.get(asset_url).status_code == 403
    verified = client.post('/api/access-code/verify', json={'code': 'sandbox-visit-code'})
    assert verified.status_code == 200
    assert 'banana_file_access=' in verified.headers['Set-Cookie']
    assert client.get(asset_url).status_code == 409  # No legacy file path for Scene.
    headers = {'X-Access-Code': 'sandbox-visit-code'}
    legacy = client.put(f'/api/projects/{project["project_id"]}/pages/{page["page_id"]}',
                        json={}, headers=headers)
    assert legacy.status_code == 409
    reference = client.get(f'/api/reference-files/project/{project["project_id"]}', headers=headers)
    assert reference.status_code == 409


def test_two_page_generation_keeps_success_when_other_page_fails(client, app, monkeypatch, tmp_path):
    project, first = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'two-page-assets')
    second = client.post(f'/api/v2/projects/{project["project_id"]}/pages',
        json={'base_project_version': 0, 'outline': {'title': '第二页', 'points': []}}, headers=key())
    assert second.status_code == 201
    second = second.json['data']
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'f' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Two pages',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    configured = client.patch(f'/api/v2/projects/{project["project_id"]}/model-config',
        json={'base_project_version': 1, 'model_config': {
            'text_model': 'fake-text', 'image_model': 'fake-image'}}, headers=key())
    assert configured.status_code == 200
    plan = client.post(f'/api/v2/projects/{project["project_id"]}/generation-plans',
        json={'base_project_version': 2}, headers=key()).json['data']
    started = client.post(f'/api/v2/projects/{project["project_id"]}/generation-tasks', json={
        'plan_id': plan['plan_id'], 'credential_id': credential['credential_id'],
        'targets': [{'page_id': p['page_id'], 'base_revision_id': p['revision_id'],
                     'base_page_version': p['page_version']} for p in (first, second)]},
        headers=key())
    assert started.status_code == 202, started.json
    def fake_chat(_url, _key, _model, _system, prompt):
        if json.loads(prompt)['page']['page_id'] == second['page_id']:
            raise SceneError('MODEL_SCENE_INVALID', 'Invalid model response')
        return {'background': {'kind': 'solid', 'color': '#FFFFFF'},
                'elements': [text_element('第一页成功')], 'image_requests': []}, 'request-ok', {}
    monkeypatch.setattr('services.scene.generation.chat_json', fake_chat)
    for _ in range(4):
        if not run_once():
            break
    with app.app_context():
        candidates = SceneCandidate.query.filter_by(project_id=project['project_id']).all()
        assert len(candidates) == 1 and candidates[0].page_id == first['page_id']
        tasks = SceneTaskItem.query.filter_by(group_id=started.json['data']['task_group_id']).all()
        assert {task.page_id: task.state for task in tasks} == {
            first['page_id']: 'succeeded', second['page_id']: 'failed'}
    group = client.get(f'/api/v2/task-groups/{started.json["data"]["task_group_id"]}').json['data']
    assert group['state'] == 'partial_failed' and group['completed_pages'] == 1


def test_cancel_during_paid_asset_call_cannot_publish_candidate(client, app, monkeypatch, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'cancel-during-call')
    parent_id = _queued_generation(client, app, project, page)
    monkeypatch.setattr('services.scene.generation.chat_json', lambda *_: (
        {'background': {'kind': 'solid', 'color': '#FFFFFF'},
         'elements': [text_element('文字')], 'image_requests': [
             {'role': 'background', 'prompt': 'abstract blue'}]}, 'text-request', {}))
    picture = BytesIO(); Image.new('RGB', (64, 64), '#224488').save(picture, format='PNG')
    assert run_once()  # Material child is now queued.
    with app.app_context():
        parent = db.session.get(SceneTaskItem, parent_id)
        child_id = SceneTaskItem.query.filter_by(group_id=parent.group_id,
            operation='generate_asset').one().id
    def cancelled_image(*_args):
        with app.app_context():
            parent = db.session.get(SceneTaskItem, parent_id)
            child = db.session.get(SceneTaskItem, child_id)
            parent.cancel_requested_at = now()
            child.cancel_requested_at = now()
            db.session.commit()
        return picture.getvalue(), 'charged-request', {}
    monkeypatch.setattr('services.scene.generation.generate_image', cancelled_image)
    assert run_once()
    with app.app_context():
        assert db.session.get(SceneTaskItem, child_id).state == 'cancelled'
        assert SceneCandidate.query.filter_by(page_id=page['page_id']).count() == 0
        assert Asset.query.filter_by(project_id=project['project_id'], kind='generated_image').count() == 0


def test_rejected_rate_limit_retries_only_asset_with_backoff(client, app, monkeypatch, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'rate-limit-assets')
    parent_id = _queued_generation(client, app, project, page)
    monkeypatch.setattr('services.scene.generation.chat_json', lambda *_: (
        {'background': {'kind': 'solid', 'color': '#FFFFFF'},
         'elements': [text_element('文字')], 'image_requests': [
             {'role': 'background', 'prompt': 'abstract blue'}]}, 'text-request', {}))
    picture = BytesIO(); Image.new('RGB', (64, 64), '#224488').save(picture, format='PNG')
    calls = {'image': 0}
    def rate_limited(*_args):
        calls['image'] += 1
        if calls['image'] == 1:
            raise SceneError('MODEL_RATE_LIMITED', 'Rejected before billing', 429,
                             {'retry_after_seconds': 12})
        return picture.getvalue(), 'second-request', {}
    monkeypatch.setattr('services.scene.generation.generate_image', rate_limited)
    assert run_once() and run_once()
    with app.app_context():
        parent = db.session.get(SceneTaskItem, parent_id)
        child = SceneTaskItem.query.filter_by(group_id=parent.group_id,
            operation='generate_asset').one()
        assert child.state == 'queued' and child.attempt_count == 1
        assert child.next_run_at is not None
        child_id = child.id
        child.next_run_at = now() - timedelta(seconds=1)
        db.session.commit()
    assert run_once() and run_once()
    with app.app_context():
        assert db.session.get(SceneTaskItem, parent_id).state == 'succeeded'
        assert db.session.get(SceneTaskItem, child_id).attempt_count == 2
        assert calls['image'] == 2


def test_scene_schema_and_hash_golden_fixture():
    root = Path(__file__).resolve().parents[3]
    fixture = json.loads((root / 'docs' / 'fixtures' / 'slide_scene_hash_v1.json').read_text(encoding='utf-8'))
    scene = normalize_scene(fixture['scene'], 960, 540)
    assert digest(scene) == fixture['sha256']
    committed = json.loads((root / 'backend' / 'schemas' / 'slide_scene_v1.schema.json').read_text(encoding='utf-8'))
    assert committed['$schema'].endswith('/draft/2020-12/schema')
    assert {k: v for k, v in committed.items() if k != '$schema'} == SlideScene.model_json_schema()


def test_outline_order_and_soft_delete_keep_snapshot(client, app):
    project, first = create(client, app)
    second = client.post(f'/api/v2/projects/{project["project_id"]}/pages',
        json={'base_project_version': 0, 'outline': {'title': '第二页', 'points': []}}, headers=key())
    assert second.status_code == 201
    updated = client.get(f'/api/v2/projects/{project["project_id"]}').json['data']
    assert updated['project_version'] == 1
    ordered = client.patch(f'/api/v2/projects/{project["project_id"]}/outline', json={
        'base_project_version': 1,
        'pages': [{'page_id': second.json['data']['page_id'], 'outline': {'title': '前页', 'points': ['要点']}},
                  {'page_id': first['page_id'], 'outline': {'title': '后页', 'points': []}}]}, headers=key())
    assert ordered.status_code == 200
    assert [p['outline_content']['title'] for p in ordered.json['data']['pages']] == ['前页', '后页']
    page = ordered.json['data']['pages'][0]
    snap = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': 2, 'pages': [{'page_id': page['page_id'],
                                        'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    assert snap.status_code == 201
    deleted = client.delete(f'/api/v2/projects/{project["project_id"]}/pages/{page["page_id"]}',
        json={'base_project_version': 2, 'base_page_version': page['page_version']}, headers=key())
    assert deleted.status_code == 200
    assert len(deleted.json['data']['pages']) == 1
    with app.app_context():
        assert db.session.get(DeckSnapshot, snap.json['data']['snapshot_id']) is not None


def _removable_reference(client, app, tmp_path):
    project, page = create(client, app)
    app.config['ASSET_STORE_ROOT'] = str(tmp_path / 'reference-removal')
    app.config['CREDENTIAL_ENCRYPTION_KEY'] = base64.urlsafe_b64encode(b'd' * 32).decode()
    app.config['MODEL_GATEWAY_ALLOWLIST'] = 'https://model.example'
    credential = client.post('/api/v2/model-credentials', json={'label': 'Deletion test',
        'base_url': 'https://model.example/v1', 'api_key': 'sandbox-test-key'}, headers=key()).json['data']
    root = f'/api/v2/projects/{project["project_id"]}'
    configured = client.patch(root + '/model-config', json={'base_project_version': project['project_version'],
        'model_config': {'text_model': 'fake-text'}}, headers=key())
    assert configured.status_code == 200
    source = BytesIO()
    Image.new('RGB', (320, 180), '#24415c').save(source, format='PNG')
    source.seek(0)
    uploaded = client.post(root + '/template-documents', data={'file': (source, 'reference.png')},
        content_type='multipart/form-data', headers=key())
    assert uploaded.status_code == 202 and run_once()
    reference = client.get(root + '/template-assets').json['data']['items'][0]
    return root, configured.json['data'], page, reference, credential


def test_template_soft_delete_preserves_scene_history_and_checks_versions(client, app, tmp_path, monkeypatch):
    from models import ProjectTemplateAsset
    root, project, page, reference, credential = _removable_reference(client, app, tmp_path)
    template_id = reference['template_asset_id']
    url = root + '/template-assets/' + template_id
    bound = client.patch(root + f'/pages/{page["page_id"]}/template', json={
        'base_page_version': page['page_version'], 'template_asset_id': template_id,
        'style_text': '保留补充说明'}, headers=key())
    assert bound.status_code == 200
    project = client.get(root).json['data']
    page = project['pages'][0]
    scene_url = root + f'/pages/{page["page_id"]}/scene'
    before = client.get(scene_url).json['data']
    plan = client.post(root + '/generation-plans', json={'base_project_version': project['project_version']}, headers=key())
    assert plan.status_code == 201
    snapshot = client.post(root + '/snapshots', json={'project_version': project['project_version'], 'pages': [{
        'page_id': page['page_id'], 'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    assert snapshot.status_code == 201
    payload = {'base_project_version': project['project_version'], 'base_analysis_revision': 0,
        'expected_pages': [{'page_id': page['page_id'], 'page_version': page['page_version'], 'revision_id': page['revision_id']}]}
    for override, code in [({'base_project_version': -1}, 'PROJECT_VERSION_CONFLICT'),
                           ({'base_analysis_revision': 99}, 'ANALYSIS_VERSION_CONFLICT'),
                           ({'expected_pages': []}, 'PAGE_VERSION_CONFLICT'),
                           ({'expected_pages': [None]}, 'INPUT_INVALID')]:
        rejected = client.delete(url, json={**payload, **override}, headers=key())
        assert rejected.status_code in (409, 422) and rejected.json['error']['code'] == code
    # Owner authorization precedes idempotency and every mutation.
    other_id = str(uuid4())
    with app.app_context():
        db.session.add(Principal(id=other_id, kind='sub2api_user', display_name='Other'))
        db.session.commit()
    with monkeypatch.context() as other:
        other.setattr('controllers.scene_controller.owner_id', lambda: other_id)
        assert client.delete(url, json=payload, headers=key()).status_code == 404
    headers = key()
    removed = client.delete(url, json=payload, headers=headers)
    assert removed.status_code == 200, removed.json
    assert removed.json['data']['project_version'] == project['project_version'] + 1
    assert removed.json['data']['pages'][0]['template_asset_id'] is None
    assert removed.json['data']['pages'][0]['template_style_text'] == '保留补充说明'
    assert removed.json['data']['pages'][0]['page_version'] == page['page_version'] + 1
    after = client.get(scene_url).json['data']
    assert after['scene'] == before['scene'] and after['revision_id'] == before['revision_id']
    assert client.delete(url, json=payload, headers=headers).json['data'] == removed.json['data']
    assert client.delete(url, json={**payload, 'base_analysis_revision': 1}, headers=headers).status_code == 409
    assert client.get(root + '/template-assets').json['data']['items'] == []
    assert client.get(reference['preview_url']).status_code == 200
    assert client.get(reference['thumbnail_url']).status_code == 200
    assert client.patch(url, json={'base_analysis_revision': 0, 'analysis': {}}, headers=key()).status_code == 404
    assert client.patch(root + f'/pages/{page["page_id"]}/template', json={
        'base_page_version': after['page_version'], 'template_asset_id': template_id}, headers=key()).status_code == 404
    analysis = client.post(root + f'/template-documents/{reference["template_document_id"]}/analyze', json={
        'credential_id': credential['credential_id'], 'selected_page_indexes': [1]}, headers=key())
    assert analysis.status_code == 422 and analysis.json['error']['code'] == 'PAGE_SELECTION_INVALID'
    with app.app_context():
        template = db.session.get(ProjectTemplateAsset, template_id)
        assert template.deleted_at is not None
        frozen = db.session.get(GenerationPlan, plan.json['data']['plan_id'])
        assert frozen.manifest_json['pages'][0]['reference']['template_asset_id'] == template_id
        assert frozen.manifest_hash == digest(frozen.manifest_json)
        assert db.session.get(DeckSnapshot, snapshot.json['data']['snapshot_id']) is not None
        assert db.session.get(Asset, template.preview_asset_id).state == 'ready'
    new_plan = client.post(root + '/generation-plans', json={
        'base_project_version': removed.json['data']['project_version']}, headers=key())
    assert new_plan.status_code == 201
    with app.app_context():
        assert db.session.get(GenerationPlan, new_plan.json['data']['plan_id']).manifest_json['pages'][0]['reference'] is None


@pytest.mark.parametrize('task_state', ['queued', 'running', 'outcome_unknown'])
def test_template_removal_waits_for_analysis_and_forbids_retry_afterwards(client, app, tmp_path, task_state):
    from models import ProjectTemplateAsset
    root, project, _, reference, credential = _removable_reference(client, app, tmp_path)
    analysis = client.post(root + f'/template-documents/{reference["template_document_id"]}/analyze', json={
        'credential_id': credential['credential_id'], 'selected_page_indexes': [1]}, headers=key())
    assert analysis.status_code == 202
    task_id = analysis.json['data']['tasks'][0]['task_id']
    with app.app_context():
        task = db.session.get(SceneTaskItem, task_id)
        task.state = task_state
        db.session.commit()
    url = root + '/template-assets/' + reference['template_asset_id']
    payload = {'base_project_version': project['project_version'], 'base_analysis_revision': 0, 'expected_pages': []}
    if task_state in ('queued', 'running'):
        denied = client.delete(url, json=payload, headers=key())
        assert denied.status_code == 409 and denied.json['error']['code'] == 'ANALYSIS_IN_PROGRESS'
        with app.app_context():
            assert db.session.get(ProjectTemplateAsset, reference['template_asset_id']).deleted_at is None
            task = db.session.get(SceneTaskItem, task_id)
            task.state = 'failed'
            task.error_code = 'MODEL_AUTH_FAILED'
            db.session.commit()
    removed = client.delete(url, json=payload, headers=key())
    assert removed.status_code == 200, removed.json
    retried = client.post('/api/v2/tasks/' + task_id + '/retry', json={'acknowledge_possible_charge': True}, headers=key())
    assert retried.status_code == 404
    with app.app_context():
        assert db.session.get(SceneTaskItem, task_id).state in ('failed', 'outcome_unknown')
        document = db.session.get(TemplateDocument, reference['template_document_id'])
        assert document.selected_page_indexes_json == [] and document.status == 'preview_ready'


@pytest.mark.parametrize('delete_during_call', [False, True])
def test_deleted_template_cannot_dispatch_or_publish_analysis(client, app, tmp_path, monkeypatch, delete_during_call):
    from models import ProjectTemplateAsset
    root, _, _, reference, credential = _removable_reference(client, app, tmp_path)
    analysis = client.post(root + f'/template-documents/{reference["template_document_id"]}/analyze', json={
        'credential_id': credential['credential_id'], 'selected_page_indexes': [1]}, headers=key())
    assert analysis.status_code == 202
    calls = []
    def mark_deleted():
        template = db.session.get(ProjectTemplateAsset, reference['template_asset_id'])
        template.deleted_at = now()
        db.session.commit()
    def model_call(*_args, **_kwargs):
        calls.append(True)
        mark_deleted()
        return {'schema_version': 1, 'role': 'content',
            'palette': {'background': '#FFFFFF', 'text': '#000000', 'accent': '#FFCC00'},
            'font_suggestions': [], 'layout_hints': [], 'decorative_hints': [],
            'content_density': 'medium', 'warnings': []}
    if not delete_during_call:
        with app.app_context():
            mark_deleted()
    monkeypatch.setattr('services.templates.style_analysis._model_call', model_call)
    assert run_once()
    result = client.get('/api/v2/tasks/' + analysis.json['data']['tasks'][0]['task_id']).json['data']
    assert result['state'] == 'failed' and result['error_code'] == 'TASK_INPUT_INVALID'
    assert bool(calls) == delete_during_call
    with app.app_context():
        template = db.session.get(ProjectTemplateAsset, reference['template_asset_id'])
        assert template.deleted_at is not None and template.analysis_revision == 0
        assert template.get_analysis() is None
        # A late worker failure must not update the deleted reference/document.
        assert template.analysis_status == 'running'
        assert db.session.get(TemplateDocument, reference['template_document_id']).status == 'analyzing'


def test_template_removal_checks_every_bound_page_before_unbinding(client, app, tmp_path):
    root, project, page, reference, _ = _removable_reference(client, app, tmp_path)
    added = client.post(root + '/pages', json={'base_project_version': project['project_version']}, headers=key())
    assert added.status_code == 201
    for item in (page, added.json['data']):
        assert client.patch(root + f'/pages/{item["page_id"]}/template', json={
            'base_page_version': item['page_version'], 'template_asset_id': reference['template_asset_id']}, headers=key()).status_code == 200
    project = client.get(root).json['data']
    payload = {'base_project_version': project['project_version'], 'base_analysis_revision': 0,
        'expected_pages': [{'page_id': item['page_id'], 'page_version': item['page_version'],
                            'revision_id': item['revision_id']} for item in project['pages']]}
    updated = patch(client, project, project['pages'][1], [{'op': 'add_element', 'element': text_element('并发保存的正文')}])
    assert updated.status_code == 200
    url = root + '/template-assets/' + reference['template_asset_id']
    denied = client.delete(url, json=payload, headers=key())
    assert denied.status_code == 409 and denied.json['error']['code'] == 'PAGE_VERSION_CONFLICT'
    current = client.get(root).json['data']
    assert all(item['template_asset_id'] == reference['template_asset_id'] for item in current['pages'])
    assert current['project_version'] == project['project_version']
    payload['expected_pages'] = [{'page_id': item['page_id'], 'page_version': item['page_version'],
                                 'revision_id': item['revision_id']} for item in current['pages']]
    removed = client.delete(url, json=payload, headers=key())
    assert removed.status_code == 200 and all(item['template_asset_id'] is None for item in removed.json['data']['pages'])
    assert [item['revision_id'] for item in removed.json['data']['pages']] == [item['revision_id'] for item in current['pages']]
    assert [item['page_version'] for item in removed.json['data']['pages']] == [item['page_version'] + 1 for item in current['pages']]

@pytest.mark.parametrize('fmt', ['pptx', 'pdf'])
def test_export_cannot_bypass_failed_text_layout(client, app, monkeypatch, fmt):
    project, page = create(client, app)
    frozen = client.post(f'/api/v2/projects/{project["project_id"]}/snapshots', json={
        'project_version': project['project_version'], 'pages': [{'page_id': page['page_id'],
        'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    assert frozen.status_code == 201
    queued = client.post(f'/api/v2/projects/{project["project_id"]}/exports', json={
        'snapshot_id': frozen.json['data']['snapshot_id'], 'format': fmt}, headers=key())
    assert queued.status_code == 202
    def unavailable(*args):
        raise ValueError('Synthetic layout failure')
    monkeypatch.setattr('workers.scene_runner.measure_text_layout', unavailable)
    assert run_once()
    result = client.get(f'/api/v2/projects/{project["project_id"]}/exports/{queued.json["data"]["export_id"]}').json['data']
    assert result['status'] == 'failed' and result['download_url'] is None
    assert result.get('report_url') is None


@pytest.mark.parametrize('code', ['PPTX_RENDER_CLIPPED', 'PPTX_RENDER_REFLOW',
    'PPTX_RENDER_TEXT_MISMATCH', 'PPTX_RENDER_FONT_MISMATCH', 'PPTX_RENDER_UNVERIFIED'])
def test_target_render_failure_is_terminal_not_a_waivable_visual_warning(client, app, monkeypatch, code):
    from backend.tests.scene_layout_fixtures import single_line_layout
    from services.exports.rendered_text import TargetRenderError
    monkeypatch.setattr('workers.scene_runner.measure_text_layout',
        lambda snapshot, revisions, assets: single_line_layout(snapshot, revisions))
    def reject(*args):
        raise TargetRenderError(code, 'Synthetic target-render failure')
    monkeypatch.setattr('workers.scene_runner.compare_export', reject)
    project, page = create(client, app)
    base = f'/api/v2/projects/{project["project_id"]}'
    snapshot = client.post(base + '/snapshots', json={'project_version': 0, 'pages': [{
        'page_id': page['page_id'], 'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    submitted = client.post(base + '/exports', json={'snapshot_id': snapshot.json['data']['snapshot_id'],
        'format': 'pptx'}, headers=key()).json['data']
    assert run_once()
    task = client.get('/api/v2/tasks/' + submitted['task_id']).json['data']
    assert task['state'] == 'failed' and task['error_code'] == code
    result = client.get(base + '/exports/' + submitted['export_id']).json['data']
    assert result['status'] == 'failed' and result['download_url'] is None
    assert result['report_url'] is None and result['review_file_url'] is None
    review = client.post(base + '/exports/' + submitted['export_id'] + '/review', json={
        'report_sha256': 'a' * 64, 'acknowledged_warning_codes': ['VISUAL_DIFF_UNCALIBRATED']}, headers=key())
    assert review.status_code == 409
    with app.app_context():
        item = db.session.get(SceneTaskItem, submitted['task_id'])
        assert item.attempt_count == 1 and item.next_run_at is None


@pytest.mark.parametrize('rendered_check', [None, {}, {'version': 1, 'status': 'failed'},
    {'version': 1, 'status': 'passed', 'rendered_pdf_sha256': 'a' * 64, 'pages': []},
    {'version': 1, 'status': 'passed', 'rendered_pdf_sha256': 'a' * 64, 'pages': ['invalid']}])
def test_unverified_legacy_pptx_report_cannot_be_accepted(client, app, monkeypatch, rendered_check):
    from backend.tests.scene_layout_fixtures import single_line_layout
    monkeypatch.setattr('workers.scene_runner.measure_text_layout',
        lambda snapshot, revisions, assets: single_line_layout(snapshot, revisions))
    project, page = create(client, app)
    base = f'/api/v2/projects/{project["project_id"]}'
    sheet = BytesIO(); Image.new('RGB', (16, 8), 'white').save(sheet, format='PNG')
    monkeypatch.setattr('workers.scene_runner.compare_export', lambda *_: ({'status': 'needs_review',
        'rendered_text_layout': rendered_check, 'pages': [{'page_id': page['page_id']}]}, [sheet.getvalue()]))
    snapshot = client.post(base + '/snapshots', json={'project_version': 0, 'pages': [{
        'page_id': page['page_id'], 'page_version': page['page_version'], 'revision_id': page['revision_id']}]}, headers=key())
    submitted = client.post(base + '/exports', json={'snapshot_id': snapshot.json['data']['snapshot_id'],
        'format': 'pptx'}, headers=key()).json['data']
    assert run_once()
    result = client.get(base + '/exports/' + submitted['export_id']).json['data']
    assert result['status'] == 'needs_review'
    response = client.post(base + '/exports/' + submitted['export_id'] + '/review', json={
        'report_sha256': result['report_sha256'],
        'acknowledged_warning_codes': ['VISUAL_DIFF_UNCALIBRATED']}, headers=key())
    assert response.status_code == 422
