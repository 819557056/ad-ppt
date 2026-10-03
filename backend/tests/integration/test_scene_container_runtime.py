"""Opt-in Linux test of the *built* Scene renderer with production restrictions.

Requires SCENE_TEST_RENDERER_IMAGE, SCENE_TEST_DOCKER_DAEMON_ID and DOCKER_HOST.
Never connects to a default Docker daemon, pulls images, mounts a Docker socket,
or reuses existing containers/volumes. Run as root on a disposable QA host so
worker-side chown can exercise the real unique-UID transport.
"""
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
from uuid import uuid4

import pytest

pytestmark = [pytest.mark.integration, pytest.mark.docker]
ROOT = Path(__file__).resolve().parents[3]
CAPS = ['CHOWN', 'DAC_OVERRIDE', 'SETUID', 'SETGID', 'KILL']


@pytest.fixture
def renderer_container(tmp_path):
    image = os.environ.get('SCENE_TEST_RENDERER_IMAGE')
    daemon = os.environ.get('SCENE_TEST_DOCKER_DAEMON_ID')
    host = os.environ.get('DOCKER_HOST')
    if not image or not daemon or not host or sys.platform != 'linux':
        pytest.skip('Explicit disposable Docker daemon, built image and Linux are required')
    assert os.geteuid() == 0, 'Real worker chown must be tested; do not silently disable UID isolation'
    binary = os.environ.get('SCENE_TEST_DOCKER_BINARY') or shutil.which('docker')
    assert binary, 'Docker CLI unavailable'

    def docker(*args, timeout=60, check=True):
        result = subprocess.run([binary, *args], capture_output=True, text=True, timeout=timeout)
        if check:
            assert result.returncode == 0, (args, result.stdout, result.stderr)
        return result

    assert docker('info', '--format', '{{.ID}}').stdout.strip() == daemon, 'Wrong Docker daemon'
    image_id = docker('image', 'inspect', image, '--format', '{{.Id}}').stdout.strip()
    assert image_id.startswith('sha256:')
    jobs = tmp_path / 'jobs'
    jobs.mkdir(mode=0o711)
    name = 'scene-test-renderer-' + uuid4().hex
    command = ['run', '--detach', '--pull=never', '--name', name, '--label', 'scene.qa=renderer-contract',
        '--network=none', '--read-only', '--tmpfs', '/tmp:size=512m,mode=1777', '--cap-drop=ALL',
        '--security-opt=no-new-privileges:true', '--pids-limit=128', '--memory=3g',
        '--mount', f'type=bind,src={jobs},dst=/jobs',
        '--env', 'SCENE_RENDER_JOB_ROOT=/jobs', '--env', 'SCENE_RENDER_REQUIRE_UNIQUE_UID=true']
    for capability in CAPS:
        command += ['--cap-add', capability]
    command.append(image_id)
    identifier = docker(*command).stdout.strip()
    assert len(identifier) == 64 and all(c in '0123456789abcdef' for c in identifier)
    try:
        deadline = time.monotonic() + 150
        while True:
            state = json.loads(docker('inspect', identifier).stdout)[0]
            if not state['State']['Running']:
                failure = docker('logs', identifier)
                pytest.fail('Renderer exited: ' + failure.stdout + failure.stderr)
            status = jobs / '.renderer_status.json'
            if status.exists():
                try:
                    ready = json.loads(status.read_text())
                except (OSError, ValueError):
                    ready = {}
                if ready.get('chromium_ok') and ready.get('libreoffice_ok') and ready.get('pptx_pdf_ok'):
                    assert ready['isolation_mode'] == 'unique_uid'
                    break
            assert time.monotonic() < deadline, ('renderer not ready', ready if status.exists() else {}, docker('logs', identifier).stdout)
            time.sleep(.5)
        yield docker, identifier, jobs, state
    finally:
        artifacts = os.environ.get('SCENE_TEST_ARTIFACTS')
        if artifacts:
            output = Path(artifacts)
            output.mkdir(parents=True, exist_ok=True)
            logs = docker('logs', identifier, check=False)
            (output / (name + '.log')).write_text(logs.stdout + logs.stderr)
        # Only remove the exact newly-created ID on the explicitly pinned daemon.
        docker('rm', '--force', identifier)


def test_built_renderer_readonly_offline_uid_isolation_and_real_jobs(renderer_container):
    from flask import Flask
    import fitz
    import pypdfium2 as pdfium
    from services.scene.render_jobs import run_render_job
    from services.scene import telemetry

    docker, identifier, jobs, inspected = renderer_container
    host = inspected['HostConfig']
    assert host['NetworkMode'] == 'none' and host['ReadonlyRootfs']
    assert host['PidsLimit'] == 128 and host['Memory'] == 3 * 1024 ** 3
    normalize = lambda value: value.upper().removeprefix('CAP_')
    assert {normalize(value) for value in host['CapAdd']} == set(CAPS)
    assert [normalize(value) for value in host['CapDrop']] == ['ALL']
    assert 'no-new-privileges:true' in host['SecurityOpt']
    env = {value.split('=', 1)[0] for value in inspected['Config']['Env']}
    assert not env.intersection({'DATABASE_URL', 'CREDENTIAL_ENCRYPTION_KEY', 'ACCESS_CODE', 'OPENAI_API_KEY', 'GOOGLE_API_KEY'})
    assert not any(mount['Destination'] in ('/var/run/docker.sock', '/app/scene_assets') for mount in inspected['Mounts'])
    probe = '''from pathlib import Path
import errno, socket
assert {p.name for p in Path('/sys/class/net').iterdir()} == {'lo'}
try:
    Path('/app/should-not-exist').write_text('forbidden')
except OSError as exc:
    assert exc.errno == errno.EROFS
else:
    raise AssertionError('root filesystem is writable')
assert not Path('/var/run/docker.sock').exists()
print('readonly_offline_ok')
'''
    assert docker('exec', identifier, '/opt/venv/bin/python', '-c', probe).stdout.strip() == 'readonly_offline_ok'
    # Two worker-prepared directories in the same volume, owned by different UIDs.
    for uid in (21001, 21002):
        directory = jobs / ('private-' + str(uid))
        directory.mkdir(mode=0o700)
        secret = directory / 'sentinel'
        secret.write_text('synthetic isolation sentinel')
        secret.chmod(0o600)
        os.chown(secret, uid, uid)
        os.chown(directory, uid, uid)
    access = '''from pathlib import Path
assert Path('/jobs/private-21001/sentinel').read_text() == 'synthetic isolation sentinel'
try:
    Path('/jobs/private-21002/sentinel').read_text()
except PermissionError:
    pass
else:
    raise AssertionError('cross-job read was allowed')
print('cross_uid_denied')
'''
    assert docker('exec', '--user', '21001:21001', identifier, '/opt/venv/bin/python', '-c', access).stdout.strip() == 'cross_uid_denied'
    # These sentinel directories are not jobs. Remove only our two files before
    # testing transport, whose orphan guard correctly refuses unfinished dirs.
    for uid in (21001, 21002):
        directory = jobs / ('private-' + str(uid))
        assert directory.resolve().parent == jobs.resolve()
        (directory / 'sentinel').unlink()
        directory.rmdir()

    app = Flask('scene-container-transport-test')
    app.config.update(SCENE_RENDER_JOB_ROOT=str(jobs), SCENE_RENDER_UNIQUE_UID=True,
        SCENE_RENDER_UID=-1, SCENE_RENDER_GID=-1, SCENE_RENDER_UID_BASE=10000, RENDER_TIMEOUT_SECONDS=150)
    text_id, request_id = str(uuid4()), str(uuid4())
    text = '容器验证 Scene 123'
    html = ('<!doctype html><meta charset="utf-8"><style>'
        '@font-face{font-family:NotoScene;src:url("file:///app/backend/fonts/NotoSansSC-Regular.ttf")}'
        '@page{size:480pt 270pt;margin:0}body{font-family:NotoScene;margin:0}'
        '.slide{width:480pt;height:270pt;position:relative}'
        '[data-scene-id]{position:absolute;left:24pt;top:24pt;width:380pt;height:60pt;font-size:24pt;line-height:1.2}'
        '</style><section class="slide"><div data-scene-id="' + text_id + '"><div data-scene-content><div data-scene-paragraph style="min-height:1lh">' + text + '</div></div></div></section>').encode()
    with app.app_context(), telemetry.scope(request_id=request_id, operation='export'):
        pdf, metadata = run_render_job('pdf', html)
        assert metadata['status'] == 'succeeded'
        with fitz.open(stream=pdf, filetype='pdf') as document:
            assert len(document) == 1 and tuple(document[0].rect)[2:] == (480.0, 270.0)
        with pdfium.PdfDocument(pdf) as document:
            page = document[0]
            try:
                textpage = page.get_textpage()
                try:
                    assert text in textpage.get_text_range()
                finally:
                    textpage.close()
            finally:
                page.close()
        layout, _ = run_render_job('text_layout', html)
        measured = json.loads(layout)
        assert measured['pages'][0][text_id]['lines'][0]['end'] == len(text)
        assert measured['pages'][0][text_id]['resolved_fonts'][0]['family_name'] == 'Noto Sans CJK SC'
        lock = json.loads((ROOT / 'deployment/scene/runtime.lock.json').read_text())
        assert measured['engine'] == 'chromium-' + lock['tools']['browser']['chromium_version']
        previews, _ = run_render_job('html_preview', html, expected_count=1)
        assert len(previews) == 1 and previews[0].startswith(b'\x89PNG\r\n\x1a\n')
        pages, office = run_render_job('pptx', (ROOT / 'renderer/smoke.pptx').read_bytes(), expected_count=1)
        assert len(pages) == 1 and pages[0].startswith(b'\x89PNG\r\n\x1a\n')
        assert office['status'] == 'succeeded'
    assert not any((path / 'ready').exists() or (path / 'running').exists() for path in jobs.iterdir() if path.is_dir())
    logs = docker('logs', identifier).stdout
    records = [json.loads(line) for line in logs.splitlines() if line.startswith('{')]
    events = [row for row in records if row.get('event') == 'scene.renderer_job']
    assert len(events) == 4 and {row['request_id'] for row in events} == {request_id}
    assert all(row['outcome'] == 'succeeded' for row in events)
    assert text not in logs and 'synthetic isolation sentinel' not in logs



@pytest.mark.skipif(not os.environ.get('SCENE_TEST_REFERENCE_PPTX'), reason='Explicit reference template opt-in is required')
def test_design_reference_template_renders_33_pages_without_changing_source(renderer_container):
    from flask import Flask
    from PIL import Image
    from config import Config
    from services.templates.scene_importer import inspect_pptx
    from services.scene.render_jobs import run_render_job
    from services.scene import telemetry

    _, _, jobs, _ = renderer_container
    source = Path(os.environ['SCENE_TEST_REFERENCE_PPTX']).resolve(strict=True)
    payload = source.read_bytes()
    expected = 'bb5e048aa4cf73898f8a603088d2e304beca5f7008b127a5026ebd6f378d8135'
    assert hashlib.sha256(payload).hexdigest() == expected, 'This is not the design reference template'
    app = Flask('scene-reference-container-test')
    app.config.from_object(Config)
    app.config.update(SCENE_RENDER_JOB_ROOT=str(jobs), SCENE_RENDER_UNIQUE_UID=True,
        SCENE_RENDER_UID=-1, SCENE_RENDER_GID=-1, SCENE_RENDER_UID_BASE=10000, RENDER_TIMEOUT_SECONDS=180)
    with app.app_context(), telemetry.scope(operation='import_template'):
        count, warnings = inspect_pptx(payload)
        assert count == 33
        assert set(warnings) >= {'EXTERNAL_RELATIONSHIP_IGNORED', 'EMBEDDED_OBJECT_STATIC_ONLY'}
        pages, metadata = run_render_job('pptx', payload, expected_count=count)
    assert len(pages) == 33 and hashlib.sha256(source.read_bytes()).hexdigest() == expected
    evidence = {'source_sha256': expected, 'count': len(pages), 'import_warnings': warnings,
                'renderer': metadata, 'pages': [], 'visual_review': 'pending', 'focus_pages': [6, 22, 30]}
    for index, rendered in enumerate(pages, 1):
        with Image.open(BytesIO(rendered)) as image:
            width, height = image.size
            image.verify()
        assert width > 0 and height > 0
        evidence['pages'].append({'page': index, 'width': width, 'height': height,
                                 'sha256': hashlib.sha256(rendered).hexdigest()})
    destination = os.environ.get('SCENE_TEST_ARTIFACTS')
    if destination:
        # New unique directory only; never overwrite an earlier visual-review record.
        output = Path(destination) / ('reference-' + uuid4().hex)
        output.mkdir(parents=True, exist_ok=False)
        for index, rendered in enumerate(pages, 1):
            (output / f'page-{index:03d}.png').write_bytes(rendered)
        (output / 'evidence.json').write_text(json.dumps(evidence, indent=2) + '\n')
        print('SCENE_REFERENCE_ARTIFACTS=' + str(output))

def test_complete_text_layout_contract_and_pdf_share_isolated_renderer(renderer_container, monkeypatch):
    """Real browser -> strict backend validation/cache -> native PPTX and PDF."""
    from types import SimpleNamespace
    from flask import Flask
    from pptx import Presentation
    from services.exports import scene_pdf
    from services.exports.scene_pptx import render_pptx, verify_pptx
    from services.scene.validation import blank_scene, digest

    docker, identifier, jobs, _ = renderer_container
    app = Flask('scene-complete-layout-contract')
    app.config.update(SCENE_RENDER_JOB_ROOT=str(jobs), SCENE_RENDER_UNIQUE_UID=True,
        SCENE_RENDER_UID=-1, SCENE_RENDER_GID=-1, SCENE_RENDER_UID_BASE=10000, RENDER_TIMEOUT_SECONDS=150)
    original_html = scene_pdf.scene_html
    # Worker and renderer normally share /app. This test's worker runs on the
    # QA host; only relocate its public, hash-checked font URI in the HTML.
    monkeypatch.setattr(scene_pdf, 'scene_html', lambda *a: original_html(*a).replace(
        scene_pdf.FONT.as_uri(), 'file:///app/backend/fonts/NotoSansSC-Regular.ttf'))
    scene = blank_scene()
    samples = ['中文 English 123\n第二行继续文本', '', '   ', '\n', 'A\n\nB\n', 'Á🄀B 中文混排 Á🄀B 中文']
    for i, sample in enumerate(samples):
        scene['elements'].append({'id': str(uuid4()), 'kind': 'text', 'role': 'body', 'text': sample,
            'frame': {'x': 20 + i * 150, 'y': 20, 'w': 130, 'h': 480, 'rotation_deg': 0},
            'style': {'font_size_pt': 20, 'font_weight': 700 if i % 2 else 400, 'line_height': 1.2,
                      'padding_pt': 4, 'align': 'left', 'vertical_align': ['top', 'middle', 'bottom'][i % 3], 'color': '#17324D'}})
    snapshot = SimpleNamespace(owner_id=str(uuid4()), project_id=str(uuid4()), manifest_json={
        'canvas': scene['canvas'], 'pages': [{'page_id': str(uuid4()), 'revision_id': 'revision', 'scene_hash': digest(scene)}]})
    revisions = {'revision': SimpleNamespace(scene_json=scene)}
    with app.app_context():
        layout = scene_pdf.measure_text_layout(snapshot, revisions, {})
        cached = scene_pdf.measure_text_layout(snapshot, revisions, {})
        assert cached == layout and cached is not layout
        payload = render_pptx(snapshot, revisions, {}, layout)
        assert verify_pptx(payload, snapshot, revisions, layout)['native_text_count'] == len(samples)
        deck = Presentation(BytesIO(payload))
        assert [s.text.replace('\v', '') for s in deck.slides[0].shapes] == samples
        pdf = scene_pdf.render_pdf(snapshot, revisions, {}, layout)
        scene_pdf.verify_pdf(pdf, snapshot, revisions)
        from services.scene.render_jobs import run_render_job
        from services.exports.rendered_text import verify_rendered_pptx
        target_pdf, target_engine = run_render_job('pptx_pdf', payload)
        rendered_check = verify_rendered_pptx(target_pdf, snapshot, revisions, layout)
        assert rendered_check['status'] == 'passed' and target_engine['engine_version'].startswith('LibreOffice ')
    events = [json.loads(line) for line in docker('logs', identifier).stdout.splitlines() if line.startswith('{')]
    assert [event['kind'] for event in events if event.get('event') == 'scene.renderer_job'] == ['text_layout', 'pdf', 'pptx_pdf']
    assert all(font['is_custom_font'] for value in layout['pages'][0].values() for font in value['resolved_fonts'])
    artifacts = os.environ.get('SCENE_TEST_ARTIFACTS')
    if artifacts:
        output = Path(artifacts)
        output.mkdir(parents=True, exist_ok=True)
        (output / 'text-layout.json').write_text(json.dumps(layout, ensure_ascii=False, indent=2))
        (output / 'text-layout.pptx').write_bytes(payload)
        (output / 'text-layout.pdf').write_bytes(pdf)
        (output / 'pptx-target.pdf').write_bytes(target_pdf)
        (output / 'pptx-rendered-check.json').write_text(json.dumps(rendered_check, indent=2))
