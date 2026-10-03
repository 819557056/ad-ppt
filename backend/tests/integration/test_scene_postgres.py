"""Opt-in real PostgreSQL contract tests; never use the supplied DB for fixtures.

SCENE_TEST_POSTGRES_URL must name a maintenance DB on a disposable test server.
The account needs CREATEDB and permission to read pg_control_system().
Each test creates/drops only its own random database.
No paid model calls, truncation of existing databases, or SQLite substitutions.
"""
import base64
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
import hashlib
import os
from pathlib import Path
import subprocess
import sys
import threading
import time
from types import SimpleNamespace
from uuid import uuid4

from alembic import command
from alembic.config import Config as AlembicConfig
from flask import Flask, request as flask_request
import pytest
from sqlalchemy import create_engine, event, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session
from sqlalchemy.pool import NullPool

from config import Config
from controllers.scene_controller import scene_bp
from models import db, Page, Project
from models.scene_v1 import (ApiIdempotencyRecord, Asset, DeckSnapshot, PageSceneRevision,
    SceneMaintenanceRun, SceneMetric, SceneTaskAttempt, SceneTaskItem, now)
from services.scene import telemetry
from services.scene.validation import digest
from services.scene.versioning import LOCAL_OWNER_ID

pytestmark = pytest.mark.integration
BACKEND = Path(__file__).resolve().parents[2]


def key(**extra):
    return {'Idempotency-Key': str(uuid4()), **extra}


@pytest.fixture
def pg(tmp_path, monkeypatch, request):
    configured = os.environ.get('SCENE_TEST_POSTGRES_URL')
    if not configured:
        pytest.skip('Set SCENE_TEST_POSTGRES_URL for an explicitly disposable PostgreSQL test server')
    parsed = make_url(configured)
    if parsed.get_backend_name() != 'postgresql':
        pytest.fail('SCENE_TEST_POSTGRES_URL must use PostgreSQL, not a compatible substitute')
    name = 'scene_test_' + uuid4().hex
    admin = create_engine(parsed, isolation_level='AUTOCOMMIT', poolclass=NullPool,
                          connect_args={'connect_timeout': 5})
    created = False
    app = None
    try:
        with admin.connect() as connection:
            connection.exec_driver_sql('CREATE DATABASE ' + name + ' TEMPLATE template0')
            created = True
        url = parsed.set(database=name).render_as_string(hide_password=False)
        cfg = AlembicConfig(str(BACKEND / 'alembic.ini'))
        cfg.set_main_option('script_location', str(BACKEND / 'migrations'))
        cfg.set_main_option('sqlalchemy.url', url.replace('%', '%%'))
        monkeypatch.setenv('BANANA_SKIP_AUTO_MIGRATE', '1')
        legacy = None
        if getattr(request, 'param', None) == 'legacy_template':
            command.upgrade(cfg, '018_add_project_title')
            from PIL import Image
            source = tmp_path / 'legacy-template.png'
            Image.new('RGB', (4, 4), '#123456').save(source)
            legacy = {'project_id': str(uuid4()), 'page_id': str(uuid4()), 'source': source,
                      'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest()}
            temporary_engine = create_engine(url, poolclass=NullPool)
            try:
                with temporary_engine.begin() as connection:
                    connection.execute(text('INSERT INTO projects (id, creation_type, status, created_at, updated_at, '
                        'template_image_path, template_style) VALUES (:id, :kind, :status, now(), now(), :path, :style)'),
                        {'id': legacy['project_id'], 'kind': 'idea', 'status': 'DRAFT',
                         'path': str(source), 'style': 'Legacy style'})
                    connection.execute(text('INSERT INTO pages (id, project_id, order_index, status, created_at, updated_at) '
                        'VALUES (:id, :project, 0, :status, now(), now())'),
                        {'id': legacy['page_id'], 'project': legacy['project_id'], 'status': 'DRAFT'})
            finally:
                temporary_engine.dispose()
        if getattr(request, 'param', None) == 'legacy_trace':
            command.upgrade(cfg, '02f8479a3c61')
            legacy = {name: str(uuid4()) for name in ('project_id', 'task_id', 'attempt_id')}
            temporary_engine = create_engine(url, poolclass=NullPool)
            try:
                with temporary_engine.begin() as connection:
                    connection.execute(text("INSERT INTO projects (id, creation_type, status, created_at, updated_at) "
                        "VALUES (:project_id, 'idea', 'DRAFT', now(), now())"), legacy)
                    connection.execute(text("INSERT INTO scene_task_items (id, owner_id, project_id, operation, resource_id, input_json, "
                        "input_hash, state, attempt_count, fence_token, result_json, created_at, updated_at) "
                        "VALUES (:task_id, '00000000-0000-4000-8000-000000000001', :project_id, 'export', :task_id, '{}', "
                        ":hash, 'failed', 1, 1, '{}', now(), now())"), dict(legacy, hash='a' * 64))
                    connection.execute(text("INSERT INTO scene_task_attempts (id, task_item_id, attempt_no, fence_token, worker_id, "
                        "started_at, dispatch_state, usage_json) VALUES (:attempt_id, :task_id, 1, 1, 'old-worker', now(), 'resolved', '{}')"), legacy)
            finally:
                temporary_engine.dispose()
        command.upgrade(cfg, 'head')
        app = Flask('scene-postgres-test')
        app.config.from_object(Config)
        app.config.update(TESTING=True, SCENE_EDITOR_ENABLED=True, SECRET_KEY='postgres-test-only',
            SQLALCHEMY_DATABASE_URI=url,
            SQLALCHEMY_ENGINE_OPTIONS={'poolclass': NullPool, 'connect_args': {
                'connect_timeout': 5, 'options': '-c lock_timeout=6000 -c statement_timeout=15000'}},
            ASSET_STORE_ROOT=str(tmp_path / 'assets'), SCENE_MIN_FREE_BYTES=0,
            CREDENTIAL_ENCRYPTION_KEY=base64.urlsafe_b64encode(b'q' * 32).decode(),
            MODEL_GATEWAY_ALLOWLIST='https://model.example',
            SCENE_CHROMIUM_PATH='', SCENE_RENDER_JOB_ROOT='', OWNER_MODEL_CONCURRENCY=1,
            SCENE_QUEUE_CAPACITY_UNITS=1024, OWNER_QUEUE_CAPACITY_UNITS=256)
        Path(app.config['ASSET_STORE_ROOT']).mkdir()
        db.init_app(app)
        telemetry.install(app)
        pids = {}
        @app.before_request
        def identify_test_connection():
            label = flask_request.headers.get('X-Test-Request')
            if label:
                pids[label] = db.session.execute(text('SELECT pg_backend_pid()')).scalar_one()
        app.register_blueprint(scene_bp)
        with app.test_client() as client:
            response = client.post('/api/v2/projects', json={'title': 'Postgres contract',
                'model_config': {'text_model': 'fake-text'}, 'pages': [{'title': 'One'}, {'title': 'Two'}]}, headers=key())
            assert response.status_code == 201, response.json
            project = response.json['data']
            credential = client.post('/api/v2/model-credentials', json={'label': 'Fake key',
                'base_url': 'https://model.example/v1', 'api_key': 'test-only-no-network'}, headers=key())
            assert credential.status_code == 201, credential.json
        engine = create_engine(url, poolclass=NullPool)
        try:
            yield SimpleNamespace(app=app, engine=engine, url=url, project=project,
                credential=credential.json['data']['credential_id'], pids=pids,
                root=Path(app.config['ASSET_STORE_ROOT']), legacy=legacy)
        finally:
            engine.dispose()
    finally:
        if app is not None and 'sqlalchemy' in app.extensions:
            with app.app_context():
                db.session.remove()
                db.engine.dispose()
        if created:
            # The name is generated above, never copied from the supplied URL.
            with admin.connect() as connection:
                connection.exec_driver_sql('DROP DATABASE ' + name + ' WITH (FORCE)')
        admin.dispose()


def submit(pg, method, path, payload, headers=None):
    with pg.app.test_client() as client:
        response = client.open(path, method=method, json=payload, headers=headers or key())
        return response.status_code, response.json


def wait_for_lock(pg, label):
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        pid = pg.pids.get(label)
        if pid:
            with pg.engine.connect() as connection:
                state = connection.execute(text('SELECT wait_event_type, wait_event FROM pg_stat_activity WHERE pid=:pid'),
                    {'pid': pid}).one_or_none()
            if state and state[0] == 'Lock':
                return state[1]
        time.sleep(.02)
    pytest.fail('The second PostgreSQL backend did not enter the expected lock wait')


def outline(pg, project=None):
    project = project or pg.project
    return {'base_project_version': project['project_version'], 'credential_id': pg.credential,
            'brief': 'Offline concurrency fixture', 'slide_count': 2}


def test_real_migration_and_jsonb_schema(pg):
    with pg.engine.connect() as connection:
        assert connection.execute(text('SHOW server_version_num')).scalar_one().startswith('16')
        assert connection.execute(text('SELECT version_num FROM alembic_version')).scalar_one() == '13e79b401ca8'
        assert connection.execute(text("SELECT data_type FROM information_schema.columns "
            "WHERE table_name='scene_task_items' AND column_name='result_json'")).scalar_one() == 'jsonb'
        assert connection.execute(text('SELECT count(*) FROM page_scene_revisions')).scalar_one() == 2


def test_two_saves_wait_then_refresh_head_and_reject_stale_input(pg, monkeypatch):
    import controllers.scene_controller as controller
    page = pg.project['pages'][0]
    path = f"/api/v2/projects/{pg.project['project_id']}/pages/{page['page_id']}/scene"
    payload = {'base_page_version': page['page_version'], 'base_revision_id': page['revision_id'],
               'commands': [{'op': 'set_background', 'background': {'kind': 'solid', 'color': '#112233'}}]}
    entered, release = threading.Event(), threading.Event()
    original = controller.save_commands
    def hold_first(*args, **kwargs):
        if flask_request.headers.get('X-Test-Request') == 'first':
            entered.set()
            assert release.wait(10)
        return original(*args, **kwargs)
    monkeypatch.setattr(controller, 'save_commands', hold_first)
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(submit, pg, 'PATCH', path, payload, key(**{'X-Test-Request': 'first'}))
        try:
            assert entered.wait(5)
            second = pool.submit(submit, pg, 'PATCH', path, payload, key(**{'X-Test-Request': 'second'}))
            wait_for_lock(pg, 'second')
        finally:
            release.set()
        results = [first.result(timeout=10), second.result(timeout=10)]
    assert [code for code, _ in results] == [200, 409], results
    assert results[1][1]['error']['code'] == 'SCENE_VERSION_CONFLICT'
    with Session(pg.engine) as session:
        current = session.get(Page, page['page_id'])
        assert current.row_version == page['page_version'] + 1
        assert session.query(PageSceneRevision).filter_by(page_id=page['page_id']).count() == 2


@pytest.mark.parametrize('changed_input', [False, True])
def test_same_idempotency_key_serializes_to_one_persisted_task(pg, monkeypatch, changed_input):
    import controllers.scene_controller as controller
    entered, release = threading.Event(), threading.Event()
    original = controller.admit_tasks
    def hold_first(*args, **kwargs):
        if flask_request.headers.get('X-Test-Request') == 'first':
            entered.set()
            assert release.wait(10)
        return original(*args, **kwargs)
    monkeypatch.setattr(controller, 'admit_tasks', hold_first)
    shared = key()
    path = '/api/v2/projects/' + pg.project['project_id'] + '/outline-tasks'
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(submit, pg, 'POST', path, outline(pg), {**shared, 'X-Test-Request': 'first'})
        try:
            assert entered.wait(5)
            other = {**outline(pg), 'brief': 'Different input'} if changed_input else outline(pg)
            second = pool.submit(submit, pg, 'POST', path, other, {**shared, 'X-Test-Request': 'second'})
            assert wait_for_lock(pg, 'second') == 'advisory'
        finally:
            release.set()
        results = [first.result(timeout=10), second.result(timeout=10)]
    assert [code for code, _ in results] == [202, 409 if changed_input else 202], results
    if changed_input:
        assert results[1][1]['error']['code'] == 'IDEMPOTENCY_CONFLICT'
    else:
        assert results[0][1]['data'] == results[1][1]['data']
    with Session(pg.engine) as session:
        assert session.query(SceneTaskItem).count() == 1
        assert session.query(ApiIdempotencyRecord).filter_by(operation='outline:' + pg.project['project_id']).count() == 1


def test_cross_project_advisory_admission_cannot_overfill_owner(pg, monkeypatch):
    import services.scene.queue_capacity as capacity
    code, response = submit(pg, 'POST', '/api/v2/projects', {'title': 'Other project',
        'model_config': {'text_model': 'fake-text'}})
    assert code == 201, response
    projects = [pg.project, response['data']]
    pg.app.config['OWNER_QUEUE_CAPACITY_UNITS'] = 1
    barrier = threading.Barrier(2)
    original = capacity._lock_admissions
    def simultaneous():
        barrier.wait(timeout=5)
        original()
    monkeypatch.setattr(capacity, '_lock_admissions', simultaneous)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(submit, pg, 'POST', '/api/v2/projects/' + project['project_id'] + '/outline-tasks',
            outline(pg, project)) for project in projects]
        results = [future.result(timeout=10) for future in futures]
    assert sorted(code for code, _ in results) == [202, 429], results
    with Session(pg.engine) as session:
        assert session.query(SceneTaskItem).count() == 1


@pytest.mark.parametrize('pg', ['legacy_template'], indirect=True)
def test_existing_legacy_template_backfill_uses_boolean_and_preserves_source(pg):
    from models import ProjectTemplateAsset
    with Session(pg.engine) as session:
        legacy = session.get(Project, pg.legacy['project_id'])
        page = session.get(Page, pg.legacy['page_id'])
        reference = session.get(ProjectTemplateAsset, page.template_asset_id)
        assert legacy.editor_mode == 'legacy_image' and legacy.owner_id == LOCAL_OWNER_ID
        assert reference.user_edited_analysis is False
        assert reference.image_path == str(pg.legacy['source'])
        assert page.template_style_text == 'Legacy style'
    assert hashlib.sha256(pg.legacy['source'].read_bytes()).hexdigest() == pg.legacy['source_sha256']


def seed_tasks(pg, count=1, **overrides):
    ids = []
    with Session(pg.engine) as session:
        for index in range(count):
            stamp = now() - timedelta(days=1) + timedelta(microseconds=index)
            values = dict(id=str(uuid4()), owner_id=LOCAL_OWNER_ID, project_id=pg.project['project_id'],
                resource_id=pg.project['project_id'], operation='generate_outline', input_json={},
                input_hash=digest({}), state='queued', created_at=stamp, updated_at=stamp)
            values.update(overrides)
            row = SceneTaskItem(**values)
            session.add(row); ids.append(row.id)
        session.commit()
    return ids


def test_task_project_no_key_update_does_not_block_worker_foreign_key_insert(pg, monkeypatch):
    import controllers.scene_controller as controller
    original = controller.admit_tasks
    entered, release = threading.Event(), threading.Event()
    def hold_request(*args, **kwargs):
        entered.set()
        assert release.wait(10)
        return original(*args, **kwargs)
    monkeypatch.setattr(controller, 'admit_tasks', hold_request)
    with ThreadPoolExecutor(max_workers=2) as pool:
        request_job = pool.submit(submit, pg, 'POST', '/api/v2/projects/' + pg.project['project_id'] + '/outline-tasks', outline(pg))
        try:
            assert entered.wait(5)
            # Must COMMIT while API still holds the project lock, not after it releases.
            worker_job = pool.submit(seed_tasks, pg, state='succeeded')
            inserted = worker_job.result(timeout=3)
            assert len(inserted) == 1
        finally:
            release.set()
        assert request_job.result(timeout=10)[0] == 202


def test_skip_locked_worker_claim_reaches_unlocked_task(pg):
    from workers.scene_runner import claim
    ids = seed_tasks(pg, 51)
    with Session(pg.engine) as blocker:
        held = blocker.execute(select(SceneTaskItem.id).order_by(SceneTaskItem.created_at, SceneTaskItem.id)
            .limit(50).with_for_update()).scalars().all()
        with pg.app.app_context():
            claimed = claim()
        assert claimed and claimed[0] == ids[-1] and claimed[0] not in held
    with Session(pg.engine) as session:
        assert session.query(SceneTaskItem).filter_by(state='running').count() == 1
        assert session.query(SceneTaskAttempt).count() == 1


def test_real_owner_row_lock_limits_concurrent_workers(pg):
    from workers.scene_runner import claim
    seed_tasks(pg, 100)
    barrier = threading.Barrier(2)
    seen, guard = set(), threading.Lock()
    with pg.app.app_context():
        engine = db.engine
    def before(_connection, _cursor, statement, _parameters, _context, _many):
        if 'FROM principals' not in statement or 'FOR UPDATE' not in statement:
            return
        tid = threading.get_ident()
        with guard:
            first = tid not in seen
            seen.add(tid)
        if first:
            barrier.wait(timeout=5)
    def worker():
        with pg.app.app_context():
            return claim()
    event.listen(engine, 'before_cursor_execute', before)
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(worker) for _ in range(2)]
            results = [future.result(timeout=10) for future in futures]
    finally:
        event.remove(engine, 'before_cursor_execute', before)
    assert len(seen) == 2 and sum(result is not None for result in results) == 1
    with Session(pg.engine) as session:
        assert session.query(SceneTaskItem).filter_by(state='running').count() == 1
        assert session.query(SceneTaskAttempt).count() == 1


def test_expired_acknowledged_call_is_unknown_and_old_fence_cannot_publish(pg):
    from services.scene.generation import _fence
    from services.scene.errors import SceneError
    from workers.scene_runner import claim, renew_lease
    identifier = seed_tasks(pg, state='running', dispatch_state='acknowledged',
        lease_owner='crashed-worker', lease_expires_at=now() - timedelta(seconds=5),
        fence_token=7, attempt_count=1)[0]
    with Session(pg.engine) as session:
        session.add(SceneTaskAttempt(task_item_id=identifier, attempt_no=1, fence_token=7,
            worker_id='crashed-worker', dispatch_state='acknowledged', started_at=now() - timedelta(minutes=1)))
        session.commit()
    with pg.app.app_context():
        assert claim() is None
        assert renew_lease(identifier, 7, 'crashed-worker', 60) is False
        with pytest.raises(SceneError) as caught:
            _fence(identifier, 7, 'crashed-worker')
        assert caught.value.code == 'TASK_FENCE_LOST'
        db.session.rollback()
    with Session(pg.engine) as session:
        item = session.get(SceneTaskItem, identifier)
        attempt = session.query(SceneTaskAttempt).one()
        assert item.state == 'outcome_unknown' and item.attempt_count == 1
        assert attempt.error_code == 'MODEL_OUTCOME_UNKNOWN' and attempt.finished_at is not None


def test_snapshot_locks_all_heads_until_manifest_is_frozen(pg, monkeypatch):
    import services.exports.quality_gate as quality
    original = quality.validate_text_fit
    entered, release = threading.Event(), threading.Event()
    def hold_snapshot(*args, **kwargs):
        if not entered.is_set():
            entered.set()
            assert release.wait(10)
        return original(*args, **kwargs)
    monkeypatch.setattr(quality, 'validate_text_fit', hold_snapshot)
    project, page = pg.project, pg.project['pages'][0]
    payload = {'project_version': project['project_version'], 'pages': [
        {name: item[name] for name in ('page_id', 'page_version', 'revision_id')} for item in project['pages']]}
    edit = {'base_page_version': page['page_version'], 'base_revision_id': page['revision_id'],
            'commands': [{'op': 'set_background', 'background': {'kind': 'solid', 'color': '#123456'}}]}
    with ThreadPoolExecutor(max_workers=2) as pool:
        frozen = pool.submit(submit, pg, 'POST', '/api/v2/projects/' + project['project_id'] + '/snapshots', payload)
        try:
            assert entered.wait(5)
            changed = pool.submit(submit, pg, 'PATCH', f"/api/v2/projects/{project['project_id']}/pages/{page['page_id']}/scene",
                edit, key(**{'X-Test-Request': 'editing'}))
            wait_for_lock(pg, 'editing')
        finally:
            release.set()
        snapshot_code, snapshot_response = frozen.result(timeout=10)
        edit_code, edit_response = changed.result(timeout=10)
    assert snapshot_code == 201 and edit_code == 200, (snapshot_response, edit_response)
    with Session(pg.engine) as session:
        snapshot = session.get(DeckSnapshot, snapshot_response['data']['snapshot_id'])
        assert [item['revision_id'] for item in snapshot.manifest_json['pages']] == [item['revision_id'] for item in project['pages']]
        assert session.get(Page, page['page_id']).head_revision_id != page['revision_id']


def test_postgres_retention_refuses_clients_and_recovers_after_db_commit(pg, monkeypatch):
    import ops.scene_retention as retention
    from services.scene.store import put_bytes, path_for
    with pg.app.app_context():
        asset = put_bytes(LOCAL_OWNER_ID, pg.project['project_id'], 'generated_image', b'fixture', 'application/octet-stream', 'bin')
        asset.created_at = now() - timedelta(days=400)
        path = path_for(asset)
        stamp = (now() - timedelta(days=400)).timestamp(); os.utime(path, (stamp, stamp))
        db.session.commit(); asset_id = asset.id
    with pg.engine.connect() as other_client:
        other_client.execute(text('SELECT 1'))
        with pytest.raises(ValueError, match='Other database clients'):
            retention.maintain(pg.url, pg.root, offline_confirmed=True)
    plan = retention.maintain(pg.url, pg.root, offline_confirmed=True)
    assert plan['records']['assets'] == [asset_id]
    def interrupt(*_args):
        raise OSError('simulated interruption after PostgreSQL commit')
    with monkeypatch.context() as patch:
        patch.setattr(retention, '_remove_file', interrupt)
        with pytest.raises(OSError):
            retention.maintain(pg.url, pg.root, 'apply', plan=plan, confirm_sha256=plan['plan_sha256'], offline_confirmed=True)
    with Session(pg.engine) as session:
        run = session.query(SceneMaintenanceRun).one(); run_id = run.id
        assert session.get(Asset, asset_id) is None and run.status == 'db_pruned'
    assert path.exists()
    result = retention.maintain(pg.url, pg.root, 'resume', run_id=run_id, offline_confirmed=True)
    assert result['status'] == 'completed' and not path.exists()


def test_jsonb_draft_publication_keeps_reserved_queue_units_constant(pg):
    parent = seed_tasks(pg, operation='generate_scene')[0]
    with pg.app.test_client() as client:
        before = client.get('/api/v2/queue-capacity').json['data']['owner']
    assert before['active_items'] == 1 and before['reserved_units'] == 7
    with Session(pg.engine) as session:
        item = session.get(SceneTaskItem, parent)
        item.state = 'waiting_assets'
        item.result_json = {'draft_asset_id': str(uuid4()), 'asset_count': 6}
        session.commit()
    seed_tasks(pg, 6, operation='generate_asset')
    with pg.app.test_client() as client:
        after = client.get('/api/v2/queue-capacity').json['data']['owner']
    assert after['active_items'] == after['reserved_units'] == 7



def test_http_worker_attempt_trace_and_aggregates_use_real_postgres(pg, monkeypatch):
    from workers.scene_runner import run_once
    draft = {'pages': [{'role': 'cover', 'title': 'Private title', 'points': [], 'facts_needed': []},
                       {'role': 'closing', 'title': 'Closing', 'points': [], 'facts_needed': []}]}
    monkeypatch.setattr('services.scene.generation.chat_json', lambda *_args, **_kw: (draft, 'PRIVATE-UPSTREAM', {}))
    with pg.app.test_client() as client:
        queued = client.post('/api/v2/projects/' + pg.project['project_id'] + '/outline-tasks',
            json=outline(pg), headers={**key(), 'X-Request-ID': 'caller-untrusted'})
    assert queued.status_code == 202, queued.json
    request_id = queued.json['request_id']
    assert queued.headers['X-Request-ID'] == request_id and request_id != 'caller-untrusted'
    with pg.app.app_context():
        assert run_once()
        assert telemetry.TRACE.get() is None
    with Session(pg.engine) as session:
        item = session.query(SceneTaskItem).one()
        attempt = session.query(SceneTaskAttempt).one()
        assert item.request_id == attempt.request_id == request_id
        assert item.state == 'succeeded' and item.result_json['draft_asset_id']
        assert attempt.queue_wait_seconds >= 0 and attempt.finished_at is not None
        assert session.query(SceneMetric).filter_by(metric='scene_task_transitions_total',
            operation='generate_outline', outcome='succeeded').one().value == 1
        assert session.query(SceneMetric).filter_by(metric='scene_queue_wait_seconds_count',
            operation='generate_outline').one().value == 1
        assert session.query(SceneMetric).filter_by(metric='scene_stage_duration_seconds_count',
            stage='model_text', operation='generate_outline', outcome='success').one().value == 1


def test_postgres_metrics_atomic_upsert_across_real_processes(pg):
    code = """from sqlalchemy import create_engine
from services.scene import telemetry as t
import sys
engine=create_engine(sys.argv[1])
for _ in range(20):
    with t.scope() as trace:
        t.add_metric('scene_retries_total',stage='retry_automatic')
        t.flush_metrics(engine,trace)
engine.dispose()
"""
    processes = [subprocess.Popen([sys.executable, '-c', code, pg.url], cwd=BACKEND,
        stdout=subprocess.PIPE, stderr=subprocess.PIPE) for _ in range(2)]
    try:
        for process in processes:
            out, err = process.communicate(timeout=40)
            assert process.returncode == 0, (out, err)
        with Session(pg.engine) as session:
            assert session.query(SceneMetric).filter_by(metric='scene_retries_total').one().value == 40
    finally:
        for process in processes:
            if process.poll() is None:
                process.kill()
                process.communicate(timeout=5)


def test_postgres_busy_metrics_drop_only_telemetry_batch(pg):
    # A scrape/other worker holding the aggregate must not indefinitely stall
    # the request, nor roll back the already-committed business transaction.
    with telemetry.scope() as trace:
        telemetry.add_metric('scene_retries_total', stage='retry_automatic')
        telemetry.flush_metrics(pg.engine, trace)
    with Session(pg.engine) as blocker:
        blocker.query(SceneMetric).filter_by(metric='scene_retries_total').with_for_update().one()
        with telemetry.scope() as trace:
            telemetry.add_metric('scene_retries_total', stage='retry_automatic')
            started = time.monotonic()
            telemetry.flush_metrics(pg.engine, trace)
            assert time.monotonic() - started < 3
            assert not trace.metrics
        # A later business commit on a separate backend is still possible.
        identifiers = seed_tasks(pg, state='succeeded')
    with Session(pg.engine) as session:
        assert session.get(SceneTaskItem, identifiers[0]).state == 'succeeded'
        assert session.query(SceneMetric).filter_by(metric='scene_retries_total').one().value == 1
    # SET LOCAL timeouts do not leak to another transaction.
    with pg.engine.connect() as connection:
        assert connection.execute(text('SHOW lock_timeout')).scalar_one() == '0'


@pytest.mark.parametrize('pg', ['legacy_trace'], indirect=True)
def test_postgres_trace_migration_preserves_old_tasks_without_inventing_queue_history(pg):
    with Session(pg.engine) as session:
        item = session.get(SceneTaskItem, pg.legacy['task_id'])
        attempt = session.get(SceneTaskAttempt, pg.legacy['attempt_id'])
        assert item.request_id == attempt.request_id == item.id
        assert item.retry_request_id is None and item.queued_at is None
        assert attempt.queue_wait_seconds is None
        assert item.state == 'failed' and item.attempt_count == attempt.attempt_no == 1


@pytest.mark.parametrize('source_type,page_count', [('image', 1), ('pptx', 3)])
def test_template_import_flushes_preview_and_thumbnail_before_reference(pg, monkeypatch, source_type, page_count):
    """Real immediate FKs catch mapper ordering that SQLite create_all can miss."""
    from io import BytesIO
    from PIL import Image
    from pptx import Presentation
    from models import ProjectTemplateAsset
    from models.scene_v1 import TemplateDocument
    from services.scene.store import checked_bytes, put_bytes
    from services.templates.scene_importer import import_document

    image = BytesIO()
    Image.new('RGB', (108, 61), '#123456').save(image, format='PNG')
    payload = image.getvalue()
    mime = 'image/png'
    if source_type == 'pptx':
        deck = Presentation()
        for _ in range(page_count):
            deck.slides.add_slide(deck.slide_layouts[6])
        output = BytesIO()
        deck.save(output)
        payload = output.getvalue()
        mime = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        monkeypatch.setattr('services.templates.scene_importer._convert_pptx',
            lambda content, count: ([image.getvalue()] * count, {'page_count': count,
                'engine_version': 'converter-fixture-only', 'warnings': []}))
    with pg.app.app_context():
        source = put_bytes(LOCAL_OWNER_ID, pg.project['project_id'], 'template_source',
            payload, mime, 'pptx' if source_type == 'pptx' else 'png')
        db.session.flush()
        document = TemplateDocument(id=str(uuid4()), owner_id=LOCAL_OWNER_ID,
            project_id=pg.project['project_id'], source_asset_id=source.id,
            source_type=source_type, status='uploaded')
        db.session.add(document)
        db.session.commit()
        result = import_document(document.id)
        db.session.commit()
        assert result['page_count'] == page_count and document.status == 'preview_ready'
        rows = ProjectTemplateAsset.query.filter_by(template_document_id=document.id).all()
        assert len(rows) == page_count
        for row in rows:
            assert checked_bytes(db.session.get(Asset, row.preview_asset_id))
            assert checked_bytes(db.session.get(Asset, row.thumbnail_asset_id))
        # Idempotent completed import does not create another set of references.
        assert import_document(document.id) == result
        db.session.commit()
        assert ProjectTemplateAsset.query.filter_by(template_document_id=document.id).count() == page_count
