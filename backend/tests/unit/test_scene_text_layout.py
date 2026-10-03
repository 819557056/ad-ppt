"""Derived layout contract, real Chromium, cache invalidation and OOXML consumers."""
import copy
import json
import os
import time
from io import BytesIO
from types import SimpleNamespace

import pytest
from flask import Flask
from pptx import Presentation
from pptx.util import Pt

from backend.tests.scene_layout_fixtures import single_line_layout, identity
from services.exports import scene_pdf
from services.exports.scene_pptx import render_pptx, verify_pptx
from services.scene import layout_contract as contract
from services.scene import render_jobs
from services.scene.validation import blank_scene, digest


def fixture(text='中文 English 123', *, vertical='top', weight=400, width=130):
    element = {'id': 'text', 'kind': 'text', 'role': 'body', 'text': text,
        'frame': {'x': 20, 'y': 20, 'w': width, 'h': 480, 'rotation_deg': 0},
        'style': {'padding_pt': 4, 'vertical_align': vertical, 'align': 'left',
                  'line_height': 1.2, 'font_size_pt': 20, 'font_weight': weight, 'color': '#17324D'}}
    scene = blank_scene()
    scene['elements'] = [element]
    snapshot = SimpleNamespace(owner_id='owner', project_id='project', manifest_json={
        'canvas': scene['canvas'], 'pages': [{'page_id': 'page', 'revision_id': 'revision'}]})
    return snapshot, {'revision': SimpleNamespace(scene_json=scene)}, element


@pytest.mark.parametrize('text', ['', ' ', '\n', 'A\n\nB\n', '中文 Á 🄀'])
def test_source_coverage_and_empty_lines(text):
    snapshot, revisions, _ = fixture(text)
    raw = single_line_layout(snapshot, revisions)
    assert contract.validate_layout(raw, snapshot, revisions) == raw
    assert raw['pages'][0]['text']['lines'][-1]['end'] == len(text)


@pytest.mark.parametrize('mutate', [
    lambda d: d.update(schema_version=True),
    lambda d: d.update(unexpected=True),
    lambda d: d.update(engine='forged'),
    lambda d: d['pages'][0].update(other=d['pages'][0]['text']),
    lambda d: d['pages'][0]['text'].update(source_length=99),
    lambda d: d['pages'][0]['text'].update(font_size_pt=19),
    lambda d: d['pages'][0]['text'].update(resolved_fonts=[]),
    lambda d: d['pages'][0]['text']['resolved_fonts'][0].update(is_custom_font=False),
    lambda d: d['pages'][0]['text']['resolved_fonts'][0].update(family_name='fallback'),
    lambda d: d['pages'][0]['text']['lines'][0].update(start=True),
    lambda d: d['pages'][0]['text']['lines'][0].update(end=2),
    lambda d: d['pages'][0]['text']['lines'][0].update(break_kind='hard'),
    lambda d: d['pages'][0]['text']['lines'][0].update(baseline_pt=float('nan')),
    lambda d: d['pages'][0]['text']['lines'][0]['line_box'].update(y=0),
    lambda d: d['pages'][0]['text']['lines'][0]['line_box'].update(w=1),
    lambda d: d.update(scene_hashes=['0' * 64]),
])
def test_tampering_rejected(mutate):
    snapshot, revisions, _ = fixture()
    raw = single_line_layout(snapshot, revisions)
    mutate(raw)
    with pytest.raises(ValueError):
        contract.validate_layout(raw, snapshot, revisions)


@pytest.mark.parametrize(('source', 'offset'), [('ÁB', 1), ('🇨🇳A', 1), ('👩\u200d💻A', 2), ('가A', 1), ('क्\u200dषA', 2)])
def test_never_split_grapheme_cluster(source, offset):
    snapshot, revisions, _ = fixture(source)
    raw = single_line_layout(snapshot, revisions)
    line = raw['pages'][0]['text']['lines'][0]
    second = copy.deepcopy(line)
    line.update(end=offset, break_kind='soft')
    second.update(start=offset)
    second['line_box']['y'] += 24
    second['baseline_pt'] += 24
    raw['pages'][0]['text']['lines'].append(second)
    with pytest.raises(ValueError, match='grapheme'):
        contract.validate_layout(raw, snapshot, revisions)


def test_lru_is_bounded_expires_and_returns_copies(monkeypatch):
    clock = [0.0]
    monkeypatch.setattr(contract.time, 'monotonic', lambda: clock[0])
    cache = contract.LayoutCache(max_bytes=40, max_entries=2, ttl=10)
    cache.put('a', {'a': [1]})
    cache.get('a')['a'].append(2)
    assert cache.get('a') == {'a': [1]}
    cache.put('b', {'b': [2]})
    cache.get('a')
    cache.put('c', {'c': [3]})
    assert cache.get('b') is None and cache.get('a') is not None
    cache.put('big', {'big': 'x' * 50})
    assert cache.get('big') is None and cache._bytes <= 40
    clock[0] = 11
    assert cache.get('a') is None and cache.get('c') is None and cache._bytes == 0


def test_cache_inputs_and_runtime_are_revalidated(monkeypatch):
    app = Flask(__name__)
    app.config['SCENE_RENDER_JOB_ROOT'] = 'synthetic-job-root'
    snapshot, revisions, element = fixture()
    engine, calls = identity(), []
    monkeypatch.setattr(scene_pdf, 'LAYOUT_CACHE', contract.LayoutCache())
    monkeypatch.setattr(scene_pdf, '_layout_identity', lambda: dict(engine))
    def render(kind, html):
        calls.append(kind)
        return json.dumps(single_line_layout(snapshot, revisions, dict(engine))).encode(), {}
    monkeypatch.setattr(render_jobs, 'run_render_job', render)
    with app.app_context():
        first = scene_pdf.measure_text_layout(snapshot, revisions, {})
        first['pages'].clear()
        assert scene_pdf.measure_text_layout(snapshot, revisions, {})['pages']
        assert len(calls) == 1
        element['text'] += '!'
        scene_pdf.measure_text_layout(snapshot, revisions, {})
        engine['script_sha256'] = 'b' * 64
        scene_pdf.measure_text_layout(snapshot, revisions, {})
        snapshot.owner_id = 'another-owner'
        scene_pdf.measure_text_layout(snapshot, revisions, {})
        snapshot.project_id = 'another-project'
        scene_pdf.measure_text_layout(snapshot, revisions, {})
        assert len(calls) == 5
        def unavailable():
            raise ValueError('renderer unavailable')
        monkeypatch.setattr(scene_pdf, '_layout_identity', unavailable)
        with pytest.raises(ValueError, match='unavailable'):
            scene_pdf.measure_text_layout(snapshot, revisions, {})
        assert len(calls) == 5


def test_engine_and_font_fingerprints_cover_compiler_and_manifest(tmp_path, monkeypatch):
    compiler = tmp_path / 'compiler.py'
    compiler.write_text('one')
    first = contract.engine_fingerprint(identity(), compiler)
    compiler.write_text('two')
    assert contract.engine_fingerprint(identity(), compiler) != first
    font = tmp_path / 'font.ttf'
    font.write_bytes(b'first')
    manifest = {'font_manifest_id': 'fonts-v1', 'font_file': 'font.ttf', 'sha256': 'invalid'}
    font.with_name('manifest.json').write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match='manifest mismatch'):
        contract.font_fingerprint(font)
    import hashlib
    manifest['sha256'] = hashlib.sha256(font.read_bytes()).hexdigest()
    font.with_name('manifest.json').write_text(json.dumps(manifest))
    before, _ = contract.font_fingerprint(font)
    manifest['family'] = 'new-family'
    font.with_name('manifest.json').write_text(json.dumps(manifest))
    after, _ = contract.font_fingerprint(font)
    assert before != after


@pytest.mark.parametrize('change', ['stale', 'unhealthy', 'identity', 'isolation'])
def test_renderer_status_blocks_stale_or_invalid_cache_identity(tmp_path, change):
    app = Flask(__name__)
    app.config.update(SCENE_RENDER_JOB_ROOT=str(tmp_path), SCENE_RENDER_UNIQUE_UID=True)
    status = {'updated_at': time.time(), 'chromium_ok': True, 'isolation_mode': 'unique_uid',
              'text_layout_identity': identity()}
    if change == 'stale':
        status['updated_at'] -= 60
    elif change == 'unhealthy':
        status['chromium_ok'] = False
    elif change == 'identity':
        status['text_layout_identity']['script_sha256'] = 'fake'
    else:
        status['isolation_mode'] = 'fixed_uid'
    (tmp_path / '.renderer_status.json').write_text(json.dumps(status))
    with app.app_context(), pytest.raises(ValueError):
        scene_pdf._layout_identity()


def test_identity_change_during_measurement_is_not_cached(monkeypatch):
    app = Flask(__name__)
    app.config['SCENE_RENDER_JOB_ROOT'] = 'synthetic-job-root'
    snapshot, revisions, _ = fixture()
    cache = contract.LayoutCache()
    monkeypatch.setattr(scene_pdf, 'LAYOUT_CACHE', cache)
    monkeypatch.setattr(scene_pdf, '_layout_identity', identity)
    raw = single_line_layout(snapshot, revisions)
    raw['engine_identity']['script_sha256'] = 'f' * 64
    monkeypatch.setattr(render_jobs, 'run_render_job', lambda *a: (json.dumps(raw).encode(), {}))
    with app.app_context(), pytest.raises(ValueError, match='identity changed'):
        scene_pdf.measure_text_layout(snapshot, revisions, {})
    assert not cache._entries


@pytest.mark.skipif(not os.environ.get('SCENE_CHROMIUM_PATH'), reason='Real Chromium opt-in')
@pytest.mark.parametrize(('text', 'vertical', 'weight'), [
    ('中文 English 123\n第二行继续文本', 'top', 400), ('', 'top', 400), ('   ', 'middle', 400),
    ('\n', 'bottom', 400), ('A\n\nB\n', 'middle', 700), ('Á🄀B 中文混排 Á🄀B 中文', 'bottom', 700),
])
def test_real_chromium_lines_baselines_fonts_and_native_pptx(text, vertical, weight):
    app = Flask(__name__)
    snapshot, revisions, element = fixture(text, vertical=vertical, weight=weight)
    before = digest(revisions['revision'].scene_json)
    with app.app_context():
        measured = scene_pdf.measure_text_layout(snapshot, revisions, {})
        data = measured['pages'][0]['text']
        assert measured['layout_engine_version'] and measured['font_manifest_hash']
        assert all(line['baseline_pt'] > line['line_box']['y'] for line in data['lines'])
        assert all(font['is_custom_font'] for font in data['resolved_fonts'])
        for line in data['lines']:
            if line['break_kind'] == 'hard':
                assert text[line['end']] == '\n'
        payload = render_pptx(snapshot, revisions, {}, measured)
        assert verify_pptx(payload, snapshot, revisions, measured)['native_text_count'] == 1
        shape = Presentation(BytesIO(payload)).slides[0].shapes[0]
        assert shape.text.replace('\v', '') == text
        assert shape.text_frame.word_wrap is False
        assert all(p.line_spacing == Pt(24) for p in shape.text_frame.paragraphs)
        pdf = scene_pdf.render_pdf(snapshot, revisions, {}, measured)
        assert scene_pdf.verify_pdf(pdf, snapshot, revisions)['searchable_text']
        assert digest(revisions['revision'].scene_json) == before


@pytest.mark.skipif(not os.environ.get('SCENE_CHROMIUM_PATH'), reason='Real Chromium opt-in')
def test_real_chromium_font_fallback_rejected():
    snapshot, revisions, _ = fixture('A\U00020000')  # Not in the fixed font's cmap.
    with Flask(__name__).app_context(), pytest.raises(ValueError, match='font fallback'):
        scene_pdf.measure_text_layout(snapshot, revisions, {})

@pytest.mark.parametrize(('source', 'exported'), [('AB 123', 'AB123'), ('AB 123', 'A B 123'), ('Á', 'Á')])
def test_pdf_verification_never_ignores_spaces_or_normalizes_unicode(source, exported):
    import fitz
    snapshot, revisions, _ = fixture(source)
    with fitz.open() as document:
        page = document.new_page(width=960, height=540)
        page.insert_text((24, 50), exported, fontsize=20)
        with pytest.raises(ValueError, match='character count mismatch'):
            scene_pdf.verify_pdf(document.tobytes(), snapshot, revisions)
