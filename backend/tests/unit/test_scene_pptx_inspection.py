"""Static PPTX package checks run before an isolated renderer sees uploads."""
from io import BytesIO
import warnings
import zipfile

import pytest
from PIL import Image
from pptx import Presentation

from services.scene.versioning import SceneError
from services.templates.scene_importer import inspect_pptx, inspect_reference_image


def _base_pptx():
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _append_entry(payload, name, data, *, symlink=False):
    output = BytesIO(payload)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)  # Deliberate duplicate entry case.
        with zipfile.ZipFile(output, 'a') as archive:
            if symlink:
                entry = zipfile.ZipInfo(name)
                entry.create_system = 3
                entry.external_attr = 0o120777 << 16
                archive.writestr(entry, data)
            else:
                archive.writestr(name, data)
    return output.getvalue()


def test_pptx_inspection_parses_external_relationship_with_single_quotes(app):
    relationship = (b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
                    b"<Relationship Id='rId9' Type='x' Target='https://fixture.invalid/' "
                    b"TargetMode = 'External'/></Relationships>")
    payload = _append_entry(_base_pptx(), 'ppt/slides/_rels/extra.xml.rels', relationship)
    with app.app_context():
        count, found = inspect_pptx(payload)
    assert count == 1
    assert 'EXTERNAL_RELATIONSHIP_IGNORED' in found


def test_pptx_inspection_refuses_xml_entities_before_conversion(app):
    malicious = (b'<!DOCTYPE Relationships [<!ENTITY copied "expanded">]>'
                 b'<Relationships>&copied;</Relationships>')
    payload = _append_entry(_base_pptx(), 'ppt/slides/_rels/extra.xml.rels', malicious)
    with app.app_context(), pytest.raises(SceneError) as raised:
        inspect_pptx(payload)
    assert raised.value.code == 'TEMPLATE_INVALID'


@pytest.mark.parametrize(('entry', 'data', 'symlink'), [
    ('ppt/vbaProject.bin', b'macro', False),
    ('ppt/activeX/activeX1.bin', b'control', False),
    ('ppt/presentation.xml', b'duplicate', False),
    ('PPT/Presentation.xml', b'case collision', False),
    ('ppt/linked.xml', b'target', True),
    ('ppt/../escape.xml', b'escape', False),
    ('ppt//ambiguous.xml', b'ambiguous', False),
])
def test_pptx_inspection_rejects_active_or_unsafe_entries(app, entry, data, symlink):
    payload = _append_entry(_base_pptx(), entry, data, symlink=symlink)
    with app.app_context(), pytest.raises(SceneError) as raised:
        inspect_pptx(payload)
    assert raised.value.code == 'TEMPLATE_INVALID'


def test_pptx_inspection_warns_about_embedded_objects(app):
    payload = _append_entry(_base_pptx(), 'ppt/embeddings/oleObject1.bin', b'fixture')
    with app.app_context():
        assert inspect_pptx(payload) == (1, ['EMBEDDED_OBJECT_STATIC_ONLY'])


def test_reference_image_inspection_sniffs_bytes_and_rejects_animation():
    source = Image.new('RGB', (8, 8), 'navy')
    buffer = BytesIO()
    source.save(buffer, format='JPEG')
    assert inspect_reference_image(buffer.getvalue()) == ('image/jpeg', 'jpg')
    buffer = BytesIO()
    source.save(buffer, format='GIF', save_all=True,
                append_images=[Image.new('RGB', (8, 8), 'red')], duration=100)
    with pytest.raises(SceneError) as raised:
        inspect_reference_image(buffer.getvalue())
    assert raised.value.code == 'TEMPLATE_INVALID'
