"""The PPTX quality gate must reject visually different but structurally valid decks."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from services.exports.scene_pptx import render_pptx, verify_pptx


def _fixture():
    first = {'id': 'title', 'kind': 'text', 'role': 'title', 'text': 'First',
             'frame': {'x': 20, 'y': 20, 'w': 400, 'h': 80, 'rotation_deg': 0},
             'style': {'padding_pt': 0, 'vertical_align': 'top', 'align': 'left',
                       'line_height': 1.2, 'font_size_pt': 32, 'font_weight': 700,
                       'color': '#17324D'}}
    second = {**first, 'id': 'subtitle', 'text': 'Second',
              'frame': {'x': 20, 'y': 100, 'w': 400, 'h': 80, 'rotation_deg': 0}}
    snapshot = SimpleNamespace(manifest_json={
        'canvas': {'width_pt': 960, 'height_pt': 540},
        'pages': [{'page_id': 'page', 'revision_id': 'revision'}]})
    revisions = {'revision': SimpleNamespace(scene_json={
        'background': {'kind': 'solid', 'color': '#FFFFFF'},
        'elements': [first, second]})}
    return snapshot, revisions


def _changed(payload, change):
    deck = Presentation(BytesIO(payload))
    change(deck)
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _reverse_shapes(deck):
    tree = deck.slides[0].shapes._spTree
    first = deck.slides[0].shapes[0]._element
    tree.remove(first)
    tree.append(first)


def test_pptx_quality_gate_checks_canvas_background_rotation_and_z_order():
    snapshot, revisions = _fixture()
    original = render_pptx(snapshot, revisions, {})
    assert verify_pptx(original, snapshot, revisions)['native_text_count'] == 2

    def cases():
        yield 'canvas size', lambda deck: setattr(deck, 'slide_width', Pt(1000))
        yield 'solid background', lambda deck: setattr(
            deck.slides[0].background.fill.fore_color, 'rgb', RGBColor(0, 0, 0))
        yield 'text rotation', lambda deck: setattr(deck.slides[0].shapes[0], 'rotation', 12)
        yield 'shape order', _reverse_shapes

    for message, mutate in cases():
        with pytest.raises(ValueError, match=message):
            verify_pptx(_changed(original, mutate), snapshot, revisions)


@pytest.mark.parametrize('isolated', [True, False])
def test_text_layout_only_requires_local_script_for_local_execution(monkeypatch, tmp_path, isolated):
    """Production backend intentionally excludes the renderer JS/build toolchain."""
    import json
    from flask import Flask
    from services.exports import scene_pdf
    from services.scene import render_jobs

    app = Flask(__name__)
    app.config['SCENE_RENDER_JOB_ROOT'] = str(tmp_path / 'jobs') if isolated else None
    snapshot, revisions = _fixture()
    monkeypatch.setattr(scene_pdf, 'LAYOUT_RENDERER', tmp_path / 'not-in-backend-image.mjs')
    calls = []
    from backend.tests.scene_layout_fixtures import single_line_layout
    result = single_line_layout(snapshot, revisions)
    if isolated:
        import time
        jobs = Path(app.config['SCENE_RENDER_JOB_ROOT'])
        jobs.mkdir()
        (jobs / '.renderer_status.json').write_text(json.dumps({'updated_at': time.time(), 'chromium_ok': True,
            'text_layout_identity': result['engine_identity']}))
    def render(kind, payload):
        calls.append((kind, payload))
        return json.dumps(result).encode(), {}
    monkeypatch.setattr(render_jobs, 'run_render_job', render)
    with app.app_context():
        if isolated:
            assert scene_pdf.measure_text_layout(snapshot, revisions, {})['pages'] == result['pages']
            assert calls[0][0] == 'text_layout' and b'First' in calls[0][1]
        else:
            with pytest.raises(ValueError, match='text layout renderer unavailable'):
                scene_pdf.measure_text_layout(snapshot, revisions, {})
            assert not calls
