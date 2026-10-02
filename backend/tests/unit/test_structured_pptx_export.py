"""Structured editable export and PPTX quality checks need no OCR services."""

from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from models import db, Page, Project
from services.structured_pptx_service import create_structured_pptx, inspect_pptx


def _page(page_id, title, points, index=0):
    return SimpleNamespace(
        id=page_id,
        part=None,
        get_outline_content=lambda: {"title": title, "points": points},
        get_description_content=lambda: {},
    )


def test_structured_export_contains_native_editable_objects(tmp_path):
    project = SimpleNamespace(project_title="Quarterly review", image_aspect_ratio="16:9")
    output = tmp_path / "deck.pptx"
    plans = create_structured_pptx(project, [
        _page("one", "Overview", ["Revenue grew", "Costs stabilized"]),
        _page("two", "Next steps", ["Ship the feature"]),
    ], output)

    report = inspect_pptx(output, expected_slides=2, require_editable=True)
    assert len(plans) == 2
    assert report["slide_count"] == 2
    assert report["editable_text_shapes"] >= 6
    assert report["native_shape_count"] > 0
    assert report["image_shapes"] == 0
    presentation = Presentation(output)
    assert any(shape.has_text_frame and "Revenue grew" in shape.text
               for shape in presentation.slides[0].shapes)


def test_quality_rejects_corrupt_package(tmp_path):
    output = tmp_path / "bad.pptx"
    output.write_bytes(b"not a pptx")
    with pytest.raises(ValueError, match="valid PPTX"):
        inspect_pptx(output)


def test_structured_export_rejects_too_many_points_instead_of_dropping_them(tmp_path):
    project = SimpleNamespace(project_title="Dense deck", image_aspect_ratio="16:9")
    with pytest.raises(ValueError, match="at most 8"):
        create_structured_pptx(project, [_page("one", "Dense", [str(n) for n in range(9)])],
                               tmp_path / "dense.pptx")


def test_structured_route_exports_selected_pages_without_images(client, app):
    with app.app_context():
        project = Project(creation_type="idea", project_title="Native deck")
        db.session.add(project)
        db.session.flush()
        first = Page(project_id=project.id, order_index=0)
        first.set_outline_content({"title": "First", "points": ["One"]})
        second = Page(project_id=project.id, order_index=1)
        second.set_outline_content({"title": "Second", "points": ["Two"]})
        db.session.add_all([first, second])
        db.session.commit()
        project_id, second_id = project.id, second.id

    response = client.get(
        f"/api/projects/{project_id}/export/structured-pptx?page_ids={second_id}"
    )
    assert response.status_code == 200
    data = response.get_json()["data"]
    assert data["method"] == "structured"
    assert data["quality"]["slide_count"] == 1
    assert data["quality"]["image_shapes"] == 0
    filename = data["download_url"].rsplit("/", 1)[-1]
    output = Path(app.config["UPLOAD_FOLDER"]) / project_id / "exports" / filename
    presentation = Presentation(output)
    texts = [shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame]
    assert any("Second" in text for text in texts)
    assert not any("First" in text for text in texts)

    quality_response = client.get(
        f"/api/projects/{project_id}/export/pptx-quality?filename={filename}"
    )
    assert quality_response.status_code == 200
    assert quality_response.get_json()["data"]["slide_count"] == 1
