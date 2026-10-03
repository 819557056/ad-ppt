"""Renderer font-risk diagnostics are advisory, never a conversion verdict."""
from zipfile import ZipFile
from io import BytesIO
from uuid import uuid4

from PIL import Image
from pptx import Presentation

from renderer.convert_template import font_warnings
from renderer.daemon import _result_metadata
from models import db
from models.scene_v1 import TemplateDocument, new_id
from services.scene.store import put_bytes
from services.scene.versioning import LOCAL_OWNER_ID
from services.templates.scene_importer import _convert_pptx, import_document


def _pptx_with_fonts(path, *families):
    runs = ''.join(f'<a:rPr><a:latin typeface="{name}"/></a:rPr>' for name in families)
    xml = (f'<root xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           f'{runs}</root>')
    with ZipFile(path, 'w') as archive:
        archive.writestr('ppt/slides/slide1.xml', xml)


def test_renderer_flags_only_uninstalled_explicit_slide_fonts(tmp_path, monkeypatch):
    source = tmp_path / 'reference.pptx'
    _pptx_with_fonts(source, 'Noto Sans CJK SC', '+mn-lt', 'Missing Fixture Font')
    monkeypatch.setattr('renderer.convert_template._installed_font_families',
                        lambda: {'noto sans cjk sc'})
    assert font_warnings(source) == ['FONT_FAMILY_UNAVAILABLE']
    monkeypatch.setattr('renderer.convert_template._installed_font_families',
                        lambda: {'noto sans cjk sc', 'missing fixture font'})
    assert font_warnings(source) == []
    monkeypatch.setattr('renderer.convert_template._installed_font_families', lambda: None)
    assert font_warnings(source) == ['FONT_AVAILABILITY_UNVERIFIED']


def test_renderer_and_importer_keep_only_known_font_warning_codes(app, monkeypatch):
    metadata = _result_metadata('pptx', 'LibreOffice 24.8\n'
        'SCENE_RENDER_WARNINGS=FONT_FAMILY_UNAVAILABLE,INJECTED\n')
    assert metadata == {'engine_version': 'LibreOffice 24.8',
                        'warnings': ['FONT_FAMILY_UNAVAILABLE']}
    monkeypatch.setattr('services.scene.render_jobs.run_render_job',
                        lambda *_args, **_kwargs: ([b'fixture'], {
                            'page_count': 1, 'engine_version': 'LibreOffice 24.8',
                            'warnings': ['FONT_FAMILY_UNAVAILABLE', 'INJECTED', {'bad': 'value'}]}))
    with app.app_context():
        monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', '/fixture-jobs')
        monkeypatch.setitem(app.config, 'SCENE_RENDER_UNIQUE_UID', True)
        pages, manifest = _convert_pptx(b'fixture', 1)
    assert pages == [b'fixture']
    assert manifest['warnings'] == ['FONT_FAMILY_UNAVAILABLE']


def test_imported_document_exposes_renderer_font_warning(client, app, tmp_path, monkeypatch):
    monkeypatch.setitem(app.config, 'SCENE_EDITOR_ENABLED', True)
    monkeypatch.setitem(app.config, 'ASSET_STORE_ROOT', str(tmp_path / 'assets'))
    monkeypatch.setitem(app.config, 'SCENE_RENDER_JOB_ROOT', str(tmp_path / 'jobs'))
    monkeypatch.setitem(app.config, 'SCENE_RENDER_UNIQUE_UID', True)
    created = client.post('/api/v2/projects', json={'title': 'Font fixture',
        'pages': [{'title': 'Page', 'points': []}]}, headers={'Idempotency-Key': str(uuid4())})
    assert created.status_code == 201
    project_id = created.json['data']['project_id']
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    pptx = BytesIO(); deck.save(pptx)
    image = BytesIO(); Image.new('RGB', (64, 36), 'navy').save(image, format='PNG')
    monkeypatch.setattr('services.scene.render_jobs.run_render_job',
                        lambda *_args, **_kwargs: ([image.getvalue()], {
                            'page_count': 1, 'engine_version': 'LibreOffice fixture',
                            'warnings': ['FONT_FAMILY_UNAVAILABLE']}))
    with app.app_context():
        source = put_bytes(LOCAL_OWNER_ID, project_id, 'template_source', pptx.getvalue(),
                           'application/vnd.openxmlformats-officedocument.presentationml.presentation', 'pptx')
        document = TemplateDocument(id=new_id(), owner_id=LOCAL_OWNER_ID, project_id=project_id,
                                    source_asset_id=source.id, source_type='pptx', status='uploaded')
        db.session.add(document); db.session.commit()
        import_document(document.id)
        db.session.commit()
        assert document.warnings_json == ['FONT_FAMILY_UNAVAILABLE']
        assert document.renderer_manifest_json['engine_version'] == 'LibreOffice fixture'
